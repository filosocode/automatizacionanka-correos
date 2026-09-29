"""Extraccion local de texto desde adjuntos (sin llamar al modelo)."""

from __future__ import annotations

import html
import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_ATTACHMENT_TEXT = 30000
TEXT_ATTACHMENT_EXTENSIONS = {
    ".csv", ".geojson", ".json", ".kml", ".md", ".tab", ".tsv", ".txt", ".xml",
}


def _kmz_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        kml_files = [name for name in archive.namelist() if name.casefold().endswith(".kml")]
        if not kml_files:
            raise ValueError("el KMZ no contiene un archivo KML")
        summaries = []
        for name in kml_files:
            raw = archive.read(name).decode("utf-8", errors="replace")
            feature_names = [
                html.unescape(re.sub(r"<[^>]+>", "", value)).strip()
                for value in re.findall(r"<name>(.*?)</name>", raw, re.IGNORECASE | re.DOTALL)
            ]
            coordinates = []
            for block in re.findall(
                r"<coordinates>(.*?)</coordinates>", raw, re.IGNORECASE | re.DOTALL
            ):
                for value in block.split():
                    parts = value.split(",")
                    if len(parts) < 2:
                        continue
                    try:
                        coordinates.append((float(parts[0]), float(parts[1])))
                    except ValueError:
                        continue
            lines = [f"Archivo KML: {name}"]
            if feature_names:
                lines.append("Elementos: " + "; ".join(dict.fromkeys(feature_names))[:3000])
            lines.append(f"Cantidad de coordenadas: {len(coordinates)}")
            if coordinates:
                west = min(point[0] for point in coordinates)
                south = min(point[1] for point in coordinates)
                east = max(point[0] for point in coordinates)
                north = max(point[1] for point in coordinates)
                lines.append(
                    "Extensión geográfica (oeste, sur, este, norte): "
                    f"{west}, {south}, {east}, {north}"
                )
            summaries.append("\n".join(lines))
        return "\n\n".join(summaries)


def extract_attachment_text(filename: str, mime_type: str, content: bytes) -> tuple[str, str]:
    """Devuelve texto y motivo de bloqueo; un motivo vacío significa extracción exitosa."""
    if len(content) > MAX_ATTACHMENT_BYTES:
        return "", f"supera el límite de {MAX_ATTACHMENT_BYTES // (1024 * 1024)} MB"
    extension = Path(filename).suffix.casefold()
    try:
        if extension in TEXT_ATTACHMENT_EXTENSIONS or mime_type.startswith("text/"):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
        elif extension == ".kmz" or mime_type == "application/vnd.google-earth.kmz":
            text = _kmz_text(content)
        elif extension == ".pdf" or mime_type == "application/pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return "", "falta instalar pypdf para leer PDF"
            page_texts = [
                page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages
            ]
            if not any(page.strip() for page in page_texts):
                return "", "PDF sin texto extraíble; puede ser un documento escaneado"
            text = "\n".join(
                f"[Página {number}]\n{page}"
                for number, page in enumerate(page_texts, 1) if page.strip()
            )
        elif extension == ".docx":
            with zipfile.ZipFile(BytesIO(content)) as document:
                root = ElementTree.fromstring(document.read("word/document.xml"))
            text = "\n".join(node.text or "" for node in root.iter()
                             if node.tag.endswith("}t"))
        elif extension in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError:
                return "", "falta instalar openpyxl para leer Excel"
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets:
                rows.append(f"Hoja: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    if any(value is not None for value in row):
                        rows.append("\t".join("" if value is None else str(value) for value in row))
            text = "\n".join(rows)
        elif extension in {".dwg", ".dxf"}:
            return "", "plano CAD requiere revisión manual"
        else:
            return "", f"formato no soportado ({extension or mime_type or 'desconocido'})"
    except Exception as exc:
        return "", f"no se pudo leer: {exc}"
    text = text.strip()
    if not text:
        return "", "el archivo no contiene texto extraíble"
    return text[:MAX_ATTACHMENT_TEXT], ""
