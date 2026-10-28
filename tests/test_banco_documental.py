"""Pruebas del banco de evaluación documental (T-71): documentos, referencias, método de comparación y aislamiento."""

from __future__ import annotations

import copy
import io
import json
import re
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from laboratorio.banco_documental import comparador, construir, datos, entrega, generar

_PDF = "application/pdf"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CORREO_O_TELEFONO = re.compile(r"[\w.+-]+@(?!ejemplo\.invalid)[\w-]+\.\w+|\b3\d{9}\b|\b\d{7,}\b")


def _bytes_de_documentos() -> dict[str, bytes]:
    return {nombre: (construir.DOCUMENTOS / nombre).read_bytes() for nombre in datos.ARCHIVOS.values()}


def _referencias() -> dict[str, dict]:
    return {caso: json.loads((construir.REFERENCIAS / f"{caso}.referencia.json").read_text(encoding="utf-8"))
            for caso in datos.CASOS}


def _extraccion_perfecta(referencia: dict) -> dict:
    """Lo que entregaría un lector impecable: sale de la referencia, solo sirve para probar el comparador."""
    return {
        "caso": referencia["caso"],
        "actividades": [{"nombre": a["nombre"], "cantidad": a["cantidad_literal"], "unidad": a["unidad_literal"],
                         "entregable": a["entregable"], "origen": {k: v for k, v in a["origen"].items() if k != "documento"}}
                        for a in referencia["actividades"]],
        "entregables": [a["entregable"] for a in referencia["actividades"]],
        "restricciones": [r["texto"] for r in referencia["restricciones"]],
        "contradicciones": [f"{c['motivo']} ({' / '.join(c['literales_clave'])})" for c in referencia["contradicciones"]],
        "preguntas_cliente": [f["motivo"] for f in referencia["faltantes"]],
        "revisiones_internas": [r["texto"] for r in referencia["revisiones_internas"]],
        "elementos_visuales": [e["descripcion"] for e in referencia["elementos_visuales"]],
    }


class BancoDocumentalTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.documentos = _bytes_de_documentos()
        cls.referencias = _referencias()

    # Cubre: T-71
    def test_01_los_documentos_son_legibles_y_validos(self) -> None:
        from openpyxl import load_workbook
        from pypdf import PdfReader

        for nombre in ("solicitud_01.pdf", "tdr_02.pdf"):
            self.assertEqual(len(PdfReader(io.BytesIO(self.documentos[nombre])).pages), 2, nombre)
        libro = load_workbook(io.BytesIO(self.documentos["cantidades_03.xlsx"]))
        self.assertEqual(libro.sheetnames, ["Resumen", "Detalle", "Supuestos", "Borrador_anterior"])

    # Cubre: T-71
    def test_02_los_manifiestos_coinciden_con_los_archivos(self) -> None:
        manifiesto = json.loads((construir.MANIFIESTO_DOCUMENTOS).read_text(encoding="utf-8"))
        for entrada in manifiesto["documentos"]:
            self.assertEqual(entrada["sha256"], construir.sha256(self.documentos[entrada["archivo"]]))
            self.assertEqual(entrada["bytes"], len(self.documentos[entrada["archivo"]]))
        enlaces = json.loads((construir.MANIFIESTO_REFERENCIAS).read_text(encoding="utf-8"))["referencias"]
        for entrada in enlaces:
            ruta = construir.REFERENCIAS / entrada["archivo"]
            self.assertEqual(entrada["sha256"], construir.sha256(ruta.read_bytes()))
            self.assertEqual(entrada["documento_sha256"], construir.sha256(self.documentos[entrada["documento"]]))
            referencia = json.loads(ruta.read_text(encoding="utf-8"))
            self.assertEqual(referencia["documento"], {"archivo": entrada["documento"],
                                                       "sha256": entrada["documento_sha256"]})

    # Cubre: T-71
    def test_03_la_generacion_es_reproducible_byte_a_byte(self) -> None:
        primera = construir.construir_en_memoria()
        segunda = construir.construir_en_memoria()
        self.assertEqual(primera, segunda)
        self.assertEqual(construir.verificar(), [], "lo guardado debe coincidir con lo que se genera hoy")

    # Cubre: T-71
    def test_04_las_cifras_y_unidades_del_pdf_digital_estan_donde_dice_la_referencia(self) -> None:
        from anka.infra.attachments import extract_attachment_text

        texto, motivo = extract_attachment_text("solicitud_01.pdf", _PDF, self.documentos["solicitud_01.pdf"])
        self.assertEqual(motivo, "")
        paginas = {int(n): t for n, t in re.findall(r"\[Página (\d+)\]\n(.*?)(?=\n\[Página |\Z)", texto, re.S)}
        plano = re.sub(r"\s+", " ", paginas[1])
        for actividad in self.referencias["C1"]["actividades"]:
            self.assertEqual(actividad["origen"]["pagina"], 1)
            self.assertRegex(plano, rf"{re.escape(actividad['cantidad_literal'])} {re.escape(actividad['unidad_literal'])}\b")
        for restriccion in self.referencias["C1"]["restricciones"]:
            pagina = restriccion["origen"]["pagina"]
            self.assertIn(restriccion["claves"][0].lower(), re.sub(r"\s+", " ", paginas[pagina]).lower(), restriccion["id"])
        for trampa in self.referencias["C1"]["trampas"]:
            self.assertIn(trampa["valor_literal"], re.sub(r"\s+", " ", paginas[trampa["origen"]["pagina"]]))
        self.assertIn("se enviará por separado", plano)        # el faltante bloqueante está en el documento

    # Cubre: T-71
    def test_05_el_pdf_escaneado_no_tiene_texto_y_requiere_ocr(self) -> None:
        import pypdfium2
        from pypdf import PdfReader

        lector = PdfReader(io.BytesIO(self.documentos["tdr_02.pdf"]))
        for pagina in lector.pages:
            self.assertEqual((pagina.extract_text() or "").strip(), "")
            self.assertEqual(len(list(pagina.images)), 1)
        pdf = pypdfium2.PdfDocument(self.documentos["tdr_02.pdf"])
        for indice in range(len(pdf)):
            self.assertEqual(pdf[indice].get_textpage().count_chars(), 0)
        self.assertTrue(self.referencias["C2"]["ocr"]["requiere"])
        self.assertTrue(all(a["requiere_ocr"] for a in self.referencias["C2"]["actividades"]))

    # Cubre: T-71
    def test_06_el_ocr_local_recupera_las_cifras_clave_y_marca_lo_visual(self) -> None:
        from anka.infra.attachments import extract_attachment_text

        texto, motivo = extract_attachment_text("tdr_02.pdf", _PDF, self.documentos["tdr_02.pdf"])
        self.assertIn("COBERTURA VISUAL PARCIAL", texto)
        self.assertIn("inspección visual", motivo)
        for esperado in ("45,0", "41,8", "6,8", "3,2", "0457"):
            self.assertIn(esperado, texto)
        # lo que el OCR no puede hacer: la tachadura y el área rayada son interpretación visual
        self.assertNotIn("tachad", texto.lower())
        self.assertNotIn("rayad", texto.lower())

    # Cubre: T-71
    def test_07_el_excel_conserva_hojas_formulas_unidades_y_localizadores(self) -> None:
        from openpyxl import load_workbook

        contenido = self.documentos["cantidades_03.xlsx"]
        formulas = load_workbook(io.BytesIO(contenido))
        valores = load_workbook(io.BytesIO(contenido), data_only=True)
        self.assertEqual(formulas["Borrador_anterior"].sheet_state, "hidden")
        self.assertIsNotNone(formulas["Resumen"]["C9"].comment)
        for actividad in self.referencias["C3"]["actividades"]:
            origen = actividad["origen"]
            celda = origen["celda"]
            hoja = origen["hoja"]
            self.assertEqual(formulas[hoja][celda].value, actividad["formula"])
            self.assertEqual(Decimal(str(valores[hoja][celda].value)), Decimal(actividad["valor"]), actividad["id"])
            self.assertEqual(valores[hoja][f"D{celda[1:]}"].value, actividad["unidad_literal"])
        self.assertEqual(valores["Detalle"]["C8"].value, 337000)             # en m2: no se convierte
        self.assertEqual(valores["Detalle"]["D8"].value, 33.7)
        self.assertEqual(valores["Borrador_anterior"]["C2"].value, 31.2)
        self.assertEqual(valores["Resumen"]["C9"].value, 7)
        self.assertEqual(sum(1 for fila in range(2, 8) if valores["Detalle"][f"A{fila}"].value), 6)

    # Cubre: T-71
    def test_08_el_lector_local_ve_formulas_valores_ocultas_y_comentarios(self) -> None:
        from anka.infra.attachments import extract_attachment_text

        texto, _ = extract_attachment_text("cantidades_03.xlsx", _XLSX, self.documentos["cantidades_03.xlsx"])
        for esperado in ("[valor guardado: 33.7]", "[valor guardado: 3.625]", "[valor guardado: 24]", "HOJA OCULTA",
                         "Pendiente de confirmar con el cliente", "C8==SUM(C2:C7) [valor guardado: 337000]"):
            self.assertIn(esperado, texto)

    # Cubre: T-71
    def test_09_los_valores_calculados_se_recalculan_con_los_datos(self) -> None:
        calculados = generar.valores_calculados()
        self.assertEqual(calculados[("Resumen", "C4")], "33.7")
        self.assertEqual(calculados[("Resumen", "C5")], "3.625")
        self.assertEqual(calculados[("Resumen", "C6")], "24")
        self.assertEqual(sum(Decimal(a) for _, _, a in datos.DETALLE_LOTES), Decimal(337000))

    # Cubre: T-71
    def test_10_cada_referencia_corresponde_a_su_documento_y_tiene_lo_exigido(self) -> None:
        for caso, referencia in self.referencias.items():
            self.assertEqual(referencia["documento"]["archivo"], datos.ARCHIVOS[caso])
            self.assertTrue(referencia["ficticio"])
            self.assertTrue(referencia["actividades"] and referencia["trampas"] and referencia["contradicciones"])
            self.assertTrue(referencia["inventario"])
            self.assertIn("ocr", referencia)
            for actividad in referencia["actividades"]:
                self.assertEqual(actividad["origen"]["documento"], datos.ARCHIVOS[caso])
                self.assertRegex(actividad["cantidad_literal"], r"^\d[\d.,]*$")
                self.assertTrue(actividad["unidad_literal"])
            bloqueantes = [f for f in referencia["faltantes"] if f["bloqueante"]]
            self.assertEqual(len(bloqueantes) > 0, caso != "C3")
            for faltante in referencia["faltantes"]:
                self.assertIn(faltante["destino"], {"cliente", "revision_tecnica_interna"})
            self.assertTrue(referencia["revisiones_internas"], "debe separar revisiones técnicas internas")
        self.assertTrue(any(a.get("ambiguo") for a in self.referencias["C1"]["actividades"]))
        self.assertTrue(self.referencias["C2"]["elementos_visuales"])

    # Cubre: T-71
    def test_11_la_referencia_t34_es_compatible_con_la_evaluacion_existente(self) -> None:
        from anka.dominio import evaluacion

        for referencia in self.referencias.values():
            esperado = referencia["esperado"]
            self.assertLessEqual(set(esperado), set(evaluacion.Esperado.__annotations__))
            self.assertTrue(set(esperado["faltante"]) <= set(evaluacion.TEMAS))

    # Cubre: T-71
    def test_12_no_hay_datos_reales_y_todo_se_declara_ficticio(self) -> None:
        textos = json.dumps({"c1": datos.C1, "c2": datos.C2, "c3": datos.C3}, ensure_ascii=False)
        self.assertEqual(_CORREO_O_TELEFONO.findall(textos), [])
        self.assertIn("ficticia", datos.C1["parametros"]["entidad"])
        self.assertIn("ficticia", datos.C2["parametros"]["entidad"])
        self.assertIn("ficticio", datos.C3["parametros"]["proyecto"])

    # Cubre: T-71
    def test_13_los_nombres_y_metadatos_de_los_documentos_no_dan_pistas(self) -> None:
        for nombre in self.documentos:
            self.assertNotRegex(nombre.lower(), r"referencia|respuesta|esperado|correct|solucion|banco")
        for nombre, contenido in self.documentos.items():
            self.assertIsNone(entrega._ID_DE_REFERENCIA.search(contenido), nombre)
            for pista in entrega._PISTAS:
                self.assertNotIn(pista, contenido, nombre)


class ComparadorTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.referencias = _referencias()

    # Cubre: T-71
    def test_20_una_extraccion_perfecta_no_tiene_errores_criticos(self) -> None:
        for caso, referencia in self.referencias.items():
            resultado = comparador.comparar(referencia, _extraccion_perfecta(referencia))
            criticos = [h for h in resultado["hallazgos"] if h["gravedad"] == comparador.CRITICO]
            self.assertEqual(criticos, [], caso)
            self.assertEqual(resultado["actividades_con_valor_y_unidad_correctos"], len(referencia["actividades"]))
            self.assertEqual(resultado["errores_no_criticos"], 0, [h for h in resultado["hallazgos"] if h["gravedad"] == "no_critico"])
        self.assertEqual(comparador.comparar(self.referencias["C2"], _extraccion_perfecta(self.referencias["C2"]))["no_evaluables"], 3)

    # Cubre: T-71
    def test_21_cumplir_el_esquema_no_basta_una_unidad_cambiada_es_critica(self) -> None:
        referencia = self.referencias["C3"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["actividades"][0].update(cantidad="337000", unidad="ha")      # número de m2 como si fueran ha
        resultado = comparador.comparar(referencia, extraccion)
        self.assertTrue(any(h["codigo"] == "asociacion_incorrecta" and h["ref"] == "C3-A1" for h in resultado["hallazgos"]))
        extraccion["actividades"][0].update(cantidad="33,7", unidad="m2")        # valor correcto con unidad cambiada
        resultado = comparador.comparar(referencia, extraccion)
        self.assertTrue(any(h["codigo"] == "cantidad_o_unidad_incorrecta" for h in resultado["hallazgos"]))
        self.assertFalse(resultado["sin_errores_criticos"])

    # Cubre: T-71
    def test_22_no_se_aceptan_conversiones_implicitas(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["actividades"][0].update(cantidad="185000", unidad="m2")      # 18,5 ha convertidas
        resultado = comparador.comparar(referencia, extraccion)
        self.assertEqual(resultado["actividades_con_valor_y_unidad_correctos"], len(referencia["actividades"]) - 1)
        self.assertEqual(resultado["errores_criticos"], 1)

    # Cubre: T-71
    def test_23_la_cifra_trampa_atada_al_concepto_equivocado_es_critica(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["actividades"][2]["cantidad"] = "12"                            # correcto
        extraccion["actividades"][0]["cantidad"] = "21,0"                           # área del predio como intervención
        resultado = comparador.comparar(referencia, extraccion)
        self.assertTrue(any(h["codigo"] == "asociacion_incorrecta" for h in resultado["hallazgos"]))
        extraccion = _extraccion_perfecta(self.referencias["C2"])
        extraccion["actividades"][3]["unidad"] = "m"                               # 6,8 km leído como 6,8 m
        resultado = comparador.comparar(self.referencias["C2"], extraccion)
        self.assertTrue(any(h["codigo"] == "asociacion_incorrecta" and h["ref"] == "C2-A4" for h in resultado["hallazgos"]))

    # Cubre: T-71
    def test_24_la_lectura_ambigua_declarada_no_es_critica_pero_se_anota(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["actividades"][1]["cantidad"] = "2350"
        extraccion["actividades"][1]["unidad"] = "km"
        resultado = comparador.comparar(referencia, extraccion)
        self.assertEqual(resultado["errores_criticos"], 0)
        self.assertTrue(any(h["codigo"] == "lectura_ambigua_aceptada" for h in resultado["hallazgos"]))
        extraccion["actividades"][1]["cantidad"] = "2,35"
        self.assertEqual(comparador.comparar(referencia, extraccion)["errores_no_criticos"], 0)

    # Cubre: T-71
    def test_25_omitir_un_faltante_bloqueante_o_una_contradiccion_critica_es_critico(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["preguntas_cliente"] = [p for p in extraccion["preguntas_cliente"] if "polígono" not in p]
        extraccion["contradicciones"] = []
        resultado = comparador.comparar(referencia, extraccion)
        codigos = {(h["codigo"], h["gravedad"]) for h in resultado["hallazgos"]}
        self.assertIn(("faltante_sin_pregunta", "critico"), codigos)
        self.assertIn(("contradiccion_omitida", "critico"), codigos)

    # Cubre: T-71
    def test_26_una_pregunta_de_mas_o_mal_destinada_no_es_critica(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["preguntas_cliente"].append("¿Cuál es el plazo de entrega?")        # el documento ya lo trae
        extraccion["preguntas_cliente"].append("¿Sobre qué hora llegan?")
        extraccion["preguntas_cliente"].append(referencia["revisiones_internas"][0]["texto"])
        extraccion["revisiones_internas"] = extraccion["revisiones_internas"][1:]
        resultado = comparador.comparar(referencia, extraccion)
        codigos = [h["codigo"] for h in resultado["hallazgos"] if h["gravedad"] == "no_critico"]
        self.assertIn("pregunta_de_mas", codigos)
        self.assertIn("revision_interna_enviada_al_cliente", codigos)
        self.assertEqual(resultado["errores_criticos"], 0)

    # Cubre: T-71
    def test_27_localizador_ausente_es_no_critico_y_sin_extraccion_es_no_evaluable(self) -> None:
        referencia = self.referencias["C3"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["actividades"][0]["origen"] = {"hoja": "Detalle", "celda": "D8"}
        del extraccion["actividades"][1]["origen"]
        resultado = comparador.comparar(referencia, extraccion)
        self.assertEqual(sum(1 for h in resultado["hallazgos"] if h["codigo"] == "localizador_ausente_o_incorrecto"), 2)
        self.assertEqual(resultado["errores_criticos"], 0)
        vacio = comparador.comparar(referencia, None)
        self.assertEqual((vacio["errores_criticos"], vacio["no_evaluables"], vacio["sin_errores_criticos"]), (0, 1, False))

    # Cubre: T-71
    def test_28_una_extraccion_que_cita_ids_de_la_referencia_se_marca_como_sospechosa(self) -> None:
        referencia = self.referencias["C1"]
        extraccion = _extraccion_perfecta(referencia)
        extraccion["restricciones"].append("según C1-R1")
        resultado = comparador.comparar(referencia, extraccion)
        self.assertTrue(any(h["codigo"] == "posible_acceso_a_la_referencia" for h in resultado["hallazgos"]))
        with self.assertRaises(ValueError):
            comparador.comparar(referencia, {**_extraccion_perfecta(referencia), "caso": "C2"})

    # Cubre: T-71
    def test_29_el_informe_usa_denominadores_reales(self) -> None:
        resultados = [comparador.comparar(r, _extraccion_perfecta(r)) for r in self.referencias.values()]
        total = comparador.informe(resultados)
        esperadas = sum(len(r["actividades"]) for r in self.referencias.values())
        self.assertEqual((total["exhaustividad"], total["precision"]), (f"{esperadas}/{esperadas}",) * 2)
        self.assertEqual(comparador.informe([comparador.comparar(self.referencias["C1"], None)])["precision"], "sin denominador")

    # Cubre: T-71
    def test_30_evaluar_carpetas_lee_extracciones_por_caso(self) -> None:
        with tempfile.TemporaryDirectory() as carpeta:
            destino = Path(carpeta)
            (destino / "C1.extraccion.json").write_text(json.dumps(_extraccion_perfecta(self.referencias["C1"])), encoding="utf-8")
            informe = comparador.evaluar_carpetas(construir.REFERENCIAS, destino)
        por_caso = {r["caso"]: r for r in informe["casos"]}
        self.assertEqual(por_caso["C1"]["errores_criticos"], 0)
        self.assertEqual(por_caso["C2"]["no_evaluables"], 1)
        self.assertEqual(por_caso["C3"]["actividades_obtenidas"], 0)


class EntregaTest(unittest.TestCase):
    # Cubre: T-71
    def test_40_la_entrega_copia_solo_documentos_y_nada_de_las_referencias(self) -> None:
        with tempfile.TemporaryDirectory() as carpeta:
            destino = Path(carpeta) / "entrega"
            copiados = entrega.preparar_entrega(destino)
            nombres = sorted(p.name for p in copiados)
            self.assertEqual(nombres, sorted([*datos.ARCHIVOS.values(), "manifiesto_documentos.json"]))
            self.assertEqual(sorted(p.name for p in destino.iterdir()), nombres)
            for ruta in destino.iterdir():
                self.assertNotIn(b"referencia", ruta.read_bytes().lower() if ruta.suffix == ".json" else b"")

    # Cubre: T-71
    def test_41_el_destino_dentro_del_banco_o_con_archivos_se_rechaza(self) -> None:
        for prohibido in (construir.REFERENCIAS / "entrega", construir.RAIZ / "otra", construir.RAIZ, construir.RAIZ.parent):
            with self.assertRaises(entrega.EntregaInsegura, msg=str(prohibido)):
                entrega.preparar_entrega(prohibido)
        self.assertFalse((construir.RAIZ / "otra").exists())
        with tempfile.TemporaryDirectory() as carpeta:
            (Path(carpeta) / "algo.txt").write_text("x", encoding="utf-8")
            with self.assertRaises(entrega.EntregaInsegura):
                entrega.preparar_entrega(Path(carpeta))

    # Cubre: T-71
    def test_42_una_referencia_colada_en_la_entrega_se_detecta(self) -> None:
        with tempfile.TemporaryDirectory() as carpeta:
            destino = Path(carpeta) / "entrega"
            entrega.preparar_entrega(destino)
            colada = destino / "C1.referencia.json"
            colada.write_bytes((construir.REFERENCIAS / "C1.referencia.json").read_bytes())
            with self.assertRaises(entrega.EntregaInsegura):
                entrega.verificar_entrega(destino)
            colada.unlink()
            manifiesto = destino / "manifiesto_documentos.json"
            datos_m = json.loads(manifiesto.read_text(encoding="utf-8"))
            datos_m["nota"] = "ver C1-A1"
            manifiesto.write_text(json.dumps(datos_m), encoding="utf-8")
            with self.assertRaises(entrega.EntregaInsegura):
                entrega.verificar_entrega(destino)

    # Cubre: T-71
    def test_43_el_comparador_no_lee_referencias_por_su_cuenta(self) -> None:
        fuente = Path(comparador.__file__).read_text(encoding="utf-8")
        for modulo in ("anka", "datos", "generar", "construir"):
            self.assertNotRegex(fuente, rf"^\s*(?:from|import)\s+(?:laboratorio\.banco_documental\.)?{modulo}\b", modulo)


class AislamientoTest(unittest.TestCase):
    # Cubre: T-71
    def test_50_el_codigo_de_produccion_no_importa_el_laboratorio(self) -> None:
        raiz = Path(__file__).resolve().parent.parent / "anka"
        for ruta in raiz.rglob("*.py"):
            self.assertNotRegex(ruta.read_text(encoding="utf-8"), r"(?m)^\s*(?:from|import)\s+laboratorio", str(ruta))

    # Cubre: T-71
    def test_51_los_datos_no_se_modifican_al_construir_referencias(self) -> None:
        antes = copy.deepcopy(datos.CASOS)
        construir.construir_en_memoria()
        self.assertEqual(antes, datos.CASOS)


if __name__ == "__main__":
    unittest.main()
