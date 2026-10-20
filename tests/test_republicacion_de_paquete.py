"""Un trabajo ya publicado se republica con el paquete vigente en su misma hoja, sin perder su historia. Cubre: T-62

La hoja declara la versión efectiva del paquete; el evento de creación conserva la de origen (D-175).
Todo con SQLite temporal y un Drive simulado.
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import io
import json
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from anka.cli import work_jobs
from anka.dominio import trabajo_work as dominio
from anka.infra import store
from anka.servicios import trabajos_work
from anka.servicios.trabajos_work import TrabajoRechazado
from tests.test_trabajos_work import CARPETA, Base

ANTIGUA, NUEVA = "aa52f0eaa46d", "1797370a19dd"


class RepublicacionDeTrabajoTests(Base):
    def publicado_con_la_antigua(self, **fila):
        self.fila("T1", "m1", **fila)
        job = trabajos_work.crear_trabajo("T1", "m1", paquete_version=ANTIGUA)["job_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=ANTIGUA)
        return job, trabajos_work.obtener(job)["sheet_file_id"]

    def campos(self, sheet_id):
        libro = load_workbook(io.BytesIO(self.drive.archivos[sheet_id]["bytes"]))
        return {f[0].value: f[1].value for f in libro["TRABAJO"].iter_rows(min_row=2) if f[0].value}

    def test_la_hoja_republicada_declara_la_version_efectiva_y_muestra_la_de_creacion(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        self.assertEqual(self.campos(hoja)["version_del_paquete"], ANTIGUA)
        self.assertNotIn("version_del_paquete_al_crear", self.campos(hoja))
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        self.assertEqual((resultado["resultado"], resultado["sheet_file_id"]), ("actualizado", hoja))
        campos = self.campos(hoja)
        self.assertEqual(campos["version_del_paquete"], NUEVA)
        self.assertEqual(campos["version_del_paquete_al_crear"], ANTIGUA)

    def test_misma_identidad_misma_hoja_sin_duplicados(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        antes = trabajos_work.obtener(job)
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        despues = trabajos_work.obtener(job)
        for campo in ("job_id", "thread_id", "message_id", "contract_version", "sheet_file_id", "ref_sha256",
                      "created_at", "published_at"):
            self.assertEqual(despues[campo], antes[campo], campo)
        self.assertEqual((despues["state"], self.drive.creados), ("publicado", 1))
        self.assertEqual([a["id"] for a in self.drive.vivos(job)], [hoja])

    def test_el_historial_original_queda_intacto_y_la_republicacion_se_anota(self):
        # Cubre: T-62
        job, _hoja = self.publicado_con_la_antigua()
        with store.lectura() as db:
            antes = [dict(r) for r in db.execute("SELECT * FROM recepcion_eventos ORDER BY id")]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        with store.lectura() as db:
            despues = [dict(r) for r in db.execute("SELECT * FROM recepcion_eventos ORDER BY id")]
        self.assertEqual(despues[:len(antes)], antes)                      # nada anterior se reescribe
        creacion, = [e for e in despues if e["event_kind"] == "trabajo_creado"]
        self.assertEqual(json.loads(creacion["detail_json"])["package_version"], ANTIGUA)
        ultimo = json.loads(despues[-1]["detail_json"])
        self.assertEqual((despues[-1]["event_kind"], ultimo["package_version"], ultimo["como"], ultimo["job_id"]),
                         ("trabajo_republicado", NUEVA, "actualizado", job))
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        self.assertEqual(len(self.eventos("trabajo_republicado", job)), 1)       # misma versión: sin duplicar

    def test_republicar_dos_veces_es_idempotente(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        campos = self.campos(hoja)
        otra = trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        self.assertEqual((otra["resultado"], otra["sheet_file_id"]), ("actualizado", hoja))
        self.assertEqual((self.drive.creados, len(self.drive.vivos(job))), (1, 1))
        self.assertEqual(self.campos(hoja), campos)
        self.assertEqual(campos["version_del_paquete"], NUEVA)

    def test_una_marca_divergente_una_hoja_en_la_papelera_o_borrada_no_se_actualizan_ni_se_reemplazan(self):
        # Cubre: T-62
        for alterar in (lambda a: a["props"].update({dominio.CLAVE_DE_MARCA: "ANKAJOB-ajeno"}),
                        lambda a: a.update({"trashed": True}),
                        lambda a: a["props"].clear()):
            with self.subTest():
                self.setUp()
                job, hoja = self.publicado_con_la_antigua()
                alterar(self.drive.archivos[hoja])
                bytes_antes = self.drive.archivos[hoja]["bytes"]
                resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA,
                                                           solo_su_hoja=True)
                self.assertEqual(resultado["resultado"], "identidad_divergente")
                self.assertEqual((self.drive.creados, self.drive.archivos[hoja]["bytes"]), (1, bytes_antes))
                trabajo = trabajos_work.obtener(job)
                self.assertEqual((trabajo["state"], trabajo["sheet_file_id"]), ("publicado", hoja))
                self.assertEqual(trabajo["claim_token"], "")               # el reclamo se suelta

    def test_una_hoja_desaparecida_no_se_recrea_al_republicar(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        del self.drive.archivos[hoja]
        self.assertEqual(trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA,
                                                        solo_su_hoja=True)["resultado"], "identidad_divergente")
        self.assertEqual(self.drive.creados, 1)

    def test_un_trabajo_conciliado_no_se_republica(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        with store.transaccion() as db:
            db.execute("UPDATE trabajos_work SET state='conciliado' WHERE job_id=?", (job,))
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        self.assertEqual(self.campos(hoja)["version_del_paquete"], ANTIGUA)

    def test_un_trabajo_antiguo_sin_version_de_creacion_sigue_igual(self):
        # Cubre: T-62
        self.fila("T1", "m1")
        job = trabajos_work.crear_trabajo("T1", "m1")["job_id"]            # sin paquete_version, como antes de T-59
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        hoja = trabajos_work.obtener(job)["sheet_file_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA, paquete_version=NUEVA, solo_su_hoja=True)
        campos = self.campos(hoja)
        self.assertNotIn("version_del_paquete", campos)
        self.assertNotIn("version_del_paquete_al_crear", campos)

    def test_publicar_sin_version_efectiva_conserva_la_de_creacion(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(self.campos(hoja)["version_del_paquete"], ANTIGUA)
        self.assertNotIn("version_del_paquete_al_crear", self.campos(hoja))
        self.assertEqual(trabajos_work.libro_del_trabajo(job, NUEVA) != trabajos_work.libro_del_trabajo(job), True)

    def test_la_orden_publicar_republica_con_el_paquete_vigente_solo_en_su_hoja(self):
        # Cubre: T-62
        job, hoja = self.publicado_con_la_antigua()
        args = SimpleNamespace(job_id=job, carpeta_id=CARPETA, confirmar_drive=True)
        with patch("anka.infra.google_services.services", return_value=(MagicMock(), self.drive)), \
                patch.object(work_jobs.paquete_work, "version_del_paquete", return_value=NUEVA):
            resultado = work_jobs._publicar(args)
            self.assertEqual((resultado["resultado"], resultado["sheet_file_id"]), ("actualizado", hoja))
            self.drive.archivos[hoja]["props"] = {dominio.CLAVE_DE_MARCA: "ANKAJOB-ajeno"}
            self.assertEqual(work_jobs._publicar(args)["resultado"], "identidad_divergente")
        self.assertEqual(self.drive.creados, 1)
