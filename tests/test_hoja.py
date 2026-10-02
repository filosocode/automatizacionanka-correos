"""El libro de cálculo del caso (anka/dominio/hoja.py, requisito T-17): la tabla vigente y lo que no debe perderse."""

from tests import _isolate  # noqa: F401  (debe ir primero)

import json
import unittest
from io import BytesIO

from anka.dominio.hoja import HOJAS, MAX_TEXTO, build_case_workbook

ORIGINAL = [{"activity": "Levantamiento", "quantity": "10", "unit": "ha", "estado": "pendiente"}]
VIGENTE = [{"activity": "Levantamiento", "quantity": "8", "unit": "ha", "estado": "confirmado"},
           {"activity": "=HYPERLINK(\"http://malo\")", "quantity": "1", "estado": "pendiente"}]
VERSIONES = [
    {"version": 1, "source": "model", "created_at": "2026-10-01T09:00:00-05:00",
     "summary": "Tabla original del modelo", "reason": "", "activities_json": json.dumps(ORIGINAL)},
    {"version": 2, "source": "dashboard", "created_at": "2026-10-01T10:00:00-05:00",
     "summary": "+1 fila(s): x; totales: ha 10 → 8", "reason": "modelo", "activities_json": json.dumps(VIGENTE)},
]
CASO = {"analysis": {"project_name": "Línea Norte", "activities": VIGENTE}, "client_name": "ACME",
        "quote_code": "26004"}
INICIAL = {"sender": "Ana <ana@acme.com>", "subject": "Cotización línea", "body": "Necesito levantar 10 ha.",
           "occurred_at": "2026-10-01T09:00:00-05:00"}


def abrir(datos):
    from openpyxl import load_workbook
    return load_workbook(BytesIO(datos))


def texto(hoja):
    return [[c.value for c in fila] for fila in hoja.iter_rows()]


# Cubre: T-17
class CaseWorkbook(unittest.TestCase):
    def setUp(self):
        self.libro = abrir(build_case_workbook(CASO, VERSIONES, INICIAL, "2026-10-02 10:30"))

    def test_it_has_the_four_sheets_in_order(self):
        self.assertEqual(tuple(self.libro.sheetnames), HOJAS)

    def test_the_first_sheet_is_the_current_table_and_nothing_else(self):
        filas = texto(self.libro["Actividades"])
        self.assertEqual(filas[0][0], "Actividad")                       # encabezados de la tabla
        self.assertEqual(filas[1][0], "Levantamiento")
        self.assertEqual(filas[1][3], "8")                               # la cantidad vigente, no la original
        self.assertEqual(self.libro["Actividades"].freeze_panes, "A2")

    def test_the_initial_request_and_the_first_analysis_table_are_kept_apart(self):
        filas = texto(self.libro["Solicitud inicial"])
        plano = " ".join(str(c) for fila in filas for c in fila if c)
        self.assertIn("ana@acme.com", plano)
        self.assertIn("Necesito levantar 10 ha.", plano)                 # el correo tal como llegó
        self.assertIn("Tabla que salió del primer análisis", plano)
        original = [fila for fila in filas if fila[0] == "Levantamiento"]
        self.assertEqual(original[0][3], "10")                           # la cantidad original, aunque ya se editó

    def test_the_history_lists_every_version_without_the_internal_reasons(self):
        filas = texto(self.libro["Historial"])
        self.assertEqual([f[0] for f in filas[1:]], [1, 2])
        self.assertEqual(filas[1][2], "Análisis automático")
        self.assertEqual(filas[2][2], "Edición en el dashboard")
        self.assertIn("ha 10 → 8", filas[2][3])
        plano = " ".join(str(c) for fila in filas for c in fila if c)
        self.assertNotIn("El modelo se lo inventó", plano)
        self.assertNotIn("modelo", " ".join(str(f[1]) for f in filas[1:]))   # el motivo interno ("modelo") no se copia

    def test_the_readme_explains_where_it_comes_from_and_that_manual_edits_are_overwritten(self):
        plano = " ".join(str(f[0]) for f in texto(self.libro["Léeme"]) if f[0])
        for esperado in ("Línea Norte", "ACME", "26004", "2026-10-02 10:30", "se sobrescribe", "Historial de versiones"):
            self.assertIn(esperado, plano)

    def test_text_that_looks_like_a_formula_is_neutralized(self):
        celdas = [c.value for fila in self.libro["Actividades"].iter_rows() for c in fila]
        self.assertNotIn('=HYPERLINK("http://malo")', celdas)
        self.assertIn("'=HYPERLINK(\"http://malo\")", celdas)

    def test_a_very_long_first_email_is_cut_and_a_case_without_history_still_builds(self):
        largo = dict(INICIAL, body="x" * (MAX_TEXTO + 500))
        libro = abrir(build_case_workbook(CASO, VERSIONES, largo, "ahora"))
        celda = libro["Solicitud inicial"]["B5"].value
        self.assertTrue(celda.endswith("[…]"))
        self.assertLessEqual(len(celda), MAX_TEXTO + 5)
        vacio = abrir(build_case_workbook({"analysis": {}}, [], None, "ahora"))
        self.assertIn("Sin versiones guardadas", " ".join(str(c.value) for fila in vacio["Solicitud inicial"].iter_rows()
                                                          for c in fila if c.value))


if __name__ == "__main__":
    unittest.main()
