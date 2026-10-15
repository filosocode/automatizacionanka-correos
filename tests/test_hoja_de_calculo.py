"""La hoja de Google Sheets del caso (anka/servicios/hoja_de_calculo.py, requisito T-17)."""

from tests import _isolate  # noqa: F401  (debe ir primero)

import tempfile
import threading
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from anka.dominio.errores import AutorizacionDeGoogle
from anka.dominio.status import QuoteStatus
from anka.infra import settings, store
from anka.servicios import ediciones, hoja_de_calculo
from anka.servicios.hoja_de_calculo import ResultadoHoja, SincronizadorDeHojas, nombre_de_hoja, publicar, sincronizar_hoja
from anka.servicios.repositorio import RepositorioDeCasos
from tests.fake_drive import FakeDriveTree
from tests.test_app import complete_analysis


def libro(datos):
    from openpyxl import load_workbook
    return load_workbook(BytesIO(datos))


# Cubre: T-17
class PublishingToDrive(unittest.TestCase):
    def setUp(self):
        self.drive = FakeDriveTree({"ACME": {"Solicitud": {}}})
        self.carpeta = self.drive.folder_id("ACME/Solicitud")
        self.nombre = nombre_de_hoja("Línea Norte")

    def test_the_first_time_it_creates_the_sheet_in_the_project_folder(self):
        hecho = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        self.assertTrue(hecho.creada)
        self.assertEqual(self.drive.uploads[-1][:3], (self.carpeta, "Tabla de actividades - Línea Norte",
                                                       "application/vnd.google-apps.spreadsheet"))
        self.assertTrue(hecho.id.startswith("sheet-"))

    def test_afterwards_it_rewrites_the_same_file_instead_of_creating_another(self):
        primero = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        segundo = publicar(self.drive, self.carpeta, self.nombre, b"dos", primero.id, primero.modificada)
        self.assertFalse(segundo.creada)
        self.assertEqual(segundo.id, primero.id)                       # el mismo archivo: Google conserva sus versiones
        self.assertEqual(len(self.drive.uploads), 1)                   # no se subió otro
        self.assertEqual(self.drive.sheets[primero.id]["bytes"], b"dos")
        self.assertEqual(segundo.aviso, "")

    def test_a_sheet_edited_directly_in_drive_is_noticed_and_the_warning_points_to_the_version_history(self):
        primero = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        self.drive.edit_sheet_in_drive(primero.id)
        segundo = publicar(self.drive, self.carpeta, self.nombre, b"dos", primero.id, primero.modificada)
        self.assertIn("Historial de versiones", segundo.aviso)
        self.assertEqual(self.drive.sheets[primero.id]["bytes"], b"dos")   # vuelve a reflejar el dashboard

    def test_the_few_seconds_drive_takes_to_finish_converting_are_not_an_external_edit(self):
        primero = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        self.drive.sheets[primero.id]["modified"] = "2026-10-02T10:00:01.000Z"
        guardada = self.drive.sheets[primero.id]["modified"]
        self.drive.settle_conversion(primero.id)                       # +58 s en la fecha: aún dentro de la tolerancia
        segundo = publicar(self.drive, self.carpeta, self.nombre, b"dos", primero.id, guardada)
        self.assertEqual(segundo.aviso, "")                            # sin falsa alarma

    def test_a_sheet_that_was_deleted_is_recreated_and_the_person_is_told(self):
        primero = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        self.drive.trash_sheet(primero.id)
        segundo = publicar(self.drive, self.carpeta, self.nombre, b"dos", primero.id, primero.modificada)
        self.assertTrue(segundo.creada)
        self.assertNotEqual(segundo.id, primero.id)
        self.assertIn("se creó una nueva", segundo.aviso)

    def test_a_sheet_with_the_same_name_already_in_the_folder_is_adopted_not_duplicated(self):
        previa = publicar(self.drive, self.carpeta, self.nombre, b"uno", "", "")
        adoptada = publicar(self.drive, self.carpeta, self.nombre, b"dos", "", "")     # sin id guardado
        self.assertEqual(adoptada.id, previa.id)
        self.assertFalse(adoptada.creada)
        self.assertEqual(len(self.drive.uploads), 1)

    def test_a_drive_error_other_than_not_found_is_not_hidden(self):
        class Roto:
            def files(self):
                raise RuntimeError("503")
        with self.assertRaises(RuntimeError):
            publicar(Roto(), self.carpeta, self.nombre, b"x", "sheet-9", "")


class CaseFixture(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.folder = Path(self._tmp.name)
        patcher = patch.object(settings, "DATA_DIR", self.folder)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.drive = FakeDriveTree({"ACME": {"Solicitud": {}}})
        self.carpeta = self.drive.folder_id("ACME/Solicitud")
        self.repo = RepositorioDeCasos()
        db = store.open_db(self.folder)
        analysis = complete_analysis()
        analysis["activities"] = [{"activity": "Levantamiento", "quantity": "10", "unit": "ha", "estado": "pendiente"}]
        store.save(db, "t-1", "m1", QuoteStatus.READY, "ana@acme.com", analysis, folder_id=self.carpeta)
        store.add_message(db, "t-1", "m1", "received", "Ana <ana@acme.com>", "Cotización línea",
                          "Necesito levantar 10 ha en la línea norte.", "2026-10-01T09:00:00-05:00")
        db.commit()
        db.close()

    def hoja(self):
        return libro(self.drive.sheets["sheet-1"]["bytes"])


# Cubre: T-17
class SyncOneCase(CaseFixture):
    def test_it_creates_the_sheet_remembers_it_and_keeps_the_original_request(self):
        resultado = sincronizar_hoja("t-1", drive=self.drive)
        self.assertEqual(resultado.estado, "creada")
        self.assertTrue(resultado.url.startswith("https://docs.google.com/spreadsheets/d/sheet-1"))
        caso = self.repo.obtener("t-1")
        self.assertEqual((caso.sheet_id, caso.sheet_error), ("sheet-1", ""))
        self.assertTrue(caso.sheet_synced_at)
        self.assertEqual(caso.sheet_url, resultado.url)
        inicial = " ".join(str(c.value) for fila in self.hoja()["Solicitud inicial"].iter_rows() for c in fila if c.value)
        self.assertIn("Necesito levantar 10 ha en la línea norte.", inicial)

    def test_an_edit_updates_the_same_sheet_and_nothing_of_the_beginning_is_lost(self):
        sincronizar_hoja("t-1", drive=self.drive)
        ediciones.save_activities("t-1", [{"activity": "Levantamiento", "quantity": "8", "unit": "ha",
                                          "estado": "confirmado"}], "correccion")
        resultado = sincronizar_hoja("t-1", drive=self.drive)
        self.assertEqual(resultado.estado, "actualizada")
        self.assertEqual(len(self.drive.sheets), 1)                     # el mismo archivo
        libro_actual = self.hoja()
        vigente = [[c.value for c in fila] for fila in libro_actual["Actividades"].iter_rows()]
        self.assertEqual(vigente[1][3], "8")                            # la hoja principal tiene la tabla de ahora
        original = [[c.value for c in fila] for fila in libro_actual["Solicitud inicial"].iter_rows()
                    if fila[0].value == "Levantamiento"]
        self.assertEqual(original[0][3], "10")                          # la solicitud inicial conserva lo original
        historial = [[c.value for c in fila] for fila in libro_actual["Historial"].iter_rows()][1:]
        self.assertGreaterEqual(len(historial), 2)                      # y el historial lista cada versión
        self.assertEqual(historial[0][2], "Análisis automático")
        self.assertEqual(len(self.repo.versiones_de_tabla("t-1")), len(historial))   # nada se borró de la base

    def test_without_a_project_folder_nothing_is_created_and_it_says_why(self):
        db = store.open_db(self.folder)
        store.save(db, "t-2", "m2", QuoteStatus.REVIEW_PENDING, "x@y.co", complete_analysis())
        db.close()
        resultado = sincronizar_hoja("t-2", drive=self.drive)
        self.assertEqual(resultado.estado, "sin_carpeta")
        self.assertIn("Excel", resultado.mensaje)
        self.assertEqual(self.drive.sheets, {})
        self.assertEqual(sincronizar_hoja("no-existe", drive=self.drive).estado, "sin_caso")

    def test_an_external_edit_is_reported_in_the_message(self):
        sincronizar_hoja("t-1", drive=self.drive)
        self.drive.edit_sheet_in_drive("sheet-1")
        resultado = sincronizar_hoja("t-1", drive=self.drive)
        self.assertIn("cambios hechos directamente en Drive", resultado.mensaje)


# Cubre: T-17
class BackgroundSync(CaseFixture):
    def test_a_save_schedules_an_update_only_when_the_case_has_a_folder(self):
        llamadas = []
        sincronizador = SincronizadorDeHojas(sincronizar=lambda hilo: llamadas.append(hilo), espera=0, en_hilo=False)
        sincronizador.programar("t-1")
        db = store.open_db(self.folder)
        store.save(db, "t-2", "m2", QuoteStatus.REVIEW_PENDING, "x@y.co", complete_analysis())
        db.close()
        sincronizador.programar("t-2")                                  # sin carpeta: no hay nada que actualizar
        sincronizador.programar("no-existe")
        self.assertEqual(llamadas, ["t-1"])

    def test_it_can_be_switched_off(self):
        llamadas = []
        sincronizador = SincronizadorDeHojas(sincronizar=lambda hilo: llamadas.append(hilo), espera=0, en_hilo=False)
        with patch.dict("os.environ", {"ANKA_HOJA_AUTOMATICA": "0"}):
            sincronizador.programar("t-1")
        self.assertEqual(llamadas, [])

    def test_changes_that_arrive_while_updating_cause_one_more_update_not_one_per_change(self):
        entrada, soltar, llamadas = threading.Event(), threading.Event(), []

        def lenta(hilo):
            llamadas.append(hilo)
            entrada.set()
            soltar.wait(5)

        sincronizador = SincronizadorDeHojas(sincronizar=lenta, espera=0)
        sincronizador.programar("t-1")
        self.assertTrue(entrada.wait(5))
        for _ in range(5):                                              # cinco ediciones más mientras se actualiza
            sincronizador.programar("t-1")
        soltar.set()
        for _ in range(100):
            if not sincronizador._activos:
                break
            threading.Event().wait(0.05)
        self.assertEqual(len(llamadas), 2)                              # la que corría y una sola repetición

    def test_a_failure_is_recorded_in_the_case_and_never_raised(self):
        def falla(hilo):
            raise RuntimeError("503 con datos tecnicos")

        sincronizador = SincronizadorDeHojas(sincronizar=falla, espera=0, en_hilo=False)
        with self.assertLogs("anka.servicios.hoja_de_calculo", "WARNING"):
            sincronizador.programar("t-1")
        error = self.repo.obtener("t-1").sheet_error
        self.assertIn("No se pudo actualizar la hoja de Google", error)
        self.assertNotIn("datos tecnicos", error)                       # el texto técnico no llega a la pantalla

    def test_an_expired_google_authorization_says_how_to_fix_it(self):
        def vencida(hilo):
            raise AutorizacionDeGoogle("Google ya no acepta el token guardado.")

        sincronizador = SincronizadorDeHojas(sincronizar=vencida, espera=0, en_hilo=False)
        with self.assertLogs("anka.servicios.hoja_de_calculo", "WARNING"):
            sincronizador.programar("t-1")
        self.assertIn("Volver a conectar Google", self.repo.obtener("t-1").sheet_error)

    def test_a_good_update_clears_the_previous_error(self):
        self.repo.registrar_error_de_hoja("t-1", "falló antes")
        sincronizar_hoja("t-1", drive=self.drive)
        self.assertEqual(self.repo.obtener("t-1").sheet_error, "")


class Result(unittest.TestCase):
    def test_the_result_carries_state_message_and_url(self):
        self.assertEqual(ResultadoHoja("creada", "m", "u").url, "u")
        self.assertIsInstance(hoja_de_calculo.sincronizador, SincronizadorDeHojas)
        self.assertTrue(MagicMock)


if __name__ == "__main__":
    unittest.main()
