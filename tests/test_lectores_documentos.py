"""Lectores de documentos con archivos reales mínimos: Word, Excel, PowerPoint, OpenDocument, PDF, imagen y ZIP.

Cada prueba construye el documento en memoria (tests/_documentos.py) y comprueba el texto y los localizadores que
recibe el modelo. Los motores de OCR se sustituyen: lo que se prueba es cómo el lector los usa y qué avisa.
Cubre: F1-03, R12, T-24
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import io
import sys
import unittest
from unittest.mock import MagicMock, patch

from anka.dominio.lectores_texto import Unreadable, hay_tabla_en_imagen, lineas_nuevas_del_ocr, reparar_ligaduras_perdidas
from anka.infra import attachments, lectores_imagen, lectores_office, lectores_pdf, lectores_zip
from tests import _documentos as doc

OCR = "anka.infra.lectores_zip.ocr_image"


def con_ocr(texto="texto de la imagen"):
    return patch(OCR, return_value=texto)


class WordTests(unittest.TestCase):
    def test_paragraphs_and_tables_keep_their_locators(self):
        cuerpo = doc.parrafo("Alcance del proyecto") + doc.tabla_docx([["Item", "Área"], ["Parque", "544 ha"]])
        texto = lectores_office.docx_text(doc.docx(cuerpo))
        self.assertIn("Párrafo 1: Alcance del proyecto", texto)
        self.assertIn("Tabla 1, fila 1: columna 1=Item | columna 2=Área", texto)
        self.assertIn("Tabla 1, fila 2: columna 1=Parque | columna 2=544 ha", texto)

    def test_headers_footers_and_notes_are_read_once(self):
        extra = {
            "word/header1.xml": f'<w:hdr xmlns:w="{doc.W}"><w:p><w:r><w:t>Encabezado ET</w:t></w:r></w:p></w:hdr>',
            "word/footer1.xml": f'<w:ftr xmlns:w="{doc.W}"><w:p><w:r><w:t>Pie ET</w:t></w:r></w:p></w:ftr>',
            "word/comments.xml": f'<w:comments xmlns:w="{doc.W}"><w:comment>{doc.parrafo("Revisar área")}</w:comment>'
                                 "</w:comments>",
            "word/footnotes.xml": f'<w:footnotes xmlns:w="{doc.W}"><w:footnote>{doc.parrafo("Nota al pie")}</w:footnote>'
                                  "</w:footnotes>",
        }
        texto = lectores_office.docx_text(doc.docx(doc.parrafo("Cuerpo"), extra))
        self.assertEqual(texto.count("Encabezado ET"), 1)  # regresión: salía duplicado
        self.assertEqual(texto.count("Pie ET"), 1)
        self.assertIn("Contenido asociado (word/comments.xml): Revisar área", texto)
        self.assertIn("Contenido asociado (word/footnotes.xml): Nota al pie", texto)

    def test_tracked_changes_are_listed_and_flagged(self):
        cuerpo = (doc.parrafo("Base") + "<w:p><w:ins><w:r><w:t>texto nuevo</w:t></w:r></w:ins></w:p>"
                  "<w:p><w:del><w:r><w:delText>texto viejo</w:delText></w:r></w:del></w:p>")
        texto = lectores_office.docx_text(doc.docx(cuerpo))
        self.assertIn("CAMBIOS RASTREADOS", texto)
        self.assertIn("Revisión rastreada (ins): texto nuevo", texto)
        self.assertIn("Revisión rastreada (del): texto viejo", texto)

    def test_embedded_images_are_read_with_ocr_and_marked_as_visual_coverage(self):
        with con_ocr("Área 12 ha en la imagen"):
            texto = lectores_office.docx_text(doc.docx(doc.parrafo("x"), {"word/media/img1.png": doc.PNG_1X1}))
        self.assertIn("[Imagen incrustada word/media/img1.png; texto OCR]\nÁrea 12 ha en la imagen", texto)
        self.assertIn("COBERTURA VISUAL PARCIAL", texto)

    def test_an_image_the_ocr_cannot_read_is_reported_not_skipped(self):
        with patch(OCR, side_effect=Unreadable("OCR no detectó texto")):
            texto = lectores_office.docx_text(doc.docx(doc.parrafo("x"), {"word/media/img1.png": doc.PNG_1X1}))
        self.assertIn("CONTENIDO GRÁFICO NO INSPECCIONADO (word/media/img1.png): OCR no detectó texto", texto)

    def test_more_images_than_the_limit_are_reported_as_partial(self):
        imagenes = {f"word/media/i{n}.png": doc.PNG_1X1 for n in range(5)}
        with con_ocr(), patch.object(lectores_zip, "MAX_OCR_IMAGES", 2):
            texto = lectores_office.docx_text(doc.docx(doc.parrafo("x"), imagenes))
        self.assertEqual(texto.count("[Imagen incrustada"), 2)
        self.assertIn("LECTURA PARCIAL: documento con más de 2 imágenes incrustadas", texto)


class PowerPointTests(unittest.TestCase):
    def test_slides_are_read_in_numeric_order(self):
        diapositivas = [f"Texto {n}" for n in range(1, 12)]
        texto = lectores_office.pptx_text(doc.pptx(diapositivas))
        self.assertLess(texto.index("[Diapositiva 2]"), texto.index("[Diapositiva 10]"))
        self.assertIn("[Diapositiva 11]\nTexto 11", texto)

    def test_notes_smartart_charts_and_images_are_covered_or_flagged(self):
        extra = {
            "ppt/notesSlides/notesSlide1.xml": f'<p:notes xmlns:p="{doc.P}" xmlns:a="{doc.A}"><a:t>Nota</a:t></p:notes>',
            "ppt/diagrams/data1.xml": f'<d xmlns:a="{doc.A}"><a:t>Nodo SmartArt</a:t></d>',
            "ppt/charts/chart1.xml": "<c/>",
            "ppt/media/i.png": doc.PNG_1X1,
        }
        with con_ocr("texto en la figura"):
            texto = lectores_office.pptx_text(doc.pptx(["Portada"], extra))
        self.assertIn("Notas del presentador (ppt/notesSlides/notesSlide1.xml): Nota", texto)
        self.assertIn("texto SmartArt (ppt/diagrams/data1.xml): Nodo SmartArt", texto)
        self.assertIn("CONTENIDO DE GRÁFICOS NO EXTRAÍDO: ppt/charts/chart1.xml", texto)
        self.assertIn("texto en la figura", texto)

    def test_a_file_without_slides_is_unreadable(self):
        with self.assertRaises(Unreadable):
            lectores_office.pptx_text(doc.zip_de({"otra/cosa.xml": "<a/>"}))


class OpenDocumentTests(unittest.TestCase):
    def test_text_document_reads_headings_paragraphs_and_images(self):
        contenido = '<office:text><text:h>Título</text:h><text:p>Primer párrafo</text:p></office:text>'
        with con_ocr("ocr odt"):
            texto = lectores_office.odt_text(doc.odf(contenido, {"Pictures/a.png": doc.PNG_1X1}))
        for esperado in ("Título", "Primer párrafo", "ocr odt", "COBERTURA VISUAL PARCIAL"):
            self.assertIn(esperado, texto)

    def test_spreadsheet_keeps_cell_locators_formulas_and_repeated_columns(self):
        contenido = (
            '<office:spreadsheet><table:table table:name="Datos"><table:table-row>'
            "<table:table-cell><text:p>Área</text:p></table:table-cell>"
            '<table:table-cell table:formula="of:=SUM([.A1])"><text:p></text:p></table:table-cell>'
            '<table:table-cell table:number-columns-repeated="2"><text:p>x</text:p></table:table-cell>'
            "</table:table-row></table:table></office:spreadsheet>"
        )
        texto = lectores_office.ods_text(doc.odf(contenido))
        self.assertIn("Hoja: Datos", texto)
        self.assertIn("fila 1, columna 1=Área", texto)
        self.assertIn("FÓRMULA SIN VALOR CALCULADO: hoja Datos, fila 1, columna 2", texto)
        self.assertIn("fila 1, columna 3=x | fila 1, columna 4=x", texto)

    def test_presentation_numbers_its_pages_and_reads_images(self):
        contenido = ("<office:presentation><draw:page><text:p>Página A</text:p></draw:page>"
                     "<draw:page><text:p>Página B</text:p></draw:page></office:presentation>")
        with con_ocr("ocr odp"):
            texto = lectores_office.odp_text(doc.odf(contenido, {"Pictures/a.png": doc.PNG_1X1}))
        self.assertIn("[Diapositiva 1]\nPágina A", texto)
        self.assertIn("[Diapositiva 2]\nPágina B", texto)
        self.assertIn("ocr odp", texto)


class ExcelTests(unittest.TestCase):
    def libro(self, hojas=("Datos",), comentario=True, **extra):
        from openpyxl import Workbook
        from openpyxl.comments import Comment

        libro = Workbook()
        primera = libro.active
        primera.title = hojas[0]
        primera["A1"], primera["B1"], primera["B2"] = "Área", 12, "=B1*2"
        if comentario:
            primera["A3"] = "nota"
            primera["A3"].comment = Comment("comentario de celda", "autor")
        for nombre in hojas[1:]:
            hoja = libro.create_sheet(nombre)
            hoja["A1"] = "secreto"
            hoja.sheet_state = "hidden"
        flujo = io.BytesIO()
        libro.save(flujo)
        contenido = flujo.getvalue()
        if extra:  # partes adicionales (imágenes, gráficos) añadidas al ZIP
            import zipfile
            entrada = zipfile.ZipFile(io.BytesIO(contenido))
            partes = {n: entrada.read(n) for n in entrada.namelist()}
            contenido = doc.zip_de({**partes, **extra})
        return contenido

    def test_values_formulas_and_hidden_sheets_keep_their_locators(self):
        texto = lectores_office.xlsx_text(self.libro(("Datos", "Oculta")))
        self.assertIn("Hoja: Datos; visibilidad: visible", texto)
        self.assertIn("A1=Área | B1=12", texto)
        self.assertIn("FÓRMULA SIN VALOR CALCULADO: hoja Datos, celda B2", texto)
        self.assertIn("Hoja: Oculta; visibilidad: hidden", texto)
        self.assertIn("A1=secreto", texto)

    def test_comments_are_listed_once_even_with_several_sheets(self):
        # regresión: el bloque de comentarios, imágenes y gráficos se repetía por cada hoja
        texto = lectores_office.xlsx_text(self.libro(("Datos", "Oculta", "Otra")))
        self.assertEqual(texto.count("comentario de celda"), 1)

    def test_images_go_through_the_ocr_once_per_workbook(self):
        with patch(OCR, return_value="texto en la imagen") as ocr:
            texto = lectores_office.xlsx_text(self.libro(("A", "B", "C"), **{"xl/media/img1.png": doc.PNG_1X1}))
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual(texto.count("texto en la imagen"), 1)

    def test_charts_and_drawings_are_reported(self):
        grafico = '<c:chartSpace xmlns:c="urn:c"><c:f>Datos!$B$1</c:f><c:v>12</c:v></c:chartSpace>'
        texto = lectores_office.xlsx_text(self.libro(**{
            "xl/charts/chart1.xml": grafico, "xl/drawings/drawing1.xml": "<d/>", "xl/pivotTables/p1.xml": "<p/>"}))
        self.assertIn("Gráfico (xl/charts/chart1.xml); referencias y valores almacenados: Datos!$B$1; 12", texto)
        self.assertIn("DIBUJOS/FORMAS DE HOJA PENDIENTES DE LECTURA VISUAL: xl/drawings/drawing1.xml", texto)
        self.assertIn("COBERTURA PARCIAL: tablas dinámicas", texto)

    def test_cell_limit_reports_the_partial_reading_once_and_stops(self):
        from openpyxl import Workbook

        libro = Workbook()
        for fila in range(1, 41):
            libro.active.append([f"A{fila}", fila, "x"])
        flujo = io.BytesIO()
        libro.save(flujo)
        with patch.object(lectores_office, "MAX_SPREADSHEET_CELLS", 10):
            texto = lectores_office.xlsx_text(flujo.getvalue())
        self.assertEqual(texto.count("LECTURA PARCIAL"), 1)  # regresión: se repetía por cada fila restante

    def test_missing_openpyxl_is_a_clear_limitation(self):
        contenido = self.libro()
        with patch.dict(sys.modules, {"openpyxl": None}), self.assertRaises(Unreadable) as caught:
            lectores_office.xlsx_text(contenido)
        self.assertIn("openpyxl", str(caught.exception))


class ExcelAntiguoTests(unittest.TestCase):
    def xlrd_falso(self, filas):
        hoja = MagicMock()
        hoja.name, hoja.nrows, hoja.ncols = "Hoja1", len(filas), len(filas[0])
        hoja.cell_value.side_effect = lambda r, c: filas[r][c]
        libro = MagicMock()
        libro.sheets.return_value = [hoja]
        modulo = MagicMock()
        modulo.open_workbook.return_value = libro
        return modulo, libro

    def test_cells_keep_row_and_column_and_empty_ones_are_skipped(self):
        modulo, libro = self.xlrd_falso([["Área", 12], ["", "ha"]])
        with patch.dict(sys.modules, {"xlrd": modulo}):
            texto = lectores_office.xls_text(b"datos")
        self.assertIn("Hoja: Hoja1; filas: 2; columnas: 2", texto)
        self.assertIn("fila 1, columna 1=Área | fila 1, columna 2=12", texto)
        self.assertIn("fila 2, columna 2=ha", texto)
        libro.release_resources.assert_called_once()

    def test_cell_limit_is_reported_as_partial(self):
        modulo, _ = self.xlrd_falso([["a", "b"]] * 20)
        with patch.dict(sys.modules, {"xlrd": modulo}), patch.object(lectores_office, "MAX_SPREADSHEET_CELLS", 5):
            texto = lectores_office.xls_text(b"datos")
        self.assertIn("LECTURA PARCIAL: se superó el límite de 5 celdas", texto)

    def test_a_corrupt_book_and_a_missing_library_are_clear_limitations(self):
        modulo = MagicMock()
        modulo.open_workbook.side_effect = ValueError("no es un libro")
        with patch.dict(sys.modules, {"xlrd": modulo}), self.assertRaises(Unreadable) as caught:
            lectores_office.xls_text(b"x")
        self.assertIn("no se pudo abrir el libro .xls", str(caught.exception))
        with patch.dict(sys.modules, {"xlrd": None}), self.assertRaises(Unreadable) as caught:
            lectores_office.xls_text(b"x")
        self.assertIn("xlrd", str(caught.exception))


class PdfTests(unittest.TestCase):
    LARGO = ("Área 544 ha del parque solar en Villavicencio, con levantamiento topográfico y LiDAR de toda "
             "la zona del proyecto en el Meta, con entregables en formato digital.")

    def test_native_text_is_read_by_page_without_needing_ocr(self):
        with patch.object(lectores_pdf, "ocr_pil_image") as ocr:
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto([self.LARGO, self.LARGO.replace("544", "61")]))
        self.assertIn("[Página 1]", texto)
        self.assertIn("[Página 2]", texto)
        self.assertIn("61 ha", texto)
        ocr.assert_not_called()

    def test_a_page_with_little_text_gets_an_ocr_pass_that_adds_only_new_lines(self):
        with patch.object(lectores_pdf, "ocr_pil_image", return_value="Cuadro de áreas\nTotal 1 ha"):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto(["Cuadro de áreas"]))
        self.assertIn("[OCR visual adicional]", texto)
        self.assertIn("Total 1 ha", texto)
        self.assertEqual(texto.count("Cuadro de áreas"), 1)  # la línea que ya estaba en el texto nativo no se repite

    def test_a_page_the_ocr_cannot_read_is_reported_as_not_inspected(self):
        with patch.object(lectores_pdf, "ocr_pil_image", side_effect=Unreadable("OCR no detectó texto")):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto(["x"]))
        self.assertIn("CONTENIDO VISUAL NO INSPECCIONADO: OCR no detectó texto", texto)

    def test_without_pdfium_a_scanned_page_says_it_was_not_inspected(self):
        contenido = doc.pdf_con_texto([""])
        with patch.dict(sys.modules, {"pypdfium2": None}):
            texto = lectores_pdf.pdf_text(contenido)
        self.assertIn("CONTENIDO VISUAL NO INSPECCIONADO: falta pypdfium2", texto)

    def test_a_page_too_large_for_the_ocr_is_not_silently_skipped(self):
        with patch.object(lectores_pdf, "MAX_OCR_PIXELS", 100):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto(["x"]))
        self.assertIn("página PDF excede el límite de píxeles para OCR", texto)

    def test_missing_pypdf_is_a_clear_limitation(self):
        with patch.dict(sys.modules, {"pypdf": None}), self.assertRaises(Unreadable) as caught:
            lectores_pdf.pdf_text(doc.pdf_con_texto(["x"]))
        self.assertIn("pypdf", str(caught.exception))


class LecturaDePdfT124Tests(unittest.TestCase):
    """T-124: el OCR no descarta cifras cortas de una tabla en imagen y las sílabas «ti» perdidas se recuperan."""

    CELDAS = "Guavio\n120\n17\nQuimbo\n150\n1\n35\nBetania\n60"

    def test_ocr_lines_equal_to_the_native_text_except_for_accents_are_not_new(self):
        # Cubre: T-124 (el OCR repite el texto nativo sin tildes: «Declaracion» por «Declaración»)
        nativo = "Declaración de Renta y Complementario\nPágina 1 de 3"
        ocr = "Declaracion de Renta y Complementario\nPagina 1 de 3\nTotal a pagar 1.250.000"
        self.assertEqual(lineas_nuevas_del_ocr(ocr, nativo), ["Total a pagar 1.250.000"])

    def test_ocr_lines_that_differ_in_more_than_accents_are_kept(self):
        # Cubre: T-124 (quitar tildes no hace iguales dos palabras distintas)
        nativo = "Cotización para el embalse"
        self.assertEqual(lineas_nuevas_del_ocr("Cotizacion para el vaso del embalse", nativo),
                         ["Cotizacion para el vaso del embalse"])

    def test_short_table_cells_contained_in_the_native_text_are_kept(self):
        # Cubre: T-124 (las celdas «Quimbo», «1» y «150» se tiraban por estar dentro de «embalse de Quimbo… 1500»)
        nativo = "Para el embalse de Quimbo, de unas 1500 hectáreas, quedo pendiente"
        nuevas = lineas_nuevas_del_ocr(self.CELDAS, nativo)
        self.assertEqual(nuevas, self.CELDAS.splitlines())

    def test_a_long_line_already_in_the_native_text_is_not_repeated(self):
        # Cubre: T-124
        nativo = "Solicito la cotizacion de los trabajos de campo con las plantillas adjuntas"
        ocr = "Solicito la cotizacion de los trabajos de campo con las plantillas adjuntas\nTotal 5 ha"
        self.assertEqual(lineas_nuevas_del_ocr(ocr, nativo), ["Total 5 ha"])

    def test_a_short_line_equal_to_a_whole_native_line_is_not_repeated(self):
        # Cubre: T-124
        self.assertEqual(lineas_nuevas_del_ocr("Gmail\nTotal 1 ha", "Gmail\nCuerpo del mensaje"), ["Total 1 ha"])

    def test_five_loose_figures_are_a_table_in_an_image_but_a_phone_number_is_not(self):
        # Cubre: T-124
        self.assertTrue(hay_tabla_en_imagen(["120", "17", "60", "23", "2,3"]))
        self.assertFalse(hay_tabla_en_imagen(["120", "17", "60", "23"]))
        self.assertFalse(hay_tabla_en_imagen(["3144861040", "3017242613", "3219560403", "3103283569", "3227671944"]))

    def test_a_null_next_to_a_letter_becomes_ti_and_an_isolated_one_is_removed(self):
        # Cubre: T-124
        self.assertEqual(reparar_ligaduras_perdidas("la co\x00zación de sep\x00embre"),
                         ("la cotización de septiembre", 2))
        self.assertEqual(reparar_ligaduras_perdidas("que \x00enen más y Mar\x00nez"), ("que tienen más y Martinez", 2))
        self.assertEqual(reparar_ligaduras_perdidas("uno \x00 dos"), ("uno  dos", 0))
        self.assertEqual(reparar_ligaduras_perdidas("sin nulos"), ("sin nulos", 0))

    def test_a_table_in_an_image_keeps_every_cell_and_is_flagged(self):
        # Cubre: T-124
        nativo = "Para el embalse de Quimbo, de unas 1500 hectareas"
        with patch.object(lectores_pdf, "ocr_pil_image", return_value=self.CELDAS):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto([nativo]))
        for celda in ("Quimbo", "150", "35"):
            self.assertIn(f"\n{celda}\n", texto)
        self.assertIn("TABLA EN IMAGEN", texto)

    def test_fewer_than_five_loose_figures_raise_no_table_warning(self):
        # Cubre: T-124
        with patch.object(lectores_pdf, "ocr_pil_image", return_value="Total\n12\n34"):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto(["Cuadro de áreas"]))
        self.assertNotIn("TABLA EN IMAGEN", texto)

    def test_a_pdf_with_lost_ti_ligatures_is_repaired_and_says_so(self):
        # Cubre: T-124
        pagina = MagicMock()
        pagina.extract_text.return_value = ("Solicito la co\x00zación de los trabajos de campo para el parque solar, "
                                            "con los estudios de sep\x00embre y las plantillas adjuntas.")
        lector = MagicMock()
        lector.pages = [pagina]
        with patch("pypdf.PdfReader", return_value=lector):
            texto = lectores_pdf.pdf_text(doc.pdf_con_texto(["x"]))
        self.assertIn("cotización de los trabajos", texto)
        self.assertIn("septiembre", texto)
        self.assertNotIn("\x00", texto)
        self.assertIn("LIGADURAS RECUPERADAS: 2", texto)

    def test_a_repair_note_does_not_make_a_nearly_empty_page_look_like_it_has_text(self):
        # Cubre: T-124 (el aviso no cuenta como texto nativo: una página casi vacía sigue pasando por OCR)
        pagina = MagicMock()
        pagina.extract_text.return_value = "Plano\x00s"
        lector = MagicMock()
        lector.pages = [pagina]
        with patch("pypdf.PdfReader", return_value=lector), \
                patch.object(lectores_pdf, "ocr_pil_image", return_value="Cuadro de áreas") as ocr:
            lectores_pdf.pdf_text(doc.pdf_con_texto(["x"]))
        ocr.assert_called_once()

    def test_the_new_notes_make_the_file_a_partial_reading_with_their_own_reason(self):
        # Cubre: T-124
        motivos = attachments._motivos_de_lectura_parcial(
            "[Página 1]\nTABLA EN IMAGEN: x\nLIGADURAS RECUPERADAS: 1 «ti» restauradas")
        self.assertTrue(any("tabla en imagen" in motivo for motivo in motivos))
        self.assertTrue(any("sílabas «ti»" in motivo for motivo in motivos))


class ImagenTests(unittest.TestCase):
    def setUp(self):
        lectores_imagen.ocr_engine.cache_clear()
        self.addCleanup(lectores_imagen.ocr_engine.cache_clear)

    def motor(self, resultado):
        return patch.object(lectores_imagen, "ocr_engine", return_value=lambda imagen: (resultado, None))

    def test_low_confidence_and_empty_lines_are_discarded(self):
        resultado = [([0], "Área 12 ha", 0.9), ([0], "ruido", 0.1), ([0], "   ", 0.99), ([0], "Plazo 40 días", 0.5)]
        with self.motor(resultado):
            texto = lectores_imagen.ocr_image(doc.PNG_1X1)
        self.assertEqual(texto, "Área 12 ha\nPlazo 40 días")

    def test_an_image_without_detectable_text_requires_visual_inspection(self):
        with self.motor(None), self.assertRaises(Unreadable) as caught:
            lectores_imagen.ocr_image(doc.PNG_1X1)
        self.assertIn("inspección visual", str(caught.exception))

    def test_image_text_names_the_file(self):
        with self.motor([([0], "Texto", 0.9)]):
            self.assertEqual(lectores_imagen.image_text("ruta/plano.png", doc.PNG_1X1),
                             "[Imagen: plano.png; texto OCR]\nTexto")

    def test_bytes_that_are_not_an_image_are_reported(self):
        with self.assertRaises(Unreadable) as caught:
            lectores_imagen.ocr_image(b"no es una imagen")
        self.assertIn("no se pudo ejecutar OCR", str(caught.exception))

    def test_an_image_over_the_pixel_limit_is_refused(self):
        with patch.object(lectores_imagen, "MAX_OCR_PIXELS", 0), self.assertRaises(Unreadable) as caught:
            lectores_imagen.ocr_image(doc.PNG_1X1)
        self.assertIn("límite de píxeles", str(caught.exception))

    def test_a_failing_engine_and_a_missing_engine_are_clear_limitations(self):
        def falla(imagen):
            raise RuntimeError("modelo dañado")

        with patch.object(lectores_imagen, "ocr_engine", return_value=falla), self.assertRaises(Unreadable):
            lectores_imagen.ocr_image(doc.PNG_1X1)
        with patch.dict(sys.modules, {"rapidocr_onnxruntime": None}), self.assertRaises(Unreadable) as caught:
            lectores_imagen.ocr_engine()
        self.assertIn("rapidocr", str(caught.exception))


class DocumentosAntiguosTests(unittest.TestCase):
    def test_without_a_converter_it_asks_to_install_one(self):
        with patch.object(lectores_imagen.shutil, "which", return_value=None), self.assertRaises(Unreadable) as caught:
            lectores_imagen.office_legacy_text(".ppt", b"x")
        self.assertIn("LibreOffice", str(caught.exception))

    def test_antiword_converts_old_word_files(self):
        resultado = MagicMock(returncode=0, stdout=b"Contenido del .doc")
        with patch.object(lectores_imagen.shutil, "which", side_effect=lambda n: "antiword" if n == "antiword" else None
                          ), patch.object(lectores_imagen.subprocess, "run", return_value=resultado):
            self.assertIn("Contenido del .doc", lectores_imagen.office_legacy_text(".doc", b"x"))
        with patch.object(lectores_imagen.shutil, "which", side_effect=lambda n: "antiword" if n == "antiword" else None
                          ), patch.object(lectores_imagen.subprocess, "run", return_value=MagicMock(returncode=1)
                                          ), self.assertRaises(Unreadable):
            lectores_imagen.office_legacy_text(".doc", b"x")

    def test_libreoffice_conversion_success_and_failure(self):
        def convierte(comando, **kwargs):
            carpeta = comando[comando.index("--outdir") + 1]
            from pathlib import Path
            (Path(carpeta) / "documento.txt").write_text("Texto convertido", encoding="utf-8")
            return MagicMock(returncode=0)

        with patch.object(lectores_imagen.shutil, "which", side_effect=lambda n: "soffice" if n == "soffice" else None
                          ), patch.object(lectores_imagen.subprocess, "run", side_effect=convierte):
            self.assertIn("Texto convertido", lectores_imagen.office_legacy_text(".xls", b"x"))
        with patch.object(lectores_imagen.shutil, "which", side_effect=lambda n: "soffice" if n == "soffice" else None
                          ), patch.object(lectores_imagen.subprocess, "run", return_value=MagicMock(returncode=1)
                                          ), self.assertRaises(Unreadable) as caught:
            lectores_imagen.office_legacy_text(".xls", b"x")
        self.assertIn("no pudo convertir", str(caught.exception))


class ShapefileTests(unittest.TestCase):
    def shapefile(self, con_prj=True, puntos=3, completo=True):
        import shapefile

        shp, shx, dbf = io.BytesIO(), io.BytesIO(), io.BytesIO()
        with shapefile.Writer(shp=shp, shx=shx, dbf=dbf) as escritor:
            escritor.field("nombre", "C")
            for n in range(puntos):
                escritor.point(-74 + n / 100, 4 + n / 100)
                escritor.record(f"punto{n}")
        archivos = {"capa.shp": shp.getvalue(), "capa.shx": shx.getvalue()}
        if completo:
            archivos["capa.dbf"] = dbf.getvalue()
        if con_prj:
            archivos["capa.prj"] = 'GEOGCS["WGS 84"]'
        return doc.zip_de(archivos)

    def test_layers_fields_extent_records_and_projection_are_reported(self):
        texto = lectores_zip.shapefile_zip_text(self.shapefile())
        self.assertIn("Capa shapefile: capa; geometría: POINT; elementos: 3; campos: nombre", texto)
        self.assertIn("Extensión XY", texto)
        self.assertIn("Registro 1: nombre=punto0", texto)
        self.assertIn("Sistema de coordenadas (PRJ): GEOGCS", texto)

    def test_a_missing_projection_is_flagged(self):
        self.assertIn("SISTEMA DE COORDENADAS NO DECLARADO", lectores_zip.shapefile_zip_text(self.shapefile(con_prj=False)))

    def test_attribute_records_are_capped(self):
        with patch.object(lectores_zip, "MAX_EVIDENCE_ITEMS", 2):
            texto = lectores_zip.shapefile_zip_text(self.shapefile(puntos=5))
        self.assertIn("LECTURA PARCIAL: se muestran 2 registros de atributos", texto)
        self.assertNotIn("Registro 3", texto)

    def test_an_incomplete_shapefile_and_a_missing_library_are_clear_limitations(self):
        incompleto, completo = self.shapefile(completo=False), self.shapefile()
        texto = lectores_zip.shapefile_zip_text(incompleto)       # U-55: se lee lo que haya y se marca lo que falta
        self.assertIn("LECTURA PARCIAL: falta el .dbf", texto)
        with patch.dict(sys.modules, {"shapefile": None}), self.assertRaises(Unreadable) as caught:
            lectores_zip.shapefile_zip_text(completo)
        self.assertIn("pyshp", str(caught.exception))


if __name__ == "__main__":
    unittest.main()
