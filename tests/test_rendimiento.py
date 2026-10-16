"""El dashboard sigue ágil con muchos casos: agrupar, Excel, índices, listas de a 200, gzip, caché y avisos (U-48).

Las pruebas de tiempo usan datos sintéticos y umbrales muy holgados (varias veces lo medido): detectan un regreso a lo
cuadrático, no una diferencia de milisegundos. Cubre: U-48
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import gzip
import http.client
import io
import json
import random
import re
import sqlite3
import tempfile
import threading
import time
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from anka.dominio.activities import ACTIVITY_COLUMNS, _cell, build_workbook
from anka.dominio.status import QuoteStatus
from anka.infra import settings, store
from anka.servicios import proyectos_repetidos as repetidos
from anka.servicios.proyectos_repetidos import agrupar, agrupar_en_cache, agrupar_lento
from anka.web import estaticos, paginas
from anka.web import server as web
from anka.web.marcador import MarcadorCompartido
from tests.test_app import complete_analysis, test_config
from tests.test_seguridad_http import atender


# --------------------------------------------------------------------------------------------- datos sintéticos
def casos_sinteticos(n: int, semilla: int = 7) -> list[dict]:
    """Casos con clientes, códigos, números y adjuntos repetidos, para que haya grupos de todos los motivos."""
    azar = random.Random(semilla)
    palabras = ["Levantamiento", "Topografía", "Calle", "Carrera", "Vía", "Predio", "Lote", "Norte", "Sur", "Batimetría"]
    clientes = ["Constructora Andina SAS", "Constructora Andina S.A.S.", "Inmobiliaria Sol", "OTACC (razón social no indicada)",
                "", "ana@x.com", "Ingeniería Delta Ltda", "Grupo Orión"]
    archivos = [("Plano 8.pdf", 41000), ("Anexo.xlsx", 1200), ("firma.png", 0), ("E.T Llanos.pdf", 9000)]
    casos = []
    for i in range(n):
        nombre = " ".join(azar.sample(palabras, 2)) + (f" {azar.randint(1, 25)}" if azar.random() < 0.7 else "")
        if azar.random() < 0.15:
            nombre = nombre.replace("Topografía", "Topografia").replace("Levantamiento", "Levantamiento")
        adjuntos = [{"filename": a, "extracted_chars": c} for a, c in azar.sample(archivos, azar.randint(0, 2))]
        casos.append({
            "thread_id": f"t{i:05d}", "client_name": azar.choice(clientes),
            "client_email": f"c{azar.randint(1, 60)}@ejemplo.test" if azar.random() < 0.8 else "",
            "quote_code": f"T-26-{azar.randint(1, 40):04d}" if azar.random() < 0.2 else "",
            "analysis": {"project_name": nombre if azar.random() < 0.95 else "", "attachments": adjuntos},
        })
    return casos


def forma(grupos) -> list[tuple[tuple[str, ...], str]]:
    return [(tuple(m["thread_id"] for m in g.miembros), g.motivo) for g in grupos]


class AgruparProyectos(unittest.TestCase):
    def test_da_los_mismos_grupos_que_comparar_todos_contra_todos(self):
        # Cubre: U-48
        casos = casos_sinteticos(800)
        separados = {frozenset((casos[i]["thread_id"], casos[i + 3]["thread_id"])) for i in range(0, 400, 7)}
        for pares in (set(), separados):
            with self.subTest(separados=len(pares)):
                rapido, lento = agrupar(casos, pares), agrupar_lento(casos, pares)
                self.assertEqual(forma(rapido), forma(lento))
        grupos = agrupar(casos)
        self.assertGreater(sum(len(g.miembros) > 1 for g in grupos), 5)                # los datos sí agrupan
        self.assertEqual({g.motivo for g in grupos if len(g.miembros) > 1},
                         {repetidos.MOTIVO_ARCHIVO, repetidos.MOTIVO_CLIENTE, repetidos.MOTIVO_NOMBRE})

    def test_la_equivalencia_se_mantiene_con_otras_semillas_y_tamanos(self):
        # Cubre: U-48
        for semilla, n in ((1, 40), (2, 150), (3, 300)):
            with self.subTest(semilla=semilla):
                casos = casos_sinteticos(n, semilla)
                self.assertEqual(forma(agrupar(casos)), forma(agrupar_lento(casos)))

    def test_presupuesto_de_tiempo_con_muchos_casos(self):
        # Cubre: U-48
        casos = casos_sinteticos(600)
        inicio = time.perf_counter()
        agrupar(casos)
        self.assertLess(time.perf_counter() - inicio, 3.0, "agrupar 600 casos debe tardar una fracción de segundo")


class AgruparEnCache(unittest.TestCase):
    def setUp(self):
        repetidos._cache = None
        self.addCleanup(setattr, repetidos, "_cache", None)

    def contando(self):
        """Parchea `agrupar` para contar cuántas veces se calcula de verdad."""
        llamadas = []
        original = repetidos.agrupar

        def contar(quotes, separados=()):
            llamadas.append(1)
            return original(quotes, separados)
        return llamadas, patch.object(repetidos, "agrupar", contar)

    def test_sin_cambios_no_vuelve_a_comparar_y_devuelve_los_casos_de_ahora(self):
        # Cubre: U-48
        casos = casos_sinteticos(120)
        llamadas, parche = self.contando()
        with parche:
            primero = agrupar_en_cache(casos, set(), "m1")
            copia = [dict(c) for c in casos]                       # mismos datos, objetos nuevos (otra lectura)
            segundo = agrupar_en_cache(copia, set(), "m1")
        self.assertEqual(len(llamadas), 1)
        self.assertEqual(forma(primero), forma(segundo))
        self.assertTrue(all(any(m is c for c in copia) for g in segundo for m in g.miembros))

    def test_cambia_el_marcador_el_cliente_el_codigo_el_nombre_un_adjunto_o_los_pares_separados(self):
        # Cubre: U-48
        casos = casos_sinteticos(120)
        llamadas, parche = self.contando()
        cambios = {
            "marcador": lambda c, s: ("m2", c, s),
            "cliente": lambda c, s: ("m1", [{**c[0], "client_name": "Otro SAS"}] + c[1:], s),
            "codigo": lambda c, s: ("m1", [{**c[0], "quote_code": "T-26-0999"}] + c[1:], s),
            "nombre": lambda c, s: ("m1", [{**c[0], "analysis": {**c[0]["analysis"], "project_name": "Nombre nuevo 99"}}] + c[1:], s),
            "adjunto": lambda c, s: ("m1", [{**c[0], "analysis": {**c[0]["analysis"], "attachments": [
                {"filename": "nuevo.pdf", "extracted_chars": 5}]}}] + c[1:], s),
            "separado": lambda c, s: ("m1", c, {frozenset((c[0]["thread_id"], c[1]["thread_id"]))}),
            "casos": lambda c, s: ("m1", c[:-1], s),
        }
        with parche:
            agrupar_en_cache(casos, set(), "m1")
            for nombre, cambio in cambios.items():
                with self.subTest(cambio=nombre):
                    antes = len(llamadas)
                    marcador, nuevos, separados = cambio(casos, set())
                    resultado = agrupar_en_cache(nuevos, separados, marcador)
                    self.assertEqual(len(llamadas), antes + 1, "debía recalcular")
                    self.assertEqual(forma(resultado), forma(agrupar_lento(nuevos, separados)))
                    agrupar_en_cache(casos, set(), "m1")                # vuelve al estado base para el siguiente cambio


# --------------------------------------------------------------------------------------------------------- Excel
def _build_workbook_anterior(quotes) -> bytes:
    """La versión de antes de U-48 (celda a celda, con `max_row`): la referencia de la prueba de equivalencia."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet

    fills = {
        "confirmado": PatternFill("solid", fgColor="D9EAD3"),
        "pendiente": PatternFill("solid", fgColor="FFF2CC"),
        "contradictorio": PatternFill("solid", fgColor="F4CCCC"),
    }
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.title = "Actividades"
    headers = ["Hilo", "Proyecto", "Cliente", "Estado del caso"] + [label for _, label in ACTIVITY_COLUMNS]
    sheet.append(headers)
    for header_cell in sheet[1]:
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.fill = PatternFill("solid", fgColor="126A62")
        header_cell.alignment = Alignment(vertical="center", wrap_text=True)
    for quote in quotes:
        analysis = quote.get("analysis") or {}
        base = [quote.get("thread_id", ""), analysis.get("project_name") or "", quote.get("client_name") or "",
                quote.get("status_label") or quote.get("status") or ""]
        for item in analysis.get("activities") or [{}]:
            row = [_cell(value) for value in base] + [
                _cell(item.get(key) or ("no indicado" if key != "estado" and item else ""))
                for key, _ in ACTIVITY_COLUMNS]
            sheet.append(row)
            state_cell = sheet.cell(row=sheet.max_row, column=len(base) + 9)
            fill = fills.get(str(state_cell.value))
            if fill:
                state_cell.fill = fill
    widths = [14, 28, 24, 26, 34, 22, 22, 12, 12, 26, 26, 30, 14, 30]
    for position, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(position)].width = width
    for sheet_row in sheet.iter_rows(min_row=2):
        for body_cell in sheet_row:
            body_cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def casos_para_excel(n: int, actividades: int = 6) -> list[dict]:
    azar = random.Random(3)
    estados = ["confirmado", "pendiente", "contradictorio", "", "raro"]
    casos = []
    for i in range(n):
        acts = [{"activity": f"=Actividad {i}-{j}" if j == 0 else f"Actividad {i}-{j}", "ubicacion": "Bogotá" if j % 2 else "",
                 "servicio_tecnica": "Topografía", "quantity": str(j * 3), "unit": "ha", "estado": azar.choice(estados),
                 "notes": "  +nota" if j == 1 else ""} for j in range(azar.randint(0, actividades))]
        casos.append({"thread_id": f"t{i:04d}", "client_name": f"Cliente {i % 7}", "status": "review",
                      "status_label": "Revisión pendiente" if i % 2 else "", "analysis": {
                          "project_name": f"Proyecto {i}", "activities": acts}})
    casos.append({"thread_id": "sin-analisis", "analysis": None})
    return casos


def estilo(celda) -> tuple:
    return (celda.font.b, celda.font.color.rgb if celda.font.color else None, celda.fill.fgColor.rgb if celda.fill.fill_type else None,
            celda.alignment.vertical, celda.alignment.wrap_text)


class ExcelLineal(unittest.TestCase):
    def test_el_contenido_y_el_formato_esenciales_son_los_de_antes(self):
        # Cubre: U-48
        casos = casos_para_excel(60)
        nuevo = load_workbook(io.BytesIO(build_workbook(casos))).active
        viejo = load_workbook(io.BytesIO(_build_workbook_anterior(casos))).active
        self.assertEqual((nuevo.title, nuevo.max_row, nuevo.max_column), (viejo.title, viejo.max_row, viejo.max_column))
        self.assertGreater(nuevo.max_row, 100)
        for fila_nueva, fila_vieja in zip(nuevo.iter_rows(), viejo.iter_rows(), strict=True):
            for a, b in zip(fila_nueva, fila_vieja, strict=True):
                self.assertEqual(a.value, b.value, a.coordinate)
                self.assertEqual(estilo(a), estilo(b), a.coordinate)
        self.assertEqual((nuevo.freeze_panes, nuevo.auto_filter.ref), (viejo.freeze_panes, viejo.auto_filter.ref))
        self.assertEqual({k: d.width for k, d in nuevo.column_dimensions.items()},
                         {k: d.width for k, d in viejo.column_dimensions.items()})

    def test_las_formulas_del_cliente_siguen_neutralizadas(self):
        # Cubre: U-48
        hoja = load_workbook(io.BytesIO(build_workbook(casos_para_excel(3, 3)))).active
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                self.assertFalse(str(celda.value or "").startswith("="), celda.coordinate)

    def test_sin_casos_es_un_libro_valido_con_solo_el_encabezado(self):
        # Cubre: U-48
        hoja = load_workbook(io.BytesIO(build_workbook([]))).active
        self.assertEqual((hoja.max_row, hoja.max_column, hoja.auto_filter.ref), (1, 14, "A1:N1"))

    def test_presupuesto_de_tiempo_y_crecimiento_en_linea_recta(self):
        # Cubre: U-48
        casos = casos_para_excel(500, 12)
        inicio = time.perf_counter()
        datos = build_workbook(casos)
        total = time.perf_counter() - inicio
        filas = load_workbook(io.BytesIO(datos)).active.max_row
        self.assertGreater(filas, 2500)
        self.assertLess(total, 8.0, f"{filas} filas deben escribirse en pocos segundos (antes: cuadrático)")
        inicio = time.perf_counter()
        build_workbook(casos[:250])
        mitad = time.perf_counter() - inicio
        self.assertLess(total, 4 * mitad + 1.0, "al doblar las filas el tiempo no debe cuadruplicarse")


# ---------------------------------------------------------------------------------------------- base de datos
class Indices(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)

    def plan(self, db, sql: str) -> str:
        return " | ".join(fila[3] for fila in db.execute("EXPLAIN QUERY PLAN " + sql))

    def test_una_base_en_el_esquema_16_pasa_al_17_sin_perder_datos(self):
        # Cubre: U-48
        db = store.open_db(self.dir)
        store.save(db, "t-1", "m-1", QuoteStatus.REVIEW_PENDING, "a@x.co", complete_analysis())
        db.commit()
        nuevos = {"idx_quotes_recientes", "idx_quotes_estado", "idx_quotes_codigo", "idx_case_events_tipo",
                  "idx_case_events_registrado", "idx_case_changes_estado", "idx_case_links_estado",
                  "idx_case_links_mensaje", "idx_notifications_sin_leer"}

        def existentes():
            return {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='index' AND name LIKE 'idx_%'")}
        self.assertTrue(nuevos <= existentes())
        for nombre in nuevos:
            db.execute(f"DROP INDEX {nombre}")
        db.execute("PRAGMA user_version = 16")
        db.commit()
        db.close()
        db = store.open_db(self.dir)                                    # migra de 16 a 17
        self.assertEqual(db.execute("PRAGMA user_version").fetchone()[0], store.SCHEMA_VERSION)
        self.assertGreaterEqual(store.SCHEMA_VERSION, 17)
        self.assertTrue(nuevos <= existentes())
        self.assertEqual(db.execute("SELECT COUNT(*) FROM quotes").fetchone()[0], 1)
        store._migration_17_indices(db)                                 # repetirla no falla (IF NOT EXISTS)
        db.close()

    def test_las_consultas_de_la_lista_y_del_marcador_usan_los_indices(self):
        # Cubre: U-48
        db = store.open_db(self.dir)
        consultas = {
            "idx_quotes_recientes": "SELECT * FROM quotes WHERE status <> 'ignored' "
                                    "ORDER BY COALESCE(updated_at, first_seen_at, '') DESC, thread_id DESC",
            "idx_case_events_tipo": "SELECT thread_id, ref, occurred_at FROM case_events WHERE kind='received' "
                                    "ORDER BY occurred_at, id",
            "idx_case_events_registrado": "SELECT MAX(recorded_at) FROM case_events",
            "idx_case_changes_estado": "SELECT * FROM case_changes WHERE state IN ('pending', 'change') ORDER BY occurred_at",
            "idx_case_links_estado": "SELECT * FROM case_links WHERE state='pending' ORDER BY occurred_at",
            "idx_quotes_estado": "SELECT thread_id FROM quotes WHERE status='review_pending'",
            "idx_quotes_codigo": "SELECT 1 FROM quotes WHERE quote_code='T-26-0001'",
            "idx_notifications_sin_leer": "SELECT COUNT(*) FROM dashboard_notifications WHERE read_at IS NULL AND kind IN ('a')",
        }
        for indice, sql in consultas.items():
            with self.subTest(indice=indice):
                self.assertIn(indice, self.plan(db, sql))
        db.close()

    def test_el_marcador_no_cambia_de_valor_con_el_indice(self):
        # Cubre: U-48
        with patch.object(settings, "DATA_DIR", self.dir):
            from anka.servicios import quotes
            vacio = quotes.change_marker()
            db = store.open_db(self.dir)
            store.save(db, "t-1", "m-1", QuoteStatus.REVIEW_PENDING, "a@x.co", complete_analysis())
            db.commit()
            db.close()
            self.assertRegex(vacio, r"^0\|\|0\|")
            self.assertRegex(quotes.change_marker(), r"^1\|20\d\d-.+\|\d+\|")

    def test_las_novedades_derivadas_conservan_su_orden_de_siempre(self):
        # Cubre: U-48
        db = store.open_db(self.dir)
        for i, estado in enumerate([QuoteStatus.REVIEW, QuoteStatus.REVIEW_PENDING, QuoteStatus.READY_REVIEW,
                                    QuoteStatus.REVIEW_PENDING]):
            store.save(db, f"t-{i}", f"m-{i}", estado, "a@x.co", complete_analysis())
        db.commit()
        store._derive_notifications(db)
        db.commit()
        orden = [r[0] for r in db.execute("SELECT thread_id FROM dashboard_notifications WHERE kind='draft_suggested' ORDER BY id")]
        self.assertEqual(orden, sorted(orden))
        db.close()


# ------------------------------------------------------------------------------------------ listas de a 200
class ListasDeA200(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        for p in (patch.object(settings, "DATA_DIR", Path(self._tmp.name)),
                  patch("anka.web.paginas.load_config", return_value=test_config()),
                  patch("anka.web.pantalla_caso.load_config", return_value=test_config())):
            p.start()
            self.addCleanup(p.stop)
        self.db = store.open_db(Path(self._tmp.name))
        self.addCleanup(self.db.close)

    def sembrar(self, n: int) -> None:
        for i in range(n):
            analisis = complete_analysis()
            analisis.update(project_name=f"Obra {i} de prueba", received_at="2026-10-01T09:00:00-05:00")
            store.save(self.db, f"t-{i:04d}", f"m-{i}", QuoteStatus.REVIEW_PENDING, f"c{i}@ejemplo.test", analisis)
        self.db.commit()

    def test_mis_casos_muestra_200_y_ofrece_ver_mas(self):
        # Cubre: U-48
        self.sembrar(450)
        pagina = paginas.render_index()
        self.assertEqual(pagina.count('class="queue-row"'), 200)
        self.assertIn("Se muestran 200 de 450 casos.", pagina)
        self.assertNotIn("miran solo los que se muestran", pagina)          # ahora buscan en todos (U-49)
        self.assertIn('href="/?etapa=atender&amp;mostrar=400"', pagina)       # el enlace conserva el filtro (U-49)
        self.assertIn("Ver 200 más", pagina)
        self.assertIn('aria-label="Más casos"', pagina)
        self.assertIn("450 casos", pagina)                                   # el resumen cuenta todos
        segunda = paginas.render_index(mostrar=400)
        self.assertEqual(segunda.count('class="queue-row"'), 400)
        self.assertIn("Ver 50 más", segunda)
        todas = paginas.render_index(mostrar=600)
        self.assertEqual(todas.count('class="queue-row"'), 450)
        self.assertNotIn("Ver ", todas.split('id="casos"')[1].split("</ol>")[1].split("<details")[0])

    def test_administrar_casos_muestra_200_y_conserva_agrupar_en_el_enlace(self):
        # Cubre: U-48
        self.sembrar(450)
        pagina = paginas.render_requests()
        self.assertEqual(len(re.findall(r'class="case case-row"', pagina)), 200)
        self.assertIn("Se muestran 200 de 450 casos.", pagina)
        self.assertIn('href="/requests?mostrar=400"', pagina)
        suelto = paginas.render_requests(agrupar_hilos=False)
        self.assertIn('href="/requests?agrupar=0&amp;mostrar=400"', suelto)
        self.assertEqual(len(re.findall(r'class="case case-row"', paginas.render_requests(mostrar=450))), 450)

    def test_con_200_o_menos_la_pagina_es_la_de_siempre(self):
        # Cubre: U-48
        self.sembrar(150)
        for pagina in (paginas.render_index(), paginas.render_requests()):
            self.assertNotIn("Se muestran", pagina)
            self.assertNotIn("Ver 200 más", pagina)
            self.assertNotIn("mostrar=", pagina)

    def test_un_valor_raro_de_mostrar_no_rompe_la_pagina(self):
        # Cubre: U-48
        self.sembrar(250)
        for crudo in (-5, 0, 10**12):
            with self.subTest(mostrar=crudo):
                self.assertIn("<ol", paginas.render_index(mostrar=crudo))
        self.assertEqual(paginas.render_index(mostrar=-5).count('class="queue-row"'), 200)


# ------------------------------------------------------------------------------ respuestas, caché y avisos
def servidor_real():
    servidor = ThreadingHTTPServer(("127.0.0.1", 0), web.ReviewHandler)
    servidor.daemon_threads = True
    hilo = threading.Thread(target=servidor.serve_forever, daemon=True)
    hilo.start()
    return servidor, hilo


class RespuestasLivianas(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        config = test_config()
        for p in (patch.object(settings, "DATA_DIR", Path(self._tmp.name)),
                  patch("anka.web.paginas.load_config", return_value=config),
                  patch("anka.web.pantalla_caso.load_config", return_value=config),
                  patch("anka.web.server.load_config", return_value=config)):
            p.start()
            self.addCleanup(p.stop)
        db = store.open_db(Path(self._tmp.name))
        for i in range(60):
            analisis = complete_analysis()
            analisis.update(project_name=f"Obra {i} de prueba")
            store.save(db, f"t-{i:03d}", f"m-{i}", QuoteStatus.REVIEW_PENDING, f"c{i}@ejemplo.test", analisis)
        db.commit()
        db.close()

    def test_gzip_cuando_el_cliente_lo_acepta_y_el_mismo_contenido(self):
        # Cubre: U-48
        for ruta in ("/", "/api/casos", "/openapi.json"):
            with self.subTest(ruta=ruta):
                _, plano, cuerpo_plano, _ = atender("GET", ruta)
                self.assertNotIn("Content-Encoding", plano)
                _, comprimido, _, handler = atender("GET", ruta, cabeceras={"Accept-Encoding": "gzip, deflate"})
                crudo = handler.wfile.getvalue().partition(b"\r\n\r\n")[2]
                self.assertEqual(comprimido["Content-Encoding"], "gzip")
                self.assertEqual(comprimido["Vary"], "Accept-Encoding")
                self.assertEqual(int(comprimido["Content-Length"]), len(crudo))
                self.assertLess(len(crudo), len(cuerpo_plano.encode("utf-8")) / 2)
                self.assertEqual(gzip.decompress(crudo).decode("utf-8").replace("\r\n", "\n"),
                                 cuerpo_plano.replace("\r\n", "\n"))
                self.assertEqual(comprimido["Cache-Control"], "no-store")

    def test_no_comprime_lo_chico_lo_que_ya_esta_comprimido_ni_si_el_cliente_lo_rechaza(self):
        # Cubre: U-48
        _, cabeceras, _, _ = atender("GET", "/state", cabeceras={"Accept-Encoding": "gzip"})
        self.assertNotIn("Content-Encoding", cabeceras)                       # menos de 1 KB
        _, cabeceras, _, _ = atender("GET", "/static/logo-anka.png", cabeceras={"Accept-Encoding": "gzip"})
        self.assertNotIn("Content-Encoding", cabeceras)
        for rechazo in ("identity", "gzip;q=0", "br", ""):
            with self.subTest(rechazo=rechazo):
                self.assertFalse(estaticos.acepta_gzip(rechazo))
        self.assertTrue(estaticos.acepta_gzip("br, GZIP;q=0.5"))
        self.assertIsNone(estaticos.comprimir(b"x" * 5000, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "gzip"))

    def test_el_json_de_la_api_es_compacto_y_sigue_siendo_el_mismo(self):
        # Cubre: U-48
        _, _, cuerpo, _ = atender("GET", "/api/casos")
        self.assertNotIn("\n", cuerpo)
        self.assertNotIn('": ', cuerpo)
        lista = json.loads(cuerpo)
        self.assertEqual(len(lista), 60)
        self.assertEqual(lista[0]["project_name"][:4], "Obra")
        self.assertEqual(json.loads(atender("GET", "/api/casos/t-001")[2])["thread_id"], "t-001")

    def test_el_css_va_aparte_con_huella_y_cache_larga(self):
        # Cubre: U-48
        _, _, pagina, _ = atender("GET", "/")
        url = estaticos.url_del_estilo()
        self.assertRegex(url, r"^/static/style\.[0-9a-f]{12}\.css$")
        self.assertIn(f'<link rel="stylesheet" href="{url}">', pagina)
        self.assertNotIn("@font-face", pagina)                                # ya no se repite en cada página
        estado, cabeceras, css, _ = atender("GET", url)
        self.assertEqual(estado, 200)
        self.assertEqual(cabeceras["Content-Type"], "text/css; charset=utf-8")
        self.assertEqual(cabeceras["Cache-Control"], "public, max-age=31536000, immutable")
        self.assertIn("@font-face", css)
        self.assertIn("X-Content-Type-Options", cabeceras)
        self.assertIn("default-src 'self'", cabeceras["Content-Security-Policy"])
        self.assertEqual(atender("GET", "/static/style.000000000000.css")[0], 404)         # solo la huella vigente

    def test_etag_y_304_en_los_archivos_fijos(self):
        # Cubre: U-48
        for ruta, cache in ((estaticos.url_del_estilo(), "public, max-age=31536000, immutable"),
                            ("/static/logo-anka.png", "public, max-age=86400"),
                            ("/static/inter-latin-wght.woff2", "public, max-age=86400")):
            with self.subTest(ruta=ruta):
                estado, cabeceras, _, _ = atender("GET", ruta)
                self.assertEqual((estado, cabeceras["Cache-Control"]), (200, cache))
                etag = cabeceras["ETag"]
                self.assertRegex(etag, r'^"[0-9a-f]{12}"$')
                estado, cabeceras, cuerpo, _ = atender("GET", ruta, cabeceras={"If-None-Match": etag})
                self.assertEqual((estado, cuerpo, cabeceras["Content-Length"]), (304, "", "0"))
                self.assertEqual(cabeceras["ETag"], etag)
                self.assertEqual(atender("GET", ruta, cabeceras={"If-None-Match": '"otra"'})[0], 200)
                self.assertEqual(atender("GET", ruta, cabeceras={"If-None-Match": f'W/{etag}, "x"'})[0], 304)

    def test_la_huella_cambia_con_el_contenido_del_css(self):
        # Cubre: U-48
        with tempfile.TemporaryDirectory() as carpeta:
            css = Path(carpeta) / "style.css"
            css.write_text("a { color: red }", encoding="utf-8")
            with patch.object(estaticos, "ESTILO", css):
                antes = estaticos.nombre_del_estilo()
                css.write_text("a { color: blue; }", encoding="utf-8")
                self.assertNotEqual(estaticos.nombre_del_estilo(), antes)
                self.assertIsNone(estaticos.buscar(antes))

    def test_http_1_1_reutiliza_la_conexion_y_cada_respuesta_tiene_longitud(self):
        # Cubre: U-48
        servidor, hilo = servidor_real()
        self.addCleanup(hilo.join, 5)
        self.addCleanup(servidor.server_close)
        self.addCleanup(servidor.shutdown)
        conexion = http.client.HTTPConnection("127.0.0.1", servidor.server_address[1], timeout=10)
        self.addCleanup(conexion.close)
        for ruta in ("/state", "/api/casos", estaticos.url_del_estilo(), "/no-existe", "/"):
            conexion.request("GET", ruta, headers={"Accept-Encoding": "gzip"})
            respuesta = conexion.getresponse()
            cuerpo = respuesta.read()
            self.assertEqual(respuesta.version, 11, ruta)
            self.assertIsNotNone(respuesta.getheader("Content-Length"), ruta)
            self.assertEqual(int(respuesta.getheader("Content-Length")), len(cuerpo), ruta)
        self.assertIsNotNone(conexion.sock, "la misma conexión sirvió las cinco peticiones")

    def test_un_post_rechazado_antes_de_leer_su_cuerpo_cierra_la_conexion(self):
        # Cubre: U-48
        estado, _, _, handler = atender("POST", "/delete", b"thread_id=x", cabeceras={"Host": "evil.example", "Origin": ""})
        self.assertEqual(estado, 403)
        self.assertTrue(handler.close_connection)
        estado, _, _, handler = atender("POST", "/delete", b"thread_id=x", cabeceras={"Origin": "http://evil.example"})
        self.assertEqual(estado, 403)
        self.assertTrue(handler.close_connection)

    def test_el_canal_de_avisos_no_tiene_longitud_y_dice_que_se_cierra(self):
        # Cubre: U-48
        from tests.test_en_vivo import AvisosDelServidor
        texto, handler = AvisosDelServidor().avisar(["a"] * 6)
        self.assertIn("Connection: close", texto)
        self.assertNotIn("Content-Length", texto)
        self.assertTrue(handler.close_connection)


class MarcadorCompartidoTests(unittest.TestCase):
    def test_veinte_lecturas_seguidas_consultan_una_vez_y_pasado_el_segundo_otra(self):
        # Cubre: U-48
        hora = [0.0]
        valores = iter(range(100))
        marcador = MarcadorCompartido(lambda: f"v{next(valores)}", 1.0, reloj=lambda: hora[0])
        self.assertEqual({marcador.valor() for _ in range(20)}, {"v0"})
        self.assertEqual(marcador.calculos, 1)
        hora[0] = 0.99
        self.assertEqual(marcador.valor(), "v0")
        hora[0] = 1.0
        self.assertEqual(marcador.valor(), "v1")
        self.assertEqual(marcador.calculos, 2)

    def test_muchas_pestanas_a_la_vez_comparten_un_solo_calculo(self):
        # Cubre: U-48
        marcador = MarcadorCompartido(lambda: (time.sleep(0.05), "x")[1], 5.0)
        resultados = []
        hilos = [threading.Thread(target=lambda: resultados.append(marcador.valor())) for _ in range(16)]
        [h.start() for h in hilos]
        [h.join() for h in hilos]
        self.assertEqual(resultados, ["x"] * 16)
        self.assertEqual(marcador.calculos, 1)

    def test_si_calcular_falla_no_se_recuerda_el_fallo(self):
        # Cubre: U-48
        intentos = iter([sqlite3.OperationalError("locked"), "ok"])

        def calcular():
            valor = next(intentos)
            if isinstance(valor, Exception):
                raise valor
            return valor
        marcador = MarcadorCompartido(calcular, 5.0)
        with self.assertRaises(sqlite3.OperationalError):
            marcador.valor()
        self.assertEqual(marcador.valor(), "ok")

    def test_el_servidor_usa_el_marcador_compartido_en_el_canal_de_avisos(self):
        # Cubre: U-48
        self.assertIsInstance(web._marcador_compartido, MarcadorCompartido)
        self.assertEqual(web.STREAM_CHECK, 1.0)


if __name__ == "__main__":
    unittest.main()
