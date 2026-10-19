"""Cola de trabajos Python → Work: identidad determinista, hoja `.xlsx` por trabajo y publicación sin duplicados (T-49, D-153).

Todo con una SQLite temporal y un Drive simulado: ninguna prueba toca Google ni datos comerciales reales.

Cubre: T-49
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import contextlib
import io
import json
import re
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from anka.cli import work_jobs
from anka.cli import work_sync as sync
from anka.dominio import trabajo_work as dominio
from anka.infra import settings, store
from anka.servicios import trabajos_work, work_import
from anka.servicios.recepcion import CUENTA
from anka.servicios.trabajos_work import TrabajoRechazado
from tests.test_importacion_aislada import GmailAislado, borrador
from tests.test_work_sync import b64, payload

SHEET_MIME = "application/vnd.google-apps.spreadsheet"
CARPETA = "carpeta-laboratorio"
PENDIENTE = ("pendiente_de_procesar", "pendiente_de_interpretacion_work")


class _Respuesta:
    def __init__(self, valor):
        self.valor = valor

    def execute(self):
        if isinstance(self.valor, Exception):
            raise self.valor
        return self.valor() if callable(self.valor) else self.valor


class ErrorHttp(Exception):
    def __init__(self, status):
        super().__init__(f"HttpError {status}")
        self.resp = SimpleNamespace(status=status)


class DriveDeLaboratorio:
    """Drive simulado con marcas (`appProperties`), papelera y ganchos para fabricar carreras y respuestas perdidas."""

    def __init__(self):
        self.archivos: dict[str, dict] = {}
        self.creados = 0
        self.actualizados: list[str] = []
        self.consultas: list[str] = []
        self.cerrojo = threading.RLock()
        self.antes_de_crear = None          # callable(drive): otro proceso crea justo antes que nosotros
        self.al_entrar_a_crear = None       # callable(): para frenar la creación
        self.perder_respuesta_de_crear = False
        self.fallar_listado: Exception | None = None
        self.ocultar_en_listado = False     # el listado aún no ve lo recién creado
        self.borrados = 0

    def files(self):
        return self

    def sembrar(self, job_id, nombre="duplicado", carpeta=CARPETA, papelera=False, marca=True):
        with self.cerrojo:
            archivo_id = f"sheet-{len(self.archivos) + 1}"
            self.archivos[archivo_id] = {"id": archivo_id, "name": nombre, "parents": [carpeta], "mime": SHEET_MIME,
                                         "props": {dominio.CLAVE_DE_MARCA: job_id} if marca else {},
                                         "trashed": papelera, "bytes": b""}
            return archivo_id

    def list(self, q, fields=None, pageSize=None, pageToken=None):
        self.consultas.append(q)
        if self.fallar_listado:
            return _Respuesta(self.fallar_listado)
        carpeta = re.match(r"^'((?:[^'\\]|\\.)*)' in parents", q).group(1)
        marca = re.search(r"appProperties has \{ key='([^']*)' and value='((?:[^'\\]|\\.)*)' \}", q)

        def listar():
            with self.cerrojo:
                if self.ocultar_en_listado:
                    return {"files": []}
                return {"files": [dict(a) for a in self.archivos.values()
                                  if carpeta in a["parents"] and not a["trashed"] and a["mime"] == SHEET_MIME
                                  and (marca is None or a["props"].get(marca.group(1)) == marca.group(2))]}
        return _Respuesta(listar)

    def get(self, fileId, fields=None):
        def leer():
            archivo = self.archivos.get(fileId)
            if archivo is None:
                raise ErrorHttp(404)
            return {"id": fileId, "trashed": archivo["trashed"], "appProperties": dict(archivo["props"])}
        return _Respuesta(leer)

    def update(self, fileId, body=None, media_body=None, fields=None):
        def actualizar():
            with self.cerrojo:
                if fileId not in self.archivos:
                    raise ErrorHttp(404)
                if media_body is not None:
                    self.archivos[fileId]["bytes"] = media_body.getbytes(0, media_body.size())
                self.actualizados.append(fileId)
                return {"id": fileId, "modifiedTime": "2026-10-17T10:00:00.000Z"}
        return _Respuesta(actualizar)

    def create(self, body, fields=None, media_body=None):
        def crear():
            if self.al_entrar_a_crear:
                self.al_entrar_a_crear()
            if self.antes_de_crear:
                self.antes_de_crear(self)
            with self.cerrojo:
                self.creados += 1
                archivo_id = f"sheet-{len(self.archivos) + 1}"
                self.archivos[archivo_id] = {
                    "id": archivo_id, "name": body["name"], "parents": list(body["parents"]),
                    "mime": body.get("mimeType"), "props": dict(body.get("appProperties") or {}), "trashed": False,
                    "bytes": media_body.getbytes(0, media_body.size()) if media_body is not None else b""}
            if self.perder_respuesta_de_crear:
                raise TimeoutError("se perdió la respuesta")
            return {"id": archivo_id, "modifiedTime": "2026-10-17T10:00:00.000Z"}
        return _Respuesta(crear)

    def vivos(self, job_id=None):
        return [a for a in self.archivos.values() if not a["trashed"]
                and (job_id is None or a["props"].get(dominio.CLAVE_DE_MARCA) == job_id)]


class Base(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.datos = Path(self.temp.name)
        for parche in (patch.object(settings, "DATA_DIR", self.datos),
                       patch.object(work_import, "programar_actualizacion_de_hoja", MagicMock()),
                       patch.object(sync, "sync_thread_activity", MagicMock())):
            parche.start()
            self.addCleanup(parche.stop)
        self.drive = DriveDeLaboratorio()

    def fila(self, hilo="T1", mensaje="m1", estado=PENDIENTE[0], disposicion=PENDIENTE[1], adjuntos=None,
             asunto="Solicitud de levantamiento", remitente="cliente@example.com", cuerpo="CUERPO SECRETO"):
        with store.transaccion() as db:
            db.execute("""INSERT INTO recepcion_mensajes(account,thread_id,message_id,state,disposition,sender,subject,
                header_date,internal_date,snippet,attachments_json,detected_at,updated_at)
                VALUES(?,?,?,?,?,?,?,'Sat, 17 Oct 2026 10:00:00 -0500','2026-10-17T15:00:00+00:00',?,?,
                '2026-10-17T10:00:00-05:00','2026-10-17T10:00:00-05:00')""",
                       (CUENTA, hilo, mensaje, estado, disposicion, remitente, asunto, cuerpo,
                        json.dumps(adjuntos if adjuntos is not None else [
                            {"filename": "plano.pdf", "mime_type": "application/pdf", "size_bytes": 1234,
                             "gmail_attachment_id": "att-1", "storage_path": None, "status": "disponible_en_gmail"}])))

    def trabajo(self, hilo="T1", mensaje="m1", **fila):
        self.fila(hilo, mensaje, **fila)
        return trabajos_work.crear_trabajo(hilo, mensaje)

    def estado(self, job_id):
        return trabajos_work.obtener(job_id)["state"]

    def eventos(self, tipo, job_id=None):
        with store.lectura() as db:
            filas = [dict(r) for r in db.execute("SELECT * FROM recepcion_eventos WHERE event_kind=? ORDER BY id",
                                                 (tipo,))]
        return [f for f in filas if job_id is None or json.loads(f["detail_json"]).get("job_id") == job_id]

    def importar(self, hilo="T1", mensaje="m1", **extra):
        work_import.dashboard_upsert(b64(json.dumps(payload(thread=hilo, message=mensaje, **extra))))


class IdentidadTests(unittest.TestCase):
    def test_el_job_id_es_estable_y_distinto_por_mensaje_tipo_y_version(self):
        base = dominio.id_de_trabajo("a@x", "T1", "m1", dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO)
        self.assertRegex(base, r"^ANKAJOB-[0-9a-f]{16}$")
        self.assertEqual(base, dominio.id_de_trabajo("a@x", "T1", "m1", dominio.TIPO_ANALIZAR,
                                                     dominio.VERSION_DEL_CONTRATO))
        otros = {dominio.id_de_trabajo("a@x", "T1", "m2", dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO),
                 dominio.id_de_trabajo("a@x", "T2", "m1", dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO),
                 dominio.id_de_trabajo("b@x", "T1", "m1", dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO),
                 dominio.id_de_trabajo("a@x", "T1", "m1", "otro_tipo", dominio.VERSION_DEL_CONTRATO),
                 dominio.id_de_trabajo("a@x", "T1", "m1", dominio.TIPO_ANALIZAR, "ANKA_DASHBOARD_JSON_V2")}
        self.assertEqual(len(otros | {base}), 6)

    def test_un_dato_vacio_no_produce_un_identificador(self):
        for datos in (("", "T1", "m1"), ("a@x", " ", "m1"), ("a@x", "T1", "")):
            with self.assertRaises(ValueError):
                dominio.id_de_trabajo(*datos, dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO)

    def test_la_huella_de_la_referencia_cambia_con_cualquier_dato(self):
        ref = dominio.referencia_del_mensaje({"thread_id": "T1", "message_id": "m1", "subject": "A"})
        otra = dominio.referencia_del_mensaje({"thread_id": "T1", "message_id": "m1", "subject": "B"})
        self.assertEqual(dominio.huella_de_la_referencia(ref), dominio.huella_de_la_referencia(dict(ref)))
        self.assertNotEqual(dominio.huella_de_la_referencia(ref), dominio.huella_de_la_referencia(otra))


class CrearTrabajoTests(Base):
    def test_crea_el_trabajo_de_un_mensaje_pendiente_con_su_referencia(self):
        creado = self.trabajo()
        self.assertTrue(creado["creado"])
        self.assertEqual(creado["state"], "pendiente_de_publicar")
        self.assertEqual(creado["job_id"], dominio.id_de_trabajo(CUENTA, "T1", "m1", dominio.TIPO_ANALIZAR,
                                                                 dominio.VERSION_DEL_CONTRATO))
        referencia = creado["referencia"]
        self.assertEqual((referencia["gmail_thread_id"], referencia["gmail_message_id"], referencia["sender"],
                          referencia["subject"]), ("T1", "m1", "cliente@example.com", "Solicitud de levantamiento"))
        self.assertEqual(referencia["attachments"], [{"filename": "plano.pdf", "mime_type": "application/pdf",
                                                      "size_bytes": 1234, "gmail_attachment_id": "att-1"}])
        self.assertEqual(creado["ref_sha256"], dominio.huella_de_la_referencia(referencia))
        self.assertEqual(len(self.eventos("trabajo_creado")), 1)

    def test_la_referencia_no_lleva_cuerpo_ni_texto_de_adjuntos(self):
        creado = self.trabajo(cuerpo="CUERPO SECRETO DEL CLIENTE")
        self.assertNotIn("SECRETO", creado["message_ref_json"])
        self.assertEqual(set(creado["referencia"]), {"gmail_thread_id", "gmail_message_id", "sender", "subject",
                                                     "internal_date", "header_date", "attachments"})

    def test_crear_dos_veces_devuelve_el_mismo_trabajo_sin_duplicar(self):
        primero = self.trabajo()
        segundo = trabajos_work.crear_trabajo("T1", "m1")
        self.assertFalse(segundo["creado"])
        self.assertEqual(primero["job_id"], segundo["job_id"])
        with store.lectura() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM trabajos_work").fetchone()[0], 1)
        self.assertEqual(len(self.eventos("trabajo_creado")), 1)

    def test_se_rechaza_lo_que_no_esta_en_la_cola_o_ya_no_esta_pendiente(self):
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.crear_trabajo("T1", "no-existe")                      # no se inventa una fila
        self.fila("T2", "m2", "procesado", "interpretado_por_work")
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.crear_trabajo("T2", "m2")
        self.fila("T3", "m3", "por_leer", "")
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.crear_trabajo("T3", "m3")
        for hilo, mensaje in (("", "m1"), ("T1", "  ")):
            with self.assertRaises(TrabajoRechazado):
                trabajos_work.crear_trabajo(hilo, mensaje)
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.crear_trabajo("T1", "m1", job_type="inventado")
        with store.lectura() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM trabajos_work").fetchone()[0], 0)

    def test_no_se_crea_trabajo_para_otros_mensajes_del_hilo(self):
        self.fila("T1", "m1")
        self.fila("T1", "m2")
        trabajos_work.crear_trabajo("T1", "m1")
        self.assertEqual(trabajos_work.trabajos_del_mensaje("T1", "m2"), [])


class HojaDelTrabajoTests(Base):
    def hojas(self, job_id):
        return load_workbook(io.BytesIO(trabajos_work.libro_del_trabajo(job_id)))

    def test_la_hoja_trae_identidad_hash_y_adjuntos_esperados(self):
        creado = self.trabajo()
        libro = self.hojas(creado["job_id"])
        self.assertEqual(libro.sheetnames, list(dominio.HOJAS))
        campos = {fila[0].value: fila[1].value for fila in libro["TRABAJO"].iter_rows(min_row=2)}
        self.assertEqual((campos["job_id"], campos["thread_id"], campos["message_id"], campos["ref_sha256"]),
                         (creado["job_id"], "T1", "m1", creado["ref_sha256"]))
        self.assertEqual(campos["version_del_contrato"], "ANKA_DASHBOARD_JSON_V1")
        self.assertEqual(campos["estado_del_trabajo"], "pendiente_de_publicar")
        self.assertIn("solo lectura", campos["instruccion_para_work"])
        adjunto, = list(libro["ADJUNTOS_ESPERADOS"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(adjunto, ("plano.pdf", "application/pdf", 1234, "att-1"))
        self.assertGreaterEqual(libro["VERIFICACION"].max_row, 5)

    def test_la_hoja_no_trae_cuerpo_ni_texto(self):
        creado = self.trabajo(cuerpo="CUERPO SECRETO DEL CLIENTE")
        libro = self.hojas(creado["job_id"])
        celdas = [str(c.value) for hoja in libro.worksheets for fila in hoja.iter_rows() for c in fila if c.value]
        self.assertFalse([c for c in celdas if "SECRETO" in c])

    def test_un_asunto_que_parece_formula_queda_como_texto(self):
        creado = self.trabajo(asunto='=HYPERLINK("http://x","clic")')
        libro = self.hojas(creado["job_id"])
        celda = next(f[1] for f in libro["TRABAJO"].iter_rows(min_row=2) if f[0].value == "asunto")
        self.assertNotEqual(celda.data_type, "f")

    def test_exportar_un_trabajo_inexistente_se_rechaza(self):
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.libro_del_trabajo("ANKAJOB-0000000000000000")


class PublicacionTests(Base):
    def test_publicar_crea_con_marca_verifica_y_deja_el_trabajo_publicado(self):
        job = self.trabajo()["job_id"]
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(resultado["resultado"], "creado")
        archivo, = self.drive.vivos(job)
        self.assertEqual((archivo["name"], archivo["parents"], archivo["mime"]),
                         (f"ANKA_JOB_{job}", [CARPETA], SHEET_MIME))
        trabajo = trabajos_work.obtener(job)
        self.assertEqual((trabajo["state"], trabajo["sheet_file_id"], trabajo["folder_id"], trabajo["attempts"]),
                         ("publicado", archivo["id"], CARPETA, 1))
        self.assertEqual((trabajo["claim_token"], trabajo["claim_until"], trabajo["last_error"]), ("", None, ""))
        self.assertTrue(trabajo["published_at"])
        self.assertEqual(len(self.eventos("trabajo_publicado", job)), 1)
        self.assertEqual(load_workbook(io.BytesIO(archivo["bytes"]))["TRABAJO"]["B2"].value, job)

    def test_republicar_actualiza_el_mismo_archivo_sin_crear_otro(self):
        job = self.trabajo()["job_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        otra = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(otra["resultado"], "actualizado")
        self.assertEqual((self.drive.creados, len(self.drive.vivos(job)), len(self.drive.actualizados)), (1, 1, 1))
        self.assertEqual(len(self.eventos("trabajo_publicado", job)), 1)

    def test_dos_publicaciones_simultaneas_dejan_un_solo_archivo(self):
        job = self.trabajo()["job_id"]
        dentro, seguir = threading.Event(), threading.Event()
        self.drive.al_entrar_a_crear = lambda: (dentro.set(), seguir.wait(10))
        resultados = []
        hilo = threading.Thread(target=lambda: resultados.append(
            trabajos_work.publicar_trabajo(self.drive, job, CARPETA)))
        hilo.start()
        self.assertTrue(dentro.wait(10))
        segundo = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)       # mientras el primero está creando
        seguir.set()
        hilo.join(10)
        self.assertEqual(segundo["resultado"], "no_reclamado")
        self.assertEqual(resultados[0]["resultado"], "creado")
        self.assertEqual((self.drive.creados, len(self.drive.vivos(job))), (1, 1))

    def test_una_respuesta_perdida_tras_crear_se_adopta_en_el_reintento(self):
        job = self.trabajo()["job_id"]
        self.drive.perder_respuesta_de_crear = True
        primera = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(primera["resultado"], "error")
        trabajo = trabajos_work.obtener(job)
        self.assertEqual(trabajo["state"], "pendiente_de_publicar")
        self.assertEqual(trabajo["last_error"], "error_de_drive:TimeoutError")
        self.assertEqual(len(self.drive.vivos(job)), 1)                          # el archivo sí se creó
        self.drive.perder_respuesta_de_crear = False
        segunda = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(segunda["resultado"], "adoptado")
        self.assertEqual((self.drive.creados, len(self.drive.vivos(job))), (1, 1))
        self.assertEqual(trabajos_work.obtener(job)["state"], "publicado")

    def test_dos_archivos_con_la_marca_dejan_el_trabajo_ambiguo_sin_borrar_nada(self):
        job = self.trabajo()["job_id"]
        a, b = self.drive.sembrar(job), self.drive.sembrar(job)
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(resultado["resultado"], "ambiguo")
        self.assertEqual(sorted(resultado["candidatos"]), sorted([a, b]))
        self.assertEqual((self.drive.creados, len(self.drive.actualizados), len(self.drive.vivos(job))), (0, 0, 2))
        self.assertEqual(self.estado(job), "ambiguo")
        self.assertEqual(len(self.eventos("trabajo_ambiguo", job)), 1)
        with self.assertRaises(TrabajoRechazado):                                # no se vuelve a publicar solo
            trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(len(self.drive.vivos(job)), 2)

    def test_otro_proceso_que_crea_justo_antes_deja_ambiguo_y_registra_lo_creado(self):
        job = self.trabajo()["job_id"]
        self.drive.antes_de_crear = lambda drive: drive.sembrar(job, nombre="del otro proceso")
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(resultado["resultado"], "ambiguo")
        self.assertEqual(len(resultado["candidatos"]), 2)
        self.assertEqual(len(self.drive.vivos(job)), 2)                           # no se elige ni se borra
        trabajo = trabajos_work.obtener(job)
        self.assertEqual(trabajo["state"], "ambiguo")
        self.assertTrue(trabajo["sheet_file_id"])

    def test_si_el_listado_no_ve_lo_creado_el_resultado_es_ambiguo_y_guarda_el_id(self):
        job = self.trabajo()["job_id"]
        self.drive.ocultar_en_listado = True
        # primer listado (antes de crear) vacío; el de verificación también: el resultado no es verificable
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(resultado["resultado"], "ambiguo")
        trabajo = trabajos_work.obtener(job)
        self.assertEqual(trabajo["state"], "ambiguo")
        self.assertEqual(trabajo["sheet_file_id"], self.drive.vivos(job)[0]["id"])

    def test_un_error_de_drive_antes_de_cualquier_efecto_deja_el_trabajo_para_reintentar(self):
        job = self.trabajo()["job_id"]
        self.drive.fallar_listado = ErrorHttp(500)
        resultado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(resultado["resultado"], "error")
        trabajo = trabajos_work.obtener(job)
        self.assertEqual((trabajo["state"], trabajo["attempts"], trabajo["claim_token"]),
                         ("pendiente_de_publicar", 1, ""))
        self.assertEqual(self.drive.creados, 0)
        self.drive.fallar_listado = None
        self.assertEqual(trabajos_work.publicar_trabajo(self.drive, job, CARPETA)["resultado"], "creado")
        self.assertEqual(trabajos_work.obtener(job)["attempts"], 2)

    def test_el_error_no_guarda_datos_del_correo(self):
        job = self.trabajo(asunto="ASUNTO CONFIDENCIAL")["job_id"]
        self.drive.fallar_listado = ErrorHttp(500)
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        trabajo = trabajos_work.obtener(job)
        self.assertEqual(trabajo["last_error"], "error_de_drive:ErrorHttp")        # solo el tipo, nunca el texto del error
        with store.lectura() as db:
            eventos = " ".join(r[0] for r in db.execute("SELECT detail_json FROM recepcion_eventos"))
        self.assertNotIn("CONFIDENCIAL", eventos)

    def test_un_reclamo_vigente_bloquea_y_uno_vencido_se_retoma(self):
        job = self.trabajo()["job_id"]
        token, _ = trabajos_work._reclamar(job, 1000.0, "t")                      # otro proceso lo reclamó a t=1000
        self.assertTrue(token)
        bloqueado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA, reloj=lambda: 1000.0 + 10)
        self.assertEqual(bloqueado["resultado"], "no_reclamado")
        self.assertEqual(self.drive.creados, 0)
        retomado = trabajos_work.publicar_trabajo(self.drive, job, CARPETA, reloj=lambda: 1000.0 + 500)
        self.assertEqual(retomado["resultado"], "creado")
        self.assertEqual(self.estado(job), "publicado")

    def test_un_archivo_guardado_en_la_papelera_o_con_otra_marca_no_se_reutiliza(self):
        job = self.trabajo()["job_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        guardado = trabajos_work.obtener(job)["sheet_file_id"]
        self.drive.archivos[guardado]["trashed"] = True
        self.assertEqual(trabajos_work.publicar_trabajo(self.drive, job, CARPETA)["resultado"], "creado")
        nuevo = trabajos_work.obtener(job)["sheet_file_id"]
        self.assertNotEqual(nuevo, guardado)
        self.drive.archivos[nuevo]["props"] = {dominio.CLAVE_DE_MARCA: "ANKAJOB-ajeno"}
        self.assertEqual(trabajos_work.publicar_trabajo(self.drive, job, CARPETA)["resultado"], "creado")
        self.assertEqual(self.drive.creados, 3)

    def test_la_busqueda_usa_la_marca_y_no_adopta_archivos_de_otro_trabajo_ni_solo_por_nombre(self):
        job = self.trabajo()["job_id"]
        self.drive.sembrar("ANKAJOB-otro", nombre=f"ANKA_JOB_{job}")             # mismo nombre, otra marca
        self.drive.sembrar(job, carpeta="otra-carpeta")                           # misma marca, otra carpeta
        self.assertEqual(trabajos_work.publicar_trabajo(self.drive, job, CARPETA)["resultado"], "creado")
        self.assertTrue(all("appProperties has" in q for q in self.drive.consultas))

    def test_no_se_publica_un_trabajo_cerrado_ni_sin_carpeta_ni_inexistente(self):
        job = self.trabajo()["job_id"]
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.publicar_trabajo(self.drive, job, " ")
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.publicar_trabajo(self.drive, "ANKAJOB-0000000000000000", CARPETA)
        with store.transaccion() as db:
            db.execute("UPDATE trabajos_work SET state='cancelado'")
        with self.assertRaises(TrabajoRechazado):
            trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.assertEqual(self.drive.creados, 0)


class RetornoTests(Base):
    def test_importar_el_borrador_concilia_el_mensaje_y_el_trabajo_una_sola_vez(self):
        job = self.trabajo()["job_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        self.importar("T1", "m1")
        trabajo = trabajos_work.obtener(job)
        self.assertEqual((trabajo["state"], bool(trabajo["conciliated_at"])), ("conciliado", True))
        self.assertEqual(len(self.eventos("trabajo_conciliado", job)), 1)
        with store.lectura() as db:
            self.assertEqual(db.execute("SELECT state FROM recepcion_mensajes").fetchone()[0], "procesado")
        despues = trabajos_work.obtener(job)
        self.importar("T1", "m1")
        self.importar("T1", "m1", project_name="otra versión")
        self.assertEqual(trabajos_work.obtener(job), despues)
        self.assertEqual(len(self.eventos("trabajo_conciliado", job)), 1)

    def test_un_trabajo_sin_publicar_tambien_se_cierra_si_el_mensaje_se_interpreto(self):
        job = self.trabajo()["job_id"]
        self.importar("T1", "m1")
        self.assertEqual(self.estado(job), "conciliado")

    def test_no_se_cierra_el_trabajo_de_otro_mensaje_del_hilo_ni_de_otro_hilo(self):
        uno, dos = self.trabajo("T1", "m1")["job_id"], self.trabajo("T1", "m2")["job_id"]
        tres = self.trabajo("T2", "m1")["job_id"]
        self.importar("T1", "m2", summary="seguimiento")
        self.assertEqual([self.estado(uno), self.estado(dos), self.estado(tres)],
                         ["pendiente_de_publicar", "conciliado", "pendiente_de_publicar"])

    def test_un_trabajo_ambiguo_o_cancelado_no_cambia_al_importar(self):
        ambiguo = self.trabajo("T1", "m1")["job_id"]
        cancelado = self.trabajo("T2", "m2")["job_id"]
        with store.transaccion() as db:
            db.execute("UPDATE trabajos_work SET state='ambiguo' WHERE job_id=?", (ambiguo,))
            db.execute("UPDATE trabajos_work SET state='cancelado' WHERE job_id=?", (cancelado,))
        self.importar("T1", "m1")
        self.importar("T2", "m2")
        self.assertEqual([self.estado(ambiguo), self.estado(cancelado)], ["ambiguo", "cancelado"])

    def test_un_caso_minimo_no_cierra_el_trabajo(self):
        job = self.trabajo()["job_id"]
        work_import.dashboard_upsert(b64(json.dumps(payload(thread="T1", message="m1", _caso_minimo=True,
                                                            _rechazo_de_work={"reason": "JSON roto"}))))
        self.assertEqual(self.estado(job), "pendiente_de_publicar")

    def test_si_falla_cerrar_el_trabajo_el_analisis_y_la_cola_no_quedan_a_medias(self):
        job = self.trabajo()["job_id"]
        with patch.object(work_import, "marcar_conciliado", side_effect=__import__("sqlite3").OperationalError("x")):
            self.importar("T1", "m1")
        with store.lectura() as db:
            self.assertEqual(db.execute("SELECT state FROM recepcion_mensajes").fetchone()[0], "pendiente_de_procesar")
            self.assertEqual(db.execute("SELECT COUNT(*) FROM quotes").fetchone()[0], 1)
        self.assertEqual(self.estado(job), "pendiente_de_publicar")
        self.assertEqual(self.eventos("trabajo_conciliado"), [])

    def test_sin_la_tabla_de_trabajos_el_analisis_se_importa(self):
        self.fila()
        with store.transaccion() as db:
            db.execute("DROP TABLE trabajos_work")
        self.importar("T1", "m1")
        with store.lectura() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM quotes").fetchone()[0], 1)

    def test_la_importacion_aislada_cierra_solo_el_trabajo_del_mensaje_elegido(self):
        elegido, otro = self.trabajo("T1", "m1")["job_id"], self.trabajo("T1", "m2")["job_id"]
        gmail = GmailAislado({"dT": borrador(hilo="T1", mensaje="m1")}, mensajes={"m1": {
            "id": "m1", "threadId": "T1", "internalDate": "1700000000000", "labelIds": ["INBOX"],
            "payload": {"mimeType": "text/plain", "headers": [{"name": "From", "value": "Ana <c@example.com>"}]}}})
        resultado = sync.importar_aislado(
            gmail, lambda raw: work_import.dashboard_upsert(raw, actualizar_hoja=False),
            self.datos / "manifiestos", "T1", "m1", {CUENTA})
        self.assertEqual(resultado["resultado"], "importado")
        self.assertEqual([self.estado(elegido), self.estado(otro)], ["conciliado", "pendiente_de_publicar"])


class CliTests(Base):
    def ejecutar(self, *argv):
        salida, error = io.StringIO(), io.StringIO()
        with contextlib.redirect_stderr(error):
            codigo = work_jobs.main(list(argv), salida=lambda texto: salida.write(texto))
        return codigo, salida.getvalue(), error.getvalue()

    def test_crear_ver_y_exportar_no_usan_drive(self):
        self.fila()
        with patch("anka.infra.google_services.services", side_effect=AssertionError("no debe tocar Google")):
            codigo, texto, _ = self.ejecutar("crear", "--thread-id", "T1", "--message-id", "m1")
            self.assertEqual(codigo, 0)
            job = json.loads(texto)["job_id"]
            self.assertNotIn("claim_token", texto)
            self.assertEqual(self.ejecutar("ver", "--job-id", job)[0], 0)
            destino = self.datos / "hoja.xlsx"
            self.assertEqual(self.ejecutar("exportar", "--job-id", job, "--salida", str(destino))[0], 0)
        self.assertEqual(load_workbook(destino)["TRABAJO"]["B2"].value, job)

    def test_publicar_se_niega_sin_confirmacion_y_no_toca_google(self):
        job = self.trabajo()["job_id"]
        with patch("anka.infra.google_services.services", side_effect=AssertionError("no debe tocar Google")):
            codigo, _, error = self.ejecutar("publicar", "--job-id", job, "--carpeta-id", CARPETA)
        self.assertEqual(codigo, 2)
        self.assertIn("--confirmar-drive", error)
        self.assertEqual(self.estado(job), "pendiente_de_publicar")

    def test_con_confirmacion_publica_con_el_drive_que_se_le_da(self):
        job = self.trabajo()["job_id"]
        with patch("anka.infra.google_services.services", return_value=(object(), self.drive)):
            codigo, texto, _ = self.ejecutar("publicar", "--job-id", job, "--carpeta-id", CARPETA, "--confirmar-drive")
        self.assertEqual(codigo, 0)
        self.assertEqual(json.loads(texto)["resultado"], "creado")

    def test_un_trabajo_inexistente_se_rechaza_con_codigo_2(self):
        self.assertEqual(self.ejecutar("ver", "--job-id", "ANKAJOB-0000000000000000")[0], 2)
        self.assertEqual(self.ejecutar("crear", "--thread-id", "T9", "--message-id", "m9")[0], 2)


class MigracionTests(Base):
    def test_la_migracion_crea_la_tabla_y_la_base_queda_integra(self):
        with store.lectura() as db:
            columnas = {r["name"] for r in db.execute("PRAGMA table_info(trabajos_work)")}
            self.assertEqual(db.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertTrue({"job_id", "thread_id", "message_id", "job_type", "state", "contract_version", "created_at",
                         "message_ref_json", "ref_sha256", "sheet_file_id", "claim_token"} <= columnas)

    def test_la_clave_unica_impide_dos_trabajos_del_mismo_mensaje_tipo_y_version(self):
        self.trabajo()
        with self.assertRaises(__import__("sqlite3").IntegrityError), store.transaccion() as db:
            db.execute("""INSERT INTO trabajos_work(job_id,account,thread_id,message_id,job_type,contract_version,state,
                created_at,updated_at) VALUES('ANKAJOB-otro',?,'T1','m1',?,?,'pendiente_de_publicar','t','t')""",
                       (CUENTA, dominio.TIPO_ANALIZAR, dominio.VERSION_DEL_CONTRATO))


if __name__ == "__main__":
    unittest.main()
