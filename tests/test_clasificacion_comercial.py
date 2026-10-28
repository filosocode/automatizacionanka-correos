"""Clasificación comercial: preliminar en Python, confirmada o corregida por Work (T-52, D-156).

SQLite temporal, Drive simulado y mensajes ficticios.

Cubre: T-52
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import json
import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

from anka.dominio import trabajo_work as dominio
from anka.dominio.clasificacion import (
    CLASIFICACIONES, NO_SOLICITUD, REVISION_MANUAL, SEGUIMIENTO, SERVICIOS_DE_ANKA, SOLICITUD_NUEVA,
    clasificar_preliminar, menciona_un_servicio,
)
from anka.dominio.status import QuoteStatus
from anka.dominio.trabajo_work import ELEGIBLE, EXCLUIDO, REVISION, Veredicto, evaluar_elegibilidad
from anka.infra import store
from anka.servicios import quotes, trabajos_work, work_import
from anka.servicios.recepcion import CUENTA
from anka.servicios.trabajos_work import NoElegible
from tests.test_elegibilidad_de_trabajos import CARPETA, TERMINOS
from tests.test_elegibilidad_de_trabajos import Base as BaseDeElegibilidad
from tests.test_work_sync import b64, payload

OK = Veredicto(ELEGIBLE, "", "")


def fila(asunto="", adjuntos=(), remitente="Ana <ana@empresa.com>"):
    return {"sender": remitente, "subject": asunto, "evidence_json": "{}",
            "attachments_json": json.dumps([{"filename": n} for n in adjuntos])}


class DominioTests(unittest.TestCase):
    def clase(self, f, veredicto=OK, expediente=False):
        return clasificar_preliminar(f, veredicto, expediente, TERMINOS)

    def test_el_ruido_de_t51_es_no_solicitud_y_no_se_envia_a_work(self):
        v = evaluar_elegibilidad({**fila("Boletín", remitente="noreply@tm.openai.com"), "evidence_json": "{}"}, TERMINOS)
        self.assertEqual(v.estado, EXCLUIDO)
        c = self.clase(fila("Batimetría"), v)
        self.assertEqual((c.clase, c.motivo, c.enviar_a_work), (NO_SOLICITUD, "notificacion_de_chatgpt", False))

    def test_la_duda_de_t51_es_revision_manual_y_no_se_envia(self):
        c = self.clase(fila("Su cotización"), Veredicto(REVISION, "remitente_desconocido", "x"))
        self.assertEqual((c.clase, c.enviar_a_work), (REVISION_MANUAL, False))

    def test_un_hilo_con_expediente_es_seguimiento_aunque_el_mensaje_no_diga_nada(self):
        for asunto in ("Re: ok", "Gracias", ""):
            c = self.clase(fila(asunto), expediente=True)
            self.assertEqual((c.clase, c.enviar_a_work), (SEGUIMIENTO, True))

    def test_cada_servicio_de_anka_se_reconoce_con_o_sin_tildes_y_en_adjuntos(self):
        for asunto in ("Cotizar topografía", "GEOMÁTICA para predio", "Vuelo LiDAR", "fotogrametria", "Batimetría de río",
                       "Inspección de puente", "Georadar en vía", "Trámite catastral", "Escáner terrestre de planta"):
            with self.subTest(asunto=asunto):
                self.assertEqual(self.clase(fila(asunto)).clase, SOLICITUD_NUEVA)
        self.assertEqual(self.clase(fila("Hola", adjuntos=["Alcance_Levantamiento.pdf"])).clase, SOLICITUD_NUEVA)
        self.assertTrue(set(SERVICIOS_DE_ANKA))

    def test_una_solicitud_indirecta_sin_palabra_de_cotizacion_se_reconoce_por_el_servicio(self):
        self.assertTrue(menciona_un_servicio(fila("Necesitamos un vuelo con dron sobre la obra")))
        self.assertEqual(self.clase(fila("Necesitamos conocer el estado del terreno de la obra")).clase, REVISION_MANUAL)

    def test_un_mensaje_humano_sin_senales_no_se_descarta_va_a_work(self):
        for asunto in ("Hola", "¿Podemos hablar mañana?", "Felicitaciones por el cumpleaños", ""):
            with self.subTest(asunto=asunto):
                c = self.clase(fila(asunto))
                self.assertEqual((c.clase, c.motivo, c.enviar_a_work), (REVISION_MANUAL, "ambiguo_para_work", True))

    def test_un_remitente_humano_no_prueba_una_solicitud(self):
        c = self.clase(fila("Reunión de equipo", remitente="Jefe <jefe@anka.com>"))
        self.assertNotEqual(c.clase, SOLICITUD_NUEVA)

    def test_nunca_se_devuelve_una_clase_fuera_de_las_cuatro_y_ninguna_palabra_clave_descarta(self):
        for asunto in ("", "Hola", "Topografía", "Boletín"):
            for expediente in (False, True):
                c = self.clase(fila(asunto), expediente=expediente)
                self.assertIn(c.clase, CLASIFICACIONES)
                self.assertNotEqual(c.clase, NO_SOLICITUD)


class ServicioTests(BaseDeElegibilidad):
    def test_cada_tipo_de_mensaje_recibe_su_clasificacion_al_crear(self):
        self.fila("S1", "a", asunto="Cotización de levantamiento topográfico")
        self.fila("S2", "b", asunto="Hola, ¿podemos hablar?")
        self.fila("S3", "c", remitente="noreply@tm.openai.com", asunto="[Task Update] ANKA 01")
        self.fila("S4", "d", remitente="no-reply@portal.com", asunto="Su cotización")
        a, b = trabajos_work.crear_trabajo("S1", "a"), trabajos_work.crear_trabajo("S2", "b")
        self.assertEqual((a["clasificacion"], a["motivo_de_clasificacion"]), (SOLICITUD_NUEVA, "servicio_de_anka_mencionado"))
        self.assertEqual((b["clasificacion"], b["motivo_de_clasificacion"]), (REVISION_MANUAL, "ambiguo_para_work"))
        for hilo, mensaje in (("S3", "c"), ("S4", "d")):
            with self.assertRaises(NoElegible):
                trabajos_work.crear_trabajo(hilo, mensaje)
        self.assertEqual(self.trabajos(), 2)                   # nada de ruido ni de dudas de T-51 genera trabajo
        clases = {e["message_id"]: json.loads(e["detail_json"])["clase"] for e in self.eventos("clasificacion_preliminar")}
        self.assertEqual(clases, {"a": SOLICITUD_NUEVA, "b": REVISION_MANUAL, "c": NO_SOLICITUD, "d": REVISION_MANUAL})

    def test_el_evento_preliminar_es_unico_y_no_lleva_asunto(self):
        self.fila("S1", "a", asunto="ASUNTO CONFIDENCIAL de topografía")
        for _ in range(3):
            trabajos_work.crear_trabajo("S1", "a")
        evento, = self.eventos("clasificacion_preliminar")
        self.assertNotIn("CONFIDENCIAL", evento["detail_json"])
        self.assertEqual((evento["thread_id"], evento["message_id"]), ("S1", "a"))

    def test_un_hilo_con_expediente_pasa_a_seguimiento_y_deja_un_evento_nuevo(self):
        self.fila("S1", "a", asunto="Topografía")
        trabajos_work.crear_trabajo("S1", "a")
        with patch.object(work_import, "programar_actualizacion_de_hoja", MagicMock()):
            work_import.dashboard_upsert(b64(json.dumps(payload("S1", "a"))))
        self.fila("S1", "b", asunto="Re: Topografía")
        self.assertEqual(trabajos_work.crear_trabajo("S1", "b")["clasificacion"], SEGUIMIENTO)
        a = [json.loads(e["detail_json"])["clase"] for e in self.eventos("clasificacion_preliminar") if e["message_id"] == "a"]
        self.assertEqual(a, [SOLICITUD_NUEVA])
        with store.lectura() as db:
            primera = db.execute("SELECT * FROM recepcion_mensajes WHERE message_id='a'").fetchone()
            self.assertEqual(trabajos_work.clasificacion_del_mensaje(db, primera).clase, SEGUIMIENTO)

    def test_la_identidad_y_la_huella_de_t49_no_dependen_de_la_clasificacion(self):
        self.fila("S1", "a", asunto="Topografía")
        self.fila("S2", "b", asunto="Hola")
        for hilo, mensaje in (("S1", "a"), ("S2", "b")):
            t = trabajos_work.crear_trabajo(hilo, mensaje)
            self.assertEqual(t["job_id"], dominio.id_de_trabajo(CUENTA, hilo, mensaje, dominio.TIPO_ANALIZAR,
                                                                dominio.VERSION_DEL_CONTRATO))
            self.assertFalse(trabajos_work.crear_trabajo(hilo, mensaje)["creado"])
        self.assertNotIn("clasificacion", json.dumps(trabajos_work.obtener(t["job_id"])["referencia"]))

    def test_la_hoja_de_work_lleva_la_clasificacion_y_ningun_cuerpo(self):
        from openpyxl import load_workbook
        self.fila("S1", "a", asunto="Cotización batimetría")
        job = trabajos_work.crear_trabajo("S1", "a")["job_id"]
        hoja = load_workbook(BytesIO(trabajos_work.libro_del_trabajo(job)))["TRABAJO"]
        campos = {f[0].value: f[1].value for f in hoja.iter_rows(min_row=2)}
        self.assertEqual((campos["clasificacion_preliminar"], campos["motivo_preliminar"]),
                         (SOLICITUD_NUEVA, "servicio_de_anka_mencionado"))
        self.assertIn("classification", campos["instruccion_para_work"])
        self.assertNotIn("cuerpo", {str(k).lower() for k in campos})

    def test_el_trabajo_publicado_lleva_la_hoja_con_la_clasificacion(self):
        from openpyxl import load_workbook
        self.fila("S2", "b", asunto="Hola")
        job = trabajos_work.crear_trabajo("S2", "b")["job_id"]
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        archivo, = self.drive.archivos.values()
        hoja = load_workbook(BytesIO(archivo["bytes"]))["TRABAJO"]
        self.assertIn(REVISION_MANUAL, {f[1].value for f in hoja.iter_rows(min_row=2)})

    def test_el_listado_incluye_la_clasificacion(self):
        self.fila("S1", "a", asunto="Georadar")
        self.fila("S2", "b", remitente="noreply@tm.openai.com", asunto="x")
        lista = {f["message_id"]: f["clasificacion"] for f in trabajos_work.clasificar_pendientes()}
        self.assertEqual(lista, {"a": SOLICITUD_NUEVA, "b": NO_SOLICITUD})


class ImportacionTests(BaseDeElegibilidad):
    def setUp(self):
        super().setUp()
        parche = patch.object(work_import, "programar_actualizacion_de_hoja", MagicMock())
        parche.start()
        self.addCleanup(parche.stop)

    def importar(self, hilo, mensaje, **extra):
        work_import.dashboard_upsert(b64(json.dumps(payload(hilo, mensaje, draft_reply="Hola, esta es la respuesta.",
                                                            missing_data=["Falta el plano"], **extra))))

    def caso(self, hilo):
        return quotes.get_quote(hilo)

    def test_work_confirma_una_solicitud_nueva_o_un_seguimiento_y_sigue_su_camino(self):
        for hilo, clase in (("I1", SOLICITUD_NUEVA), ("I2", SEGUIMIENTO), ("I3", None)):
            with self.subTest(clase=clase):
                self.fila(hilo, "m")
                self.importar(hilo, "m", **({"classification": clase} if clase else {}))
                analisis = self.caso(hilo)["analysis"]
                self.assertEqual(len(analisis["activities"]), 1)
                self.assertTrue(analisis["is_quote_request"])
                self.assertEqual(analisis.get("classification"), clase)
                self.assertTrue(analisis["work_reply"]["text"] if "work_reply" in analisis else True)

    def test_no_solicitud_descarta_el_caso_nuevo_sin_actividades_ni_respuesta_y_cierra_la_cola(self):
        self.fila("I1", "m")
        self.importar("I1", "m", classification=NO_SOLICITUD, classification_reason="Es un saludo")
        caso = self.caso("I1")
        self.assertEqual(caso["status"], QuoteStatus.IGNORED.value)
        self.assertEqual((caso["analysis"]["activities"], caso["analysis"]["other_questions"],
                          caso["analysis"]["work_reply"]), ([], [], {}))
        self.assertFalse(caso["analysis"]["is_quote_request"])
        with store.lectura() as db:
            self.assertEqual(tuple(db.execute("SELECT state,disposition FROM recepcion_mensajes").fetchone()),
                             ("procesado", "interpretado_por_work"))

    def test_no_solicitud_en_un_hilo_con_expediente_no_lo_descarta_queda_en_revision(self):
        self.fila("I1", "m1")
        self.importar("I1", "m1", classification=SOLICITUD_NUEVA)
        self.fila("I1", "m2")
        self.importar("I1", "m2", classification=NO_SOLICITUD)
        self.assertEqual(self.caso("I1")["status"], QuoteStatus.REVIEW.value)

    def test_revision_manual_no_avanza_a_actividades_ni_a_respuesta(self):
        self.fila("I1", "m")
        self.importar("I1", "m", classification=REVISION_MANUAL)
        caso = self.caso("I1")
        self.assertEqual(caso["status"], QuoteStatus.REVIEW.value)
        self.assertEqual((caso["analysis"]["activities"], caso["analysis"]["other_questions"],
                          caso["analysis"]["work_reply"]), ([], [], {}))

    def test_la_clasificacion_de_work_queda_auditada_una_sola_vez_aunque_se_reimporte(self):
        self.fila("I1", "m")
        self.importar("I1", "m", classification=SOLICITUD_NUEVA)
        self.importar("I1", "m", classification=SOLICITUD_NUEVA, project_name="Otra versión")
        evento, = self.eventos("clasificacion_work")
        self.assertEqual(json.loads(evento["detail_json"]), {"clase": SOLICITUD_NUEVA})
        interpretado, = self.eventos("interpretado_por_work")
        self.assertEqual(json.loads(interpretado["detail_json"])["classification"], SOLICITUD_NUEVA)

    def test_un_valor_ilegible_va_a_revision_manual_y_nunca_avanza_como_si_no_hubiera_campo(self):
        for hilo, valor in (("I1", "TAL_VEZ"), ("I2", 7), ("I3", "")):
            with self.subTest(valor=valor):
                self.fila(hilo, "m")
                self.importar(hilo, "m", classification=valor)
                caso = self.caso(hilo)
                self.assertEqual(caso["analysis"]["classification"], REVISION_MANUAL)
                self.assertEqual(caso["analysis"]["activities"], [])
                self.assertEqual(caso["status"], QuoteStatus.REVIEW.value)
                self.assertTrue(any("cuatro permitidas" in a for a in caso["analysis"]["question_quality_warnings"]))

    def test_se_toleran_mayusculas_espacios_y_null(self):
        self.fila("I1", "m")
        self.importar("I1", "m", classification=" solicitud nueva ")
        self.assertEqual(self.caso("I1")["analysis"]["classification"], SOLICITUD_NUEVA)
        self.fila("I2", "m")
        self.importar("I2", "m", classification=None)
        self.assertNotIn("classification", self.caso("I2")["analysis"])
        self.assertEqual(len(self.caso("I2")["analysis"]["activities"]), 1)

    def test_el_contrato_sin_el_campo_sigue_siendo_valido(self):
        self.fila("I1", "m")
        self.importar("I1", "m")
        self.assertNotIn("classification", self.caso("I1")["analysis"])


if __name__ == "__main__":
    unittest.main()
