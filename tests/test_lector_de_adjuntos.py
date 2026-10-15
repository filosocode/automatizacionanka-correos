"""U-55: el lector de adjuntos lee lo que antes rechazaba y dice con claridad lo que no puede leer.

Todos los archivos son sintéticos: se construyen aquí mismo; no hay archivos de clientes.
"""

from tests import _isolate  # noqa: F401  (debe ir primero)
import base64
import io
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from anka.dominio.attachments import MARCA_FORMATO_ABIERTO, hermanos_de_shapefile
from anka.dominio.lectores_texto import Unreadable
from anka.infra import lectores_imagen, lectores_pdf, mail
from anka.infra.attachments import extract_attachment_text, extract_attachment_text_for_work_packet
from anka.infra.presupuesto_ocr import presupuesto_de_caso
from anka.servicios.work_attachment_packet import prepare_work_attachment_packet
from tests import _documentos as doc

LARGO = ("Área 544 ha del parque solar en Villavicencio, con levantamiento topográfico y LiDAR de toda "
         "la zona del proyecto en el Meta, con entregables en formato digital.")

KML_SIN_DECLARAR = """<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
<Document><name>Proyecto sintético</name><atom:author><atom:name>Anka</atom:name></atom:author>
<Placemark><name>Eje principal</name><description>Eje de prueba</description>
<LineString><coordinates>-74.0,4.0,0 -74.0,4.01,0</coordinates></LineString></Placemark>
<Placemark><name>Recorrido</name><gx:Track><when>2026-10-01T10:00:00Z</when><gx:coord>-74.1 4.1 0</gx:coord>
<when>2026-10-01T10:05:00Z</when><gx:coord>-74.1 4.11 0</gx:coord></gx:Track></Placemark>
<Placemark><name>Punto con atributo</name><ExtendedData><Data name="vereda"><value>La Esperanza</value></Data>
</ExtendedData><Point><coordinates>-74.2,4.2,0</coordinates></Point></Placemark>
</Document></kml>"""


def kml_declarado() -> str:
    return KML_SIN_DECLARAR.replace(
        '<kml xmlns="http://www.opengis.net/kml/2.2">',
        '<kml xmlns="http://www.opengis.net/kml/2.2" xmlns:atom="http://www.w3.org/2005/Atom" '
        'xmlns:gx="http://www.google.com/kml/ext/2.2">')


class KmlConPrefijosSinDeclararTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_a_kml_with_undeclared_prefixes_opens_and_keeps_every_geometry_and_attribute(self):
        texto, motivo = extract_attachment_text("mapa.kml", "", KML_SIN_DECLARAR.encode())
        referencia, _ = extract_attachment_text("mapa.kml", "", kml_declarado().encode())
        self.assertIn("Elementos geográficos: 3", texto)
        for esperado in ("Eje principal", "Recorrido", "Punto con atributo", "vereda=La Esperanza",
                         "Marcas de tiempo KML: 2", "-74.1"):
            self.assertIn(esperado, texto)
        self.assertIn("Lectura tolerante: prefijos XML sin declarar (atom, gx)", texto)
        # sin esa línea, el texto es idéntico al del mismo archivo bien declarado
        self.assertEqual("\n".join(linea for linea in texto.splitlines() if "Lectura tolerante" not in linea),
                         referencia)
        self.assertEqual(motivo, "")

    def test_a_kmz_with_undeclared_prefixes_opens_too(self):
        texto, motivo = extract_attachment_text("mapa.kmz", "", doc.zip_de({"doc.kml": KML_SIN_DECLARAR}))
        self.assertIn("Elementos geográficos: 3", texto)
        self.assertEqual(motivo, "")

    def test_a_declared_kml_is_not_touched(self):
        texto, _ = extract_attachment_text("mapa.kml", "", kml_declarado().encode())
        self.assertNotIn("Lectura tolerante", texto)

    def test_a_truly_broken_kml_still_fails_with_a_clear_reason(self):
        texto, motivo = extract_attachment_text("roto.kml", "", b"<kml><Document><Placemark></kml>")
        self.assertEqual(texto, "")
        self.assertIn("no se pudo leer", motivo)

    def test_entities_are_never_resolved_or_expanded(self):
        with tempfile.TemporaryDirectory() as carpeta:
            secreto = Path(carpeta) / "secreto.txt"
            secreto.write_text("CONTENIDO-PRIVADO", encoding="utf-8")
            externa = (f'<?xml version="1.0"?><!DOCTYPE kml [<!ENTITY x SYSTEM "{secreto.as_uri()}">]>'
                       '<kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>&x;</name></Document></kml>')
            interna = ('<?xml version="1.0"?><!DOCTYPE kml [<!ENTITY a "AAAA"><!ENTITY b "&a;&a;&a;&a;">]>'
                       '<kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>&b;</name></Document></kml>')
            for contenido in (externa, interna, externa.encode("utf-16")):
                texto, motivo = extract_attachment_text(
                    "ent.kml", "", contenido if isinstance(contenido, bytes) else contenido.encode())
                self.assertEqual(texto, "")
                self.assertIn("entidades", motivo)
                self.assertNotIn("CONTENIDO-PRIVADO", motivo)


def shapefile_en_bytes(puntos=3, con_dbf=True, con_prj=True):
    import shapefile

    shp, shx, dbf = io.BytesIO(), io.BytesIO(), io.BytesIO()
    with shapefile.Writer(shp=shp, shx=shx, dbf=dbf) as escritor:
        escritor.field("nombre", "C")
        escritor.field("area_ha", "N", decimal=1)
        for n in range(puntos):
            escritor.point(-74 + n / 100, 4 + n / 100)
            escritor.record(f"punto{n}", 10.5 + n)
    archivos = {".shp": shp.getvalue(), ".shx": shx.getvalue()}
    if con_dbf:
        archivos[".dbf"] = dbf.getvalue()
    if con_prj:
        archivos[".prj"] = b'GEOGCS["WGS 84"]'
    return archivos


class ShapefileSueltoTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_a_lone_shp_gives_geometry_extent_and_record_count_and_is_partial(self):
        archivos = shapefile_en_bytes()
        texto, motivo = extract_attachment_text("capa.shp", "", archivos[".shp"])
        self.assertIn("Capa shapefile: capa; geometría: Point; elementos: 3", texto)
        self.assertIn("Extensión XY (mín X, mín Y, máx X, máx Y): -74.0, 4.0, -73.98, 4.02", texto)
        self.assertIn("falta el .dbf", texto)
        self.assertIn("falta el .shx", texto)
        self.assertIn("parcial", motivo)
        self.assertIn("sin .prj", motivo)

    def test_with_its_dbf_the_fields_and_records_are_read(self):
        archivos = shapefile_en_bytes()
        texto, motivo = extract_attachment_text("capa.shp", "", archivos[".shp"], {".dbf": archivos[".dbf"]})
        self.assertIn("campos: nombre, area_ha", texto)
        self.assertIn("Registro 1: nombre=punto0; area_ha=10.5", texto)
        self.assertNotIn("falta el .dbf", texto)
        self.assertIn("parcial", motivo)        # aún faltan el .shx y el .prj

    def test_with_every_sibling_the_reading_is_complete_in_both_modes(self):
        archivos = shapefile_en_bytes()
        companeros = {extension: contenido for extension, contenido in archivos.items() if extension != ".shp"}
        for leer in (extract_attachment_text, extract_attachment_text_for_work_packet):
            texto, motivo = leer("capa.shp", "", archivos[".shp"], companeros)
            self.assertIn("Sistema de coordenadas (PRJ): GEOGCS", texto)
            self.assertEqual(motivo, "")

    def test_a_dbf_alone_lists_its_fields_and_is_partial(self):
        texto, motivo = extract_attachment_text("capa.dbf", "", shapefile_en_bytes()[".dbf"])
        self.assertIn("Tabla dBase capa; registros: 3; campos: nombre, area_ha", texto)
        self.assertIn("parcial", motivo)

    def test_a_dbf_with_a_different_record_count_is_flagged(self):
        archivos = shapefile_en_bytes(puntos=3)
        otro = shapefile_en_bytes(puntos=5)
        texto, _ = extract_attachment_text("capa.shp", "", archivos[".shp"], {".dbf": otro[".dbf"]})
        self.assertIn("el .dbf declara 5 registros y el .shp tiene 3", texto)

    def test_a_truncated_shp_reports_what_it_could_count(self):
        archivos = shapefile_en_bytes(puntos=3)
        texto, motivo = extract_attachment_text("capa.shp", "", archivos[".shp"][:-10])
        self.assertIn("elementos: 2", texto)
        self.assertIn("termina incompleto", texto)
        self.assertIn("parcial", motivo)

    def test_a_file_that_is_not_a_shapefile_says_so(self):
        texto, motivo = extract_attachment_text("falso.shp", "", b"x" * 200)
        self.assertEqual(texto, "")
        self.assertIn("no tiene la cabecera de un shapefile", motivo)

    def test_a_zip_with_an_incomplete_shapefile_reads_it_and_flags_what_is_missing(self):
        archivos = shapefile_en_bytes()
        contenido = doc.zip_de({"capa.shp": archivos[".shp"], "capa.prj": archivos[".prj"], "nota.txt": "hola"})
        texto, motivo = extract_attachment_text("paquete.zip", "", contenido)
        self.assertIn("Capa shapefile: capa; geometría: Point; elementos: 3", texto)
        self.assertIn("falta el .dbf", texto)
        self.assertIn("parcial", motivo)

    def test_a_shapefile_and_its_siblings_are_one_entry(self):
        nombres = ["Capa.SHP", "capa.dbf", "CAPA.shx", "capa.prj", "otra.dbf", "informe.pdf", "capa.xml"]
        self.assertEqual(hermanos_de_shapefile(nombres),
                         {"Capa.SHP": {".dbf": "capa.dbf", ".shx": "CAPA.shx", ".prj": "capa.prj"}})

    def test_in_a_message_the_siblings_are_read_with_the_shp_and_not_listed_apart(self):
        archivos = shapefile_en_bytes()
        partes = [{"filename": f"capa{extension}", "mimeType": "application/octet-stream",
                   "body": {"data": base64.urlsafe_b64encode(contenido).decode(), "size": len(contenido)}}
                  for extension, contenido in archivos.items()]
        partes.append({"filename": "informe.txt", "mimeType": "text/plain",
                       "body": {"data": base64.urlsafe_b64encode(b"hola").decode(), "size": 4}})
        resultados = mail.extract_message_attachments(None, {"id": "m1", "payload": {"parts": partes}})
        por_nombre = {r["filename"]: r for r in resultados}
        self.assertEqual(por_nombre["capa.shp"]["coverage"], "completa")
        self.assertIn("campos: nombre, area_ha", por_nombre["capa.shp"]["text"])
        for hermano in ("capa.shx", "capa.dbf", "capa.prj"):
            self.assertEqual(por_nombre[hermano]["status"], "ignorado")
            self.assertIn("parte del shapefile capa.shp", por_nombre[hermano]["reason"])
        self.assertEqual(por_nombre["informe.txt"]["coverage"], "completa")

    def test_in_a_work_packet_the_siblings_are_one_source_file(self):
        archivos = shapefile_en_bytes()
        with tempfile.TemporaryDirectory() as carpeta:
            rutas = []
            for extension, contenido in archivos.items():
                ruta = Path(carpeta) / f"capa{extension}"
                ruta.write_bytes(contenido)
                rutas.append(str(ruta))
            resultado = prepare_work_attachment_packet(rutas, str(Path(carpeta) / "paquete"))
            self.assertEqual(resultado["source_files"], 1)
            self.assertEqual(resultado["extracted_files"], 1)
            import json
            manifiesto = json.loads((Path(carpeta) / "paquete" / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual([a["filename"] for a in manifiesto["attachments"]], ["capa.shp"])
            self.assertEqual(manifiesto["attachments"][0]["coverage"], "completa")


class ShapefileEnElExpedienteDeDriveTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_in_the_drive_dossier_the_siblings_are_not_records_of_their_own(self):
        from unittest.mock import MagicMock
        from anka.infra.drive_exchange import DriveExchange
        from anka.servicios import work_drive as work
        from tests.fake_work_drive import MemoryDrive

        def parte(nombre, contenido):
            return {"filename": nombre, "mimeType": "application/octet-stream",
                    "body": {"data": base64.urlsafe_b64encode(contenido).decode(), "size": len(contenido)}}

        archivos = shapefile_en_bytes()
        mensaje = {"id": "m1", "internalDate": "1", "labelIds": ["INBOX"], "payload": {
            "headers": [{"name": "From", "value": "cliente@example.com"}, {"name": "Subject", "value": "Capa"}],
            "parts": [{"mimeType": "text/plain", "body": {"data": base64.urlsafe_b64encode(b"Correo").decode()}},
                      *[parte(f"capa{extension}", contenido) for extension, contenido in archivos.items()]]}}
        drive = MemoryDrive()
        work.prepare_thread(MagicMock(), DriveExchange(drive), "reception", {"id": "t1", "messages": [mensaje]},
                            {"anka@example.com"})
        carpeta = DriveExchange(drive).find("reception", "thread:t1")
        manifiesto = DriveExchange(drive).read_json(DriveExchange(drive).find(carpeta["id"], "current"))
        estados = {r["filename"]: r["status"] for r in manifiesto["attachments"]}
        self.assertEqual(estados, {"capa.shp": "extraido", "capa.shx": "ignorado", "capa.dbf": "ignorado",
                                   "capa.prj": "ignorado"})
        self.assertEqual([r["coverage"] for r in manifiesto["attachments"] if r["status"] != "ignorado"],
                         ["completa"])
        self.assertFalse(manifiesto["review_required"])


class PdfLargoTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_a_long_pdf_is_read_in_part_and_says_which_pages(self):
        contenido = doc.pdf_con_texto([f"{LARGO} Marca numero {n}" for n in range(1, 8)])
        with patch.object(lectores_pdf, "MAX_PDF_PAGINAS_LEIDAS", 5), \
                patch.object(lectores_pdf, "ocr_pil_image") as ocr:
            texto, motivo = extract_attachment_text("largo.pdf", "application/pdf", contenido)
        self.assertIn("Marca numero 5", texto)
        self.assertNotIn("Marca numero 6", texto)
        self.assertIn("COBERTURA PARCIAL: PDF de 7 páginas; se leyeron las páginas 1–5", texto)
        self.assertIn("faltan las páginas 6–7", texto)
        self.assertIn("parcial", motivo)
        ocr.assert_not_called()

    def test_ocr_is_capped_per_file_and_the_rest_is_listed(self):
        contenido = doc.pdf_con_texto(["x"] * 4)
        with patch.object(lectores_pdf, "MAX_OCR_PAGES", 2), \
                patch.object(lectores_pdf, "ocr_pil_image", return_value="texto del OCR") as ocr:
            texto, _ = extract_attachment_text("p.pdf", "application/pdf", contenido)
        self.assertEqual(ocr.call_count, 2)
        self.assertIn("OCR limitado a 2 páginas por archivo; páginas sin OCR: 3, 4", texto)


def motor_lento(segundos=0.06):
    def motor(_imagen):
        time.sleep(segundos)
        return [([0], "texto reconocido", 0.9)], None
    return patch.object(lectores_imagen, "ocr_engine", return_value=motor)


class PresupuestoDeOcrTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_by_default_every_page_gets_ocr_as_always(self):
        contenido = doc.pdf_con_texto(["x"] * 3)
        with motor_lento(0.01), presupuesto_de_caso({}):
            texto, motivo = extract_attachment_text("p.pdf", "application/pdf", contenido)
        self.assertEqual(texto.count("texto reconocido"), 3)
        self.assertNotIn("omitido", texto)
        self.assertNotIn("omitido", motivo)

    def test_a_per_file_budget_skips_the_ocr_that_does_not_fit_and_says_which_pages(self):
        contenido = doc.pdf_con_texto(["x"] * 4)
        with motor_lento(), presupuesto_de_caso({"ocr_segundos_por_archivo": 0.05}):
            texto, motivo = extract_attachment_text("p.pdf", "application/pdf", contenido)
        self.assertEqual(texto.count("texto reconocido"), 1)
        self.assertIn("COBERTURA PARCIAL: OCR omitido por tiempo (presupuesto 0.05 s por archivo); "
                      "páginas sin OCR: 2–4", texto)
        self.assertIn("OCR omitido por tiempo", motivo)

    def test_a_per_case_budget_is_shared_by_all_the_files_of_the_case(self):
        contenido = doc.pdf_con_texto([""])
        with motor_lento(), presupuesto_de_caso({"ocr_segundos_por_caso": 0.05}):
            primero, _ = extract_attachment_text("a.pdf", "application/pdf", contenido)
            segundo, motivo = extract_attachment_text("b.pdf", "application/pdf", contenido)
        self.assertIn("texto reconocido", primero)
        self.assertEqual(segundo, "")                         # nada leído: el caso ya gastó su tiempo
        self.assertIn("OCR omitido por tiempo", motivo)

    def test_the_file_budget_restarts_for_each_file(self):
        contenido = doc.pdf_con_texto(["x"])
        with motor_lento(), presupuesto_de_caso({"ocr_segundos_por_archivo": 0.05}):
            for nombre in ("a.pdf", "b.pdf"):
                texto, _ = extract_attachment_text(nombre, "application/pdf", contenido)
                self.assertIn("texto reconocido", texto)

    def test_the_policy_marks_a_scanned_pdf_as_needing_vision_without_running_ocr(self):
        contenido = doc.pdf_con_texto([""] * 3)
        with patch.object(lectores_pdf, "ocr_pil_image") as ocr, presupuesto_de_caso({"ocr_omitir_escaneados": True}):
            texto, motivo = extract_attachment_text("escaneo.pdf", "application/pdf", contenido)
        ocr.assert_not_called()
        self.assertEqual(texto, "")
        self.assertIn("requiere visión", motivo)

    def test_the_policy_keeps_the_native_text_of_a_mixed_pdf_and_lists_the_scanned_pages(self):
        contenido = doc.pdf_con_texto([LARGO] * 3 + [""] * 4)
        with patch.object(lectores_pdf, "ocr_pil_image") as ocr, presupuesto_de_caso({"ocr_omitir_escaneados": True}):
            texto, motivo = extract_attachment_text("mixto.pdf", "application/pdf", contenido)
        ocr.assert_not_called()
        self.assertIn("544 ha del parque", texto)
        self.assertIn("REQUIERE VISIÓN: 4 de 7 páginas parecen escaneadas", texto)
        self.assertIn("páginas sin OCR: 4–7", texto)
        self.assertIn("requiere visión", motivo)

    def test_the_policy_does_not_apply_to_a_pdf_that_is_mostly_text(self):
        contenido = doc.pdf_con_texto([LARGO] * 3 + [""])
        with patch.object(lectores_pdf, "ocr_pil_image", side_effect=Unreadable("sin texto")) as ocr, \
                presupuesto_de_caso({"ocr_omitir_escaneados": True}):
            texto, _ = extract_attachment_text("casi.pdf", "application/pdf", contenido)
        self.assertEqual(ocr.call_count, 1)
        self.assertNotIn("REQUIERE VISIÓN", texto)

    def test_the_policy_marks_a_loose_image_as_needing_vision(self):
        with motor_lento(0.0) as motor, presupuesto_de_caso({"ocr_omitir_escaneados": True}):
            texto, motivo = extract_attachment_text("foto.png", "image/png", doc.PNG_1X1)
        motor.assert_not_called()
        self.assertEqual(texto, "")
        self.assertIn("requiere visión", motivo)

    def test_the_defaults_in_config_keep_the_old_behaviour(self):
        from anka.infra.presupuesto_ocr import presupuesto_desde_config
        presupuesto = presupuesto_desde_config({})
        self.assertEqual((presupuesto.por_archivo, presupuesto.por_caso, presupuesto.omitir_escaneados),
                         (0.0, 0.0, False))
        self.assertFalse(presupuesto.agotado())

    def test_the_config_keys_are_validated(self):
        from anka.infra.settings import validate_config
        base = {"required_fields": [], "llm": {"provider": "ollama", "base_url": "http://x"}}
        self.assertEqual(validate_config({**base, "ocr_segundos_por_archivo": 30, "ocr_segundos_por_caso": 0,
                                          "ocr_omitir_escaneados": True}), [])
        problemas = validate_config({**base, "ocr_segundos_por_archivo": -1, "ocr_segundos_por_caso": "x",
                                     "ocr_omitir_escaneados": "si"})
        self.assertEqual(len(problemas), 3)


class ExcelTests(unittest.TestCase):
    """Cubre: U-55"""

    def libro(self):
        from openpyxl import Workbook
        libro = Workbook()
        visible = libro.active
        visible.title = "Datos"
        visible.append(["Área", 12])
        visible["B2"] = "=B1*2"
        oculta = libro.create_sheet("Interna")
        oculta.sheet_state = "hidden"
        oculta.append(["precio base", "reservado"])
        muy_oculta = libro.create_sheet("Secreta")
        muy_oculta.sheet_state = "veryHidden"
        muy_oculta.append(["nota", "solo para el equipo"])
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()

    def test_hidden_and_very_hidden_sheets_are_read_and_marked(self):
        texto, motivo = extract_attachment_text("libro.xlsx", "", self.libro())
        self.assertIn("Hoja: Interna; visibilidad: hidden; HOJA OCULTA;", texto)
        self.assertIn("Hoja: Secreta; visibilidad: veryHidden; HOJA OCULTA;", texto)
        self.assertIn("solo para el equipo", texto)
        self.assertNotIn("Hoja: Datos; visibilidad: visible; HOJA OCULTA", texto)

    def test_formulas_are_shown_with_their_saved_value_and_never_recalculated(self):
        texto, _ = extract_attachment_text("libro.xlsx", "", self.libro())
        self.assertIn("B2==B1*2 [valor guardado: None]", texto)
        self.assertIn("FÓRMULA SIN VALOR CALCULADO: hoja Datos, celda B2", texto)


class SinLectorTests(unittest.TestCase):
    """Cubre: U-55"""

    def test_rar_and_7z_are_not_readable_and_ask_for_an_open_format(self):
        for nombre, contenido, nombre_formato in (("planos.rar", b"Rar!\x1a\x07\x00", "RAR"),
                                                  ("planos.7z", b"7z\xbc\xaf\x27\x1c", "7z")):
            texto, motivo = extract_attachment_text(nombre, "application/octet-stream", contenido)
            self.assertEqual(texto, "")
            self.assertIn(nombre_formato, motivo)
            self.assertIn(MARCA_FORMATO_ABIERTO, motivo)
            self.assertEqual(extract_attachment_text_for_work_packet(nombre, "", contenido), (texto, motivo))

    def test_dwg_says_its_version_and_asks_for_an_open_format(self):
        texto, motivo = extract_attachment_text("plano.dwg", "", b"AC1032" + b"\x00" * 50)
        self.assertEqual(texto, "")
        self.assertIn("plano DWG (AutoCAD 2018, AC1032)", motivo)
        self.assertIn(MARCA_FORMATO_ABIERTO, motivo)
        _, otro = extract_attachment_text("plano.dwg", "", b"zzzzzz")
        self.assertIn(MARCA_FORMATO_ABIERTO, otro)

    DXF = ("0\nSECTION\n2\nHEADER\n9\n$ACADVER\n1\nAC1027\n9\n$INSUNITS\n70\n6\n0\nENDSEC\n"
           "0\nSECTION\n2\nTABLES\n0\nTABLE\n2\nLAYER\n0\nLAYER\n2\nEJES\n70\n0\n0\nLAYER\n2\nTEXTOS\n70\n0\n"
           "0\nENDTAB\n0\nENDSEC\n0\nSECTION\n2\nENTITIES\n0\nLINE\n8\nEJES\n0\nLINE\n8\nEJES\n"
           "0\nTEXT\n8\nTEXTOS\n1\nPortal 1\n0\nMTEXT\n8\nTEXTOS\n1\n{\\fArial|b0;Cota 12 m}\\PNorte\n"
           "0\nENDSEC\n0\nEOF\n")

    def test_a_text_dxf_gives_version_layers_entities_and_texts_and_is_partial(self):
        texto, motivo = extract_attachment_text("plano.dxf", "", self.DXF.encode())
        self.assertIn("Archivo DXF: plano.dxf; versión: AC1027 (AutoCAD 2013); unidades (código 6)", texto)
        self.assertIn("Capas (2): EJES, TEXTOS", texto)
        self.assertIn("Entidades por tipo: LINE=2, TEXT=1, MTEXT=1", texto)
        self.assertIn("Texto 1: Portal 1", texto)
        self.assertIn("Texto 2: Cota 12 m Norte", texto)
        self.assertIn("parcial", motivo)

    def test_a_binary_dxf_or_a_fake_dxf_is_not_readable_and_asks_for_an_open_format(self):
        for contenido in (b"AutoCAD Binary DXF\r\n\x1a\x00" + b"\x00" * 20, b"esto no es un dxf"):
            texto, motivo = extract_attachment_text("plano.dxf", "", contenido)
            self.assertEqual(texto, "")
            self.assertIn(MARCA_FORMATO_ABIERTO, motivo)

    def test_inside_a_zip_the_unreadable_formats_are_listed_with_the_mark(self):
        contenido = doc.zip_de({"a.rar": b"Rar!", "b.dwg": b"AC1015", "nota.txt": "hola"})
        texto, motivo = extract_attachment_text("paquete.zip", "", contenido)
        self.assertIn("Sin lector local: a.rar, b.dwg", texto)
        self.assertIn(MARCA_FORMATO_ABIERTO, texto)
        self.assertIn("parcial", motivo)


if __name__ == "__main__":
    unittest.main()
