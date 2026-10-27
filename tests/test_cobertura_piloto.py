"""Herramienta de cobertura del piloto: cuántas cifras de los adjuntos llegan a la tabla de Work (PLAN_PILOTO_LECTURA §5)."""

from tests import _isolate  # noqa: F401  (debe ir primero)

import contextlib
import hashlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from anka.dominio.areas_calculadas import AVISO_DE_DATO
from anka.dominio.status import QuoteStatus
from anka.infra import instantaneas_de_manifiesto, store
from laboratorio.piloto import cobertura
from tests.test_app import complete_analysis

CORREO = "Necesitamos levantar un predio de 41,5 ha. El campo dura 12 días."
PDF = ("Garantía de 6 días. " + "Texto de relleno del documento. " * 5
       + "El área del lote es de 13,8 ha. Secciones cada 50 metros.")
KMZ = (f"Polígono «A»: área geodésica aproximada calculada 415000.0 m² (41.5000 ha); perímetro 9.0 m; "
       f"anillos interiores: 0. {AVISO_DE_DATO}\nÁrea total calculada de 2 polígonos: 500000.0 m² (50.0000 ha). {AVISO_DE_DATO}")


def _adjuntos() -> list[dict]:
    return [{"filename": "tdr.pdf", "text": PDF, "coverage": "completa", "source_id": "x"},
            {"filename": "lote.kmz", "text": KMZ, "coverage": "parcial", "source_id": "y"},
            {"filename": "plano.pdf", "text": "", "coverage": "no_legible", "source_id": "z"}]


def _analisis(**extra) -> dict:
    return {**complete_analysis(), "activities": [{"activity": "Levantamiento", "quantity": "41,5", "unit": "ha"}],
            **extra}


# Cubre: T-132
class CifrasDelHiloTests(unittest.TestCase):
    def test_toma_las_cifras_del_correo_y_de_cada_adjunto_sin_repetir(self) -> None:
        cifras = cobertura.cifras_del_hilo({"correo": CORREO + " Otra vez 41,5 ha.", "adjuntos": _adjuntos()})
        self.assertEqual([(c["archivo"], c["cifra"]) for c in cifras],
                         [("correo", "41,5 ha"), ("correo", "12 días"), ("tdr.pdf", "13,8 ha"), ("tdr.pdf", "50 m")])

    def test_no_cuenta_lo_que_calculo_el_programa_de_un_kml(self) -> None:
        cifras = cobertura.cifras_del_hilo({"correo": "", "adjuntos": [_adjuntos()[1]]})
        self.assertEqual(cifras, [])

    def test_un_adjunto_sin_texto_o_mal_formado_no_rompe(self) -> None:
        self.assertEqual(cobertura.cifras_del_hilo({"adjuntos": [_adjuntos()[2], None, "x"]}), [])   # type: ignore[list-item]
        self.assertEqual(cobertura.cifras_del_hilo({}), [])


# Cubre: T-132
class MedirHiloTests(unittest.TestCase):
    TEXTOS = {"correo": CORREO, "adjuntos": _adjuntos()}

    def test_cuenta_las_cifras_que_la_tabla_no_menciona_y_marca_el_detalle(self) -> None:
        m = cobertura.medir_hilo(_analisis(), self.TEXTOS)
        self.assertEqual((m["cifras"], m["cifras_que_faltan"]), (4, 3))
        self.assertAlmostEqual(m["cobertura"], 0.25)
        self.assertEqual([(c["cifra"], c["en_tabla"]) for c in m["detalle"]],
                         [("41,5 ha", True), ("12 días", False), ("13,8 ha", False), ("50 m", False)])

    def test_la_cobertura_sube_cuando_la_tabla_menciona_mas(self) -> None:
        tabla = [{"activity": "Levantamiento", "quantity": "41,5", "notes": "campo 12 días, lote de 13,8 ha, cada 50 m"}]
        m = cobertura.medir_hilo(_analisis(activities=tabla), self.TEXTOS)
        self.assertEqual(m["cifras_que_faltan"], 0)
        self.assertEqual(m["cobertura"], 1.0)

    def test_sin_cifras_la_cobertura_no_se_inventa(self) -> None:
        self.assertIsNone(cobertura.medir_hilo(_analisis(), {"correo": "Hola", "adjuntos": []})["cobertura"])

    def test_cuenta_adjuntos_y_toma_la_fidelidad_de_t38(self) -> None:
        evidencia = {"comprobadas": 7, "diferencias": [{"x": 1}], "no_comprobables": [1, 2], "estado": "con_diferencias"}
        m = cobertura.medir_hilo(_analisis(evidence_check=evidencia), self.TEXTOS)
        self.assertEqual((m["adjuntos"], m["adjuntos_no_legibles"], m["adjuntos_parciales"]), (3, 1, 1))
        self.assertEqual((m["comprobadas"], m["diferencias"], m["no_comprobables"], m["estado_t38"]),
                         (7, 1, 2, "con_diferencias"))
        self.assertEqual(cobertura.medir_hilo(_analisis(), self.TEXTOS)["estado_t38"], "sin comprobar")


class ConBase(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.datos = Path(self.temp.name)
        db = store.open_db(self.datos)
        store.save(db, "t-ok", "m1", QuoteStatus.READY_REVIEW, "a@b.co", _analisis())
        store.save(db, "t-sin-textos", "m2", QuoteStatus.READY_REVIEW, "a@b.co", _analisis())
        store.save(db, "t-sin-tabla", "m3", QuoteStatus.READY_REVIEW, "a@b.co", {**_analisis(), "activities": []})
        store.save(db, "t-incompleto", "m4", QuoteStatus.READY_REVIEW, "a@b.co", _analisis())
        db.commit()
        db.close()
        directorio = self.datos / "manifiestos"
        instantaneas_de_manifiesto.guardar_textos(directorio, "t-ok", "r1", CORREO, _adjuntos())
        instantaneas_de_manifiesto.guardar_textos(directorio, "t-sin-tabla", "r1", CORREO, _adjuntos())
        instantaneas_de_manifiesto.guardar_textos(
            directorio, "t-incompleto", "r1", "x" * (instantaneas_de_manifiesto.MAX_TEXTO_LOCAL + 1), _adjuntos())

    def hash_de_la_base(self) -> str:
        return hashlib.sha256((self.datos / "quotes.sqlite3").read_bytes()).hexdigest()


# Cubre: T-132
class MedirLaBaseTests(ConBase):
    def medir(self, *hilos: str):
        return cobertura.medir(cobertura.copia_de_la_base(self.datos / "quotes.sqlite3"),
                               self.datos / "manifiestos", hilos)

    def test_mide_los_hilos_con_tabla_y_textos_y_dice_por_que_omite_los_demas(self) -> None:
        mediciones, omitidos = self.medir()
        self.assertEqual(list(mediciones), ["t-ok"])
        self.assertEqual(omitidos, {"t-sin-textos": "sin textos guardados del hilo", "t-sin-tabla": "sin tabla de actividades",
                                    "t-incompleto": "textos incompletos (un adjunto no se pudo bajar)"})

    def test_se_pueden_pedir_hilos_concretos_y_uno_que_no_existe(self) -> None:
        mediciones, omitidos = self.medir("t-ok", "t-fantasma")
        self.assertEqual((list(mediciones), omitidos), (["t-ok"], {"t-fantasma": "no está en la base"}))

    def test_leer_la_base_no_la_modifica(self) -> None:
        antes = self.hash_de_la_base()
        self.medir()
        self.assertEqual(self.hash_de_la_base(), antes)

    def test_la_copia_es_en_memoria_y_no_deja_escribir_en_la_original(self) -> None:
        copia = cobertura.copia_de_la_base(self.datos / "quotes.sqlite3")
        copia.execute("DELETE FROM quotes")
        copia.commit()
        otra = cobertura.copia_de_la_base(self.datos / "quotes.sqlite3")
        self.assertEqual(otra.execute("SELECT COUNT(*) FROM quotes").fetchone()[0], 4)

    def test_el_resumen_solo_trae_conteos_nunca_texto_del_cliente(self) -> None:
        mediciones, omitidos = self.medir()
        salida = "\n".join(cobertura.resumen(mediciones, omitidos))
        self.assertIn("Hilos medidos: 1 | omitidos: 3", salida)
        self.assertIn("cobertura global: 25%", salida)
        for frase in ("predio", "lote", "Secciones", "t-ok"):
            self.assertNotIn(frase, salida)


# Cubre: T-132
class HojaDelPilotoTests(ConBase):
    def hoja(self) -> Path:
        mediciones, _ = cobertura.medir(cobertura.copia_de_la_base(self.datos / "quotes.sqlite3"),
                                        self.datos / "manifiestos")
        ruta = self.datos / "salida" / "cobertura_piloto.xlsx"
        cobertura.escribir_hoja(ruta, mediciones)
        return ruta

    def test_la_hoja_trae_resumen_y_detalle_con_columna_para_clasificar(self) -> None:
        libro = load_workbook(self.hoja())
        self.assertEqual(libro.sheetnames, ["Resumen", "Cifras"])
        detalle = list(libro["Cifras"].iter_rows(values_only=True))
        self.assertEqual(detalle[0][-1], "Clasificación (persona)")
        self.assertEqual([f[4] for f in detalle[1:]], ["sí", "no", "no", "no"])
        self.assertEqual(list(libro["Resumen"].iter_rows(min_row=2, values_only=True))[0][:5], ("t-ok", 1, 4, 3, 0.25))

    def test_cuenta_lo_que_una_persona_clasifico_solo_entre_las_que_faltan(self) -> None:
        ruta = self.hoja()
        libro = load_workbook(ruta)
        hoja = libro["Cifras"]
        for fila, clase in ((2, "real"), (3, "real"), (4, "contexto"), (5, "ruido")):   # la fila 2 está en la tabla
            hoja.cell(row=fila, column=6, value=clase)
        libro.save(ruta)
        self.assertEqual(cobertura.clasificadas(ruta), {"real": 1, "contexto": 1, "ruido": 1})


# Cubre: T-132
class PrincipalTests(ConBase):
    def correr(self, *args: str) -> tuple[int, str, str]:
        salida, error = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
            codigo = cobertura.principal(["--datos", str(self.datos), "--salida", str(self.datos / "salida"), *args])
        return codigo, salida.getvalue(), error.getvalue()

    def test_mide_imprime_conteos_y_deja_la_hoja(self) -> None:
        codigo, salida, _ = self.correr()
        self.assertEqual(codigo, 0)
        self.assertIn("Hilos medidos: 1", salida)
        self.assertTrue((self.datos / "salida" / "cobertura_piloto.xlsx").is_file())

    def test_sin_hilos_medibles_no_escribe_hoja(self) -> None:
        codigo, salida, _ = self.correr("t-sin-tabla")
        self.assertEqual(codigo, 0)
        self.assertIn("Hilos medidos: 0", salida)
        self.assertFalse((self.datos / "salida").exists())

    def test_sin_base_dice_que_falta_y_termina_con_error(self) -> None:
        with tempfile.TemporaryDirectory() as vacio:
            salida, error = io.StringIO(), io.StringIO()
            with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(error):
                codigo = cobertura.principal(["--datos", vacio])
        self.assertEqual(codigo, 1)
        self.assertIn("No encuentro la base", error.getvalue())

    def test_cuenta_la_clasificacion_de_una_hoja_llena(self) -> None:
        self.correr()
        ruta = self.datos / "salida" / "cobertura_piloto.xlsx"
        libro = load_workbook(ruta)
        libro["Cifras"].cell(row=3, column=6, value="real")
        libro["Cifras"].cell(row=4, column=6, value="ruido")
        libro.save(ruta)
        codigo, salida, _ = self.correr("--clasificadas", str(ruta))
        self.assertEqual(codigo, 0)
        self.assertIn("real: 1, contexto: 0, ruido: 1", salida)
        self.assertIn("50%", salida)


if __name__ == "__main__":
    unittest.main()
