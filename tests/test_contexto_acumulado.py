"""T-59 Fase C: el contexto acumulado de un SEGUIMIENTO llega a Work en la hoja del trabajo.

SQLite temporal, Gmail falso y Drive simulado. No se publica nada real ni se crea ninguna tarea de Work.

Cubre: T-59
"""

from tests import _isolate  # noqa: F401  (debe ir primero)

import copy
import json
import re
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from anka.cli import paquete_work
from anka.dominio import contexto_previo, preguntas
from anka.dominio.status import QuoteStatus
from anka.dominio.trabajo_work import HOJAS
from anka.dominio.work_contract import validate
from anka.infra import store
from anka.servicios import quotes, trabajos_work
from tests.test_continuidad_de_casos import Contexto, fila
from tests.test_elegibilidad_de_trabajos import CARPETA
from tests.test_elegibilidad_de_trabajos import Base as BaseDeElegibilidad

ENVIO = "2026-09-02T09:00:00-05:00"
RESPUESTA = "2026-09-03T09:00:00-05:00"
PAQUETE = "abc123def456"
LEVANTAMIENTO = fila("Levantamiento topográfico", "10", "ha", "confirmado", entregable="Plano")
BATIMETRIA = fila("Batimetría", servicio="Batimetría", ubicacion="Embalse")


def hojas_de(contenido):
    libro = load_workbook(BytesIO(contenido))
    return {h.title: [[c.value for c in f] for f in h.iter_rows()] for h in libro.worksheets}


def campos(hoja):
    return {f[0]: f[1] for f in hoja[1:]}


class Base(BaseDeElegibilidad):
    def expediente(self, hilo, mensaje="a", **analisis):
        base = {"project_name": "Predio El Roble", "client_name": "Ana", "received_at": "2026-09-01T09:00:00-05:00",
                "activities": [copy.deepcopy(LEVANTAMIENTO), copy.deepcopy(BATIMETRIA)], **analisis}
        with store.transaccion() as db:
            store.save(db, hilo, mensaje, QuoteStatus.REVIEW, "cliente@example.com", base)

    def libro(self, hilo, mensaje, paquete=""):
        self.fila(hilo, mensaje, asunto="Re: Topografía")
        trabajo = trabajos_work.crear_trabajo(hilo, mensaje, paquete_version=paquete)
        return trabajo["job_id"], hojas_de(trabajos_work.libro_del_trabajo(trabajo["job_id"]))


# Cubre: T-59
class ContextoEnLaHoja(Base):
    def test_un_seguimiento_trae_el_contexto_del_expediente_en_tres_hojas(self):
        self.expediente("S1")
        _, hojas = self.libro("S1", "b")
        self.assertEqual(list(hojas), [*HOJAS, *contexto_previo.HOJAS_DE_CONTEXTO])
        previo = campos(hojas["CONTEXTO_PREVIO"])
        self.assertEqual((previo["expediente_thread_id"], previo["mensaje_por_analizar"],
                          previo["ultimo_mensaje_analizado"], previo["mensaje_ya_analizado"]), ("S1", "b", "a", "no"))
        self.assertEqual((previo["actividades_vigentes"], previo["tabla_editada_por_una_persona"],
                          previo["propuesta_pendiente_de_revision"]), ("2", "no", "no"))
        actividades = hojas["ACTIVIDADES_VIGENTES"]
        self.assertEqual(actividades[0], list(contexto_previo.COLUMNAS_DE_ACTIVIDADES))
        self.assertEqual([f[0] for f in actividades[1:]], ["Levantamiento topográfico", "Batimetría"])
        self.assertEqual(actividades[1][3:5], ["10", "ha"])
        self.assertIn("instruccion_de_seguimiento", campos(hojas["TRABAJO"]))
        self.assertIn("provisional", previo["aviso"])

    def test_una_solicitud_nueva_conserva_el_formato_anterior_sin_contexto_ficticio(self):
        self.fila("S9", "m", asunto="Cotización de batimetría")
        job = trabajos_work.crear_trabajo("S9", "m")["job_id"]
        hojas = hojas_de(trabajos_work.libro_del_trabajo(job))
        self.assertEqual(list(hojas), list(HOJAS))
        self.assertNotIn("instruccion_de_seguimiento", campos(hojas["TRABAJO"]))

    def test_la_hoja_publicada_en_drive_lleva_el_contexto(self):
        self.expediente("S1")
        job, _ = self.libro("S1", "b")
        trabajos_work.publicar_trabajo(self.drive, job, CARPETA)
        archivo, = self.drive.archivos.values()
        self.assertIn("ACTIVIDADES_VIGENTES", hojas_de(archivo["bytes"]))

    def test_construir_la_hoja_no_escribe_en_el_expediente_y_es_idempotente(self):
        self.expediente("S1")
        job, primera = self.libro("S1", "b")
        with store.lectura() as db:
            antes = [tuple(r) for r in db.execute("SELECT thread_id, analysis_json, status FROM quotes")]
            eventos = db.execute("SELECT COUNT(*) FROM case_events").fetchone()[0]
        segunda = hojas_de(trabajos_work.libro_del_trabajo(job))
        with store.lectura() as db:
            self.assertEqual([tuple(r) for r in db.execute("SELECT thread_id, analysis_json, status FROM quotes")], antes)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM case_events").fetchone()[0], eventos)
        self.assertEqual(campos(primera["CONTEXTO_PREVIO"])["huella_del_contexto"],
                         campos(segunda["CONTEXTO_PREVIO"])["huella_del_contexto"])

    def test_el_mensaje_ya_analizado_se_dice(self):
        self.expediente("S1", mensaje="b")
        _, hojas = self.libro("S1", "b")
        self.assertEqual(campos(hojas["CONTEXTO_PREVIO"])["mensaje_ya_analizado"], "si")


# Cubre: T-59
class SinMezclaNiSecretos(Base):
    def test_el_contexto_es_solo_del_expediente_del_hilo(self):
        self.expediente("S1")
        self.expediente("S2", activities=[fila("SECRETO_DE_OTRO_EXPEDIENTE", "1", "u", "confirmado")],
                        project_name="Otro proyecto")
        _, hojas = self.libro("S1", "b")
        todo = json.dumps(hojas, ensure_ascii=False)
        self.assertNotIn("SECRETO_DE_OTRO_EXPEDIENTE", todo)
        self.assertNotIn("Otro proyecto", todo)
        self.assertNotIn("S2", json.dumps([hojas[h] for h in contexto_previo.HOJAS_DE_CONTEXTO]))

    def test_no_incluye_correos_tokens_ni_rutas(self):
        self.expediente("S1", client_email="privado@example.com", token="ya29.SECRETO", ruta=r"C:\Users\x")
        _, hojas = self.libro("S1", "b")
        for nombre in contexto_previo.HOJAS_DE_CONTEXTO:
            texto = json.dumps(hojas[nombre], ensure_ascii=False)
            for prohibido in ("@", "ya29", "SECRETO", "C:\\", "token"):
                self.assertNotIn(prohibido, texto.replace("Estado provisional", ""))

    def test_los_textos_largos_y_las_listas_se_recortan_con_aviso(self):
        muchas = [fila(f"Actividad {n}", str(n), "u", "confirmado", notes="x" * 900) for n in range(70)]
        self.expediente("S1", activities=muchas)
        _, hojas = self.libro("S1", "b")
        self.assertEqual(len(hojas["ACTIVIDADES_VIGENTES"]) - 1, contexto_previo.MAX_FILAS)
        self.assertTrue(all(len(str(f[9])) <= contexto_previo.MAX_TEXTO for f in hojas["ACTIVIDADES_VIGENTES"][1:]))
        self.assertEqual(campos(hojas["CONTEXTO_PREVIO"])["contexto_recortado"], "si")


# Cubre: T-59
class PreguntasEnElContexto(Base):
    def preguntas(self, hojas):
        return [dict(zip(hojas["PREGUNTAS"][0], f, strict=True)) for f in hojas["PREGUNTAS"][1:]]

    def test_enviada_propuesta_e_interna_se_distinguen(self):
        analisis = {"question_audit": [
            {"question": "¿Cuál es el datum?", "disposition": "ask_client"},
            {"question": "¿Plazo de entrega?", "disposition": "ask_client"},
            {"question": "Revisar la contradicción de áreas", "disposition": "internal_review"},
            {"question": "¿Nombre legal?", "disposition": "do_not_ask"}]}
        preguntas.anotar_envio(analisis, ENVIO, "s-1", ["¿Cuál es el datum?"], False)
        self.expediente("S1", **analisis)
        _, hojas = self.libro("S1", "b")
        tipos = {p["pregunta"]: p["tipo"] for p in self.preguntas(hojas)}
        self.assertEqual(tipos, {"¿Cuál es el datum?": "enviada", "¿Plazo de entrega?": "propuesta_no_enviada",
                                 "Revisar la contradicción de áreas": "interna"})
        enviada = next(p for p in self.preguntas(hojas) if p["tipo"] == "enviada")
        self.assertEqual((enviada["ronda"], enviada["enviada_el"], enviada["mensaje_gmail"], enviada["estado"],
                          enviada["respondida_en_mensaje"]), ("1", ENVIO, "s-1", "enviada", None))

    def test_una_pregunta_solo_propuesta_nunca_figura_como_enviada(self):
        self.expediente("S1", question_audit=[{"question": "¿Cuál es el datum?", "disposition": "ask_client"}])
        _, hojas = self.libro("S1", "b")
        self.assertEqual({p["tipo"] for p in self.preguntas(hojas)}, {"propuesta_no_enviada"})

    def test_la_respuesta_ligada_y_la_pregunta_sin_respuesta_demostrada(self):
        analisis = {"question_audit": [{"question": "¿Plazo de entrega?", "disposition": "ask_client"}]}
        preguntas.anotar_envio(analisis, ENVIO, "s-1", ["¿Cuál es el datum?", "¿Plazo de entrega?"], False)
        preguntas.conservar_rondas(analisis, analisis, "m-2")          # sin llegada posterior: no se liga nada
        analisis["received_at"] = RESPUESTA
        preguntas.conservar_rondas(analisis, analisis, "m-2")
        self.expediente("S1", **analisis)
        _, hojas = self.libro("S1", "b")
        por_texto = {p["pregunta"]: p for p in self.preguntas(hojas)}
        self.assertEqual(por_texto["¿Cuál es el datum?"]["estado"], "sin_confirmar")
        self.assertEqual(por_texto["¿Plazo de entrega?"]["estado"], "sigue_abierta")
        self.assertEqual({p["respondida_en_mensaje"] for p in por_texto.values()}, {"m-2"})
        self.assertNotIn("respondida", {p["estado"] for p in por_texto.values()})

    def test_un_expediente_anterior_a_las_rondas_tambien_se_lista(self):
        self.expediente("S1", preguntas_enviadas={"en": ENVIO, "preguntas": ["¿Cuál es el datum?"], "esperada": False})
        _, hojas = self.libro("S1", "b")
        self.assertEqual([(p["tipo"], p["mensaje_gmail"]) for p in self.preguntas(hojas)], [("enviada", None)])


# Cubre: T-59
class RevisionHumanaEnElContexto(Base):
    def test_tabla_editada_propuesta_pendiente_y_motivos(self):
        self.expediente("S1", manual_edit={"at": ENVIO, "by": "persona"},
                        pending_model_activities=[copy.deepcopy(LEVANTAMIENTO)],
                        revision_de_tabla=[{"tipo": "dato_distinto", "actividad": "Levantamiento topográfico",
                                            "campo": "quantity", "anterior": "10", "propuesto": "12", "fuente": "c2"}])
        _, hojas = self.libro("S1", "b")
        previo = campos(hojas["CONTEXTO_PREVIO"])
        self.assertEqual((previo["tabla_editada_por_una_persona"], previo["propuesta_pendiente_de_revision"]),
                         ("si", "si"))
        self.assertEqual(previo["motivos_de_revision"], "dato_distinto: Levantamiento topográfico")

    def test_las_contradicciones_abiertas_se_cuentan(self):
        self.expediente("S1", activities=[fila("A", "1", "u", "contradictorio"), fila("B", "1", "u", "confirmado")],
                        contradictions=[{"parameter": "área"}])
        _, hojas = self.libro("S1", "b")
        self.assertEqual(campos(hojas["CONTEXTO_PREVIO"])["contradicciones_abiertas"], "2")


# Cubre: T-59
class VersionDelPaquete(Base):
    def test_un_trabajo_nuevo_lleva_la_version_del_paquete_en_el_evento_y_en_la_hoja(self):
        self.fila("S1", "a", asunto="Topografía")
        trabajo = trabajos_work.crear_trabajo("S1", "a", paquete_version=PAQUETE)
        evento = self.eventos("trabajo_creado")[0]
        self.assertEqual(json.loads(evento["detail_json"])["package_version"], PAQUETE)
        hojas = hojas_de(trabajos_work.libro_del_trabajo(trabajo["job_id"]))
        self.assertEqual(campos(hojas["TRABAJO"])["version_del_paquete"], PAQUETE)

    def test_un_trabajo_anterior_no_cambia_su_hoja_ni_su_identidad(self):
        self.fila("S1", "a", asunto="Topografía")
        antiguo = trabajos_work.crear_trabajo("S1", "a")
        nuevo = trabajos_work.crear_trabajo("S1", "a", paquete_version=PAQUETE)         # el mismo trabajo: no se rehace
        self.assertFalse(nuevo["creado"])
        self.assertEqual(antiguo["job_id"], nuevo["job_id"])
        self.assertNotIn("package_version", self.eventos("trabajo_creado")[0]["detail_json"])
        hojas = hojas_de(trabajos_work.libro_del_trabajo(antiguo["job_id"]))
        self.assertNotIn("version_del_paquete", campos(hojas["TRABAJO"]))

    def test_la_version_calculada_es_la_del_manifiesto(self):
        import tempfile
        with tempfile.TemporaryDirectory() as temp:
            manifiesto = paquete_work.construir(Path(temp))
        self.assertEqual(paquete_work.version_del_paquete(), manifiesto["paquete_version"])
        self.assertRegex(manifiesto["paquete_version"], r"^[0-9a-f]{12}$")

    def test_el_despacho_pasa_la_version_al_crear_el_trabajo(self):
        from anka.servicios import despacho
        self.fila("S1", "a", asunto="Topografía")
        with patch.object(trabajos_work, "publicar_trabajo", MagicMock(return_value={"resultado": "creado"})):
            despacho.despachar_mensaje(self.drive, despacho.Alcance("S1", "a"),
                                       {despacho.NOMBRE_PENDIENTES: CARPETA}, paquete=PAQUETE)
        self.assertEqual(json.loads(self.eventos("trabajo_creado")[0]["detail_json"])["package_version"], PAQUETE)


# Cubre: T-59, T-52
class ElPrompt(unittest.TestCase):
    PROMPT = Path(paquete_work.PROMPT).read_text(encoding="utf-8")

    def test_trae_las_reglas_del_seguimiento(self):
        texto = self.PROMPT
        for frase in ("CONTEXTO_PREVIO, ACTIVIDADES_VIGENTES y PREGUNTAS", "estado provisional", "ACUMULADA",
                      "activity exacto", "no la elimina", "ya no se requiere", "contradictorio", "nunca lo sobrescribas",
                      "«enviada»", "propuesta_no_enviada", "sigue_abierta", "sin_confirmar", "do_not_ask",
                      "no autoriza nada comercial"):
            self.assertIn(frase, texto)

    def test_la_version_y_el_cierre_son_los_nuevos(self):
        self.assertIn("PRECISO_V6.1.3", self.PROMPT)
        self.assertNotIn("PRECISO_V6.1.1", self.PROMPT)
        self.assertTrue(self.PROMPT.rstrip().endswith("FIN_PROMPT_ANKA_V6.1.3"))

    def test_no_se_pierde_la_regla_de_la_revision_interna_sin_bloqueo(self):
        # Cubre: T-57
        self.assertIn("REVISIÓN INTERNA SIN BLOQUEO", self.PROMPT)
        self.assertIn("datos ficticios o de prueba, laboratorio", self.PROMPT)
        self.assertIn("NO vacían la salida", self.PROMPT)

    def test_el_esquema_publicado_no_cambia_y_el_ejemplo_de_seguimiento_lo_cumple(self):
        esquema = json.loads(paquete_work.ESQUEMA.read_text(encoding="utf-8"))
        self.assertEqual(esquema["required"],
                         ["thread_id", "message_id", "client_email", "project_name", "status", "engine"])
        ejemplo = paquete_work.ejemplos(self.PROMPT)["SEGUIMIENTO"]
        self.assertEqual(validate(ejemplo), [])
        self.assertEqual(ejemplo["skill_version"], "PRECISO_V6.1.3")

    def test_el_paquete_nuevo_se_reconstruye_y_se_verifica(self):
        import tempfile
        with tempfile.TemporaryDirectory() as temp:
            carpeta = Path(temp)
            manifiesto = paquete_work.construir(carpeta)
            self.assertEqual(paquete_work.verificar_paquete(carpeta), [])
            self.assertEqual(manifiesto["prompt"], {"file": "PRECISO_V6.1.3.txt", "version": "PRECISO_V6.1.3",
                                                    "sha256": manifiesto["prompt"]["sha256"]})
            self.assertIn("Seguimiento con contexto", (carpeta / "T-52_CLASIFICACION.md").read_text(encoding="utf-8"))
            self.assertNotIn("{{", (carpeta / "00_LEEME.md").read_text(encoding="utf-8"))
            self.assertTrue(re.search(manifiesto["paquete_version"], (carpeta / "00_LEEME.md").read_text("utf-8")))


# Cubre: T-59
class RetornosDeWorkConLaTablaAcumulada(Contexto):
    """Lo que Work devolvería con el prompt nuevo, importado con el importador de siempre."""

    def test_una_tabla_completa_e_igual_no_cambia_nada(self):
        antes = self.tabla()
        versiones = self.versiones()
        despues = self.seguimiento(copy.deepcopy(antes))
        self.assertEqual(despues["activities"], antes)
        self.assertEqual((self.versiones(), self.eventos_de("table_merge")), (versiones, []))

    def test_una_tabla_completa_con_una_magnitud_cambiada_pide_revision_y_no_pierde_nada(self):
        completa = [fila("Levantamiento topográfico", "12", "ha", "confirmado", entregable="Plano", fuente="correo 2",
                         notes="Cambio: 10 ha → 12 ha"), copy.deepcopy(BATIMETRIA)]
        despues = self.seguimiento(completa)
        self.assertEqual([a["quantity"] for a in despues["activities"]], ["10", ""])
        self.assertEqual(despues["revision_de_tabla"][0]["tipo"], "dato_distinto")
        self.assertEqual(len(despues["pending_model_activities"]), 2)

    def test_una_tabla_solo_con_lo_modificado_equivale_a_la_acumulada(self):
        parcial = self.seguimiento([fila("Batimetría", "5", "km", "confirmado", servicio="Batimetría",
                                         ubicacion="Embalse", fuente="correo 2")])
        self.assertEqual([a["activity"] for a in parcial["activities"]], ["Levantamiento topográfico", "Batimetría"])
        self.assertEqual(parcial["activities"][1]["quantity"], "5")

    def test_un_nombre_distinto_no_se_fusiona_sino_que_se_revisa(self):
        despues = self.seguimiento([fila("Levantamiento topográfico y altimétrico", "10", "ha", "confirmado",
                                         entregable="Plano")])
        self.assertEqual(despues["revision_de_tabla"][0]["tipo"], "sustitucion_ambigua")
        self.assertEqual(len(despues["activities"]), 2)

    def test_la_eliminacion_con_la_nota_del_prompt_se_detecta_y_no_se_aplica_sola(self):
        despues = self.seguimiento([fila("Batimetría", servicio="Batimetría", ubicacion="Embalse",
                                         notes="El cliente indica que ya no se requiere", fuente="correo 2")])
        self.assertEqual(despues["revision_de_tabla"][0]["tipo"], "eliminacion_pedida")
        self.assertEqual(len(despues["activities"]), 2)

    def test_una_contradiccion_queda_en_revision(self):
        despues = self.seguimiento([fila("Levantamiento topográfico", "10", "ha", "contradictorio", entregable="Plano",
                                         notes="El correo dice 14 ha")])
        self.assertEqual(despues["activities"][0]["estado"], "contradictorio")
        self.assertEqual(despues["revision_de_tabla"][0]["tipo"], "contradiccion")

    def test_la_revision_interna_con_datos_de_laboratorio_no_vacia_la_tabla(self):
        # Cubre: T-57
        despues = self.seguimiento([fila("Batimetría", "5", "km", "pendiente", servicio="Batimetría",
                                         ubicacion="Embalse")], status="Requiere revisión técnica interna",
                                   analysis_complete=False)
        self.assertEqual(len(despues["activities"]), 2)

    def test_la_edicion_humana_pendiente_no_se_pisa(self):
        with store.transaccion() as db:
            analisis = json.loads(db.execute("SELECT analysis_json FROM quotes WHERE thread_id='t-1'").fetchone()[0])
            analisis["manual_edit"] = {"at": ENVIO, "by": "persona"}
            analisis["activities"][1]["quantity"] = "9"
            db.execute("UPDATE quotes SET analysis_json=? WHERE thread_id='t-1'", (json.dumps(analisis),))
        mia = self.tabla()
        despues = self.seguimiento([fila("Batimetría", "5", "km", "confirmado", servicio="Batimetría",
                                         ubicacion="Embalse")])
        self.assertEqual(despues["activities"], mia)

    def test_una_importacion_historica_de_v6_1_1_sigue_entrando(self):
        # Cubre: T-57
        despues = self.seguimiento([], skill_version="PRECISO_V6.1.1")
        self.assertEqual(despues["work_skill_version"], "PRECISO_V6.1.1")
        self.assertEqual(len(despues["activities"]), 2)

    def test_el_contexto_que_vio_work_corresponde_a_lo_importado(self):
        # el contexto de la hoja se arma del mismo análisis que luego se integra: mismas actividades vigentes
        contexto = contexto_previo.contexto_de_seguimiento(self.analisis(), thread_id="t-1", estado="x", codigo_t="",
                                                           ultimo_mensaje="m-1", mensaje_nuevo="m-2")
        self.assertEqual([f[0] for f in contexto.actividades], [a["activity"] for a in self.tabla()])
        self.assertEqual(quotes.get_quote("t-1")["analysis"]["activities"], self.tabla())


if __name__ == "__main__":
    unittest.main()
