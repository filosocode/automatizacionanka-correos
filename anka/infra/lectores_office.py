"""Lectores de Word, Excel y PowerPoint (OOXML y .xls) y de OpenDocument, con localizadores citables."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from anka.dominio.attachments import MAX_SPREADSHEET_CELLS
from anka.dominio.lectores_geograficos import local_name
from anka.dominio.lectores_texto import Unreadable, plain_text
from anka.infra.lectores_zip import append_archive_image_ocr, zip_acotado


def xls_text(content: bytes) -> str:
    """Texto de un libro .xls antiguo, con fila y columna de cada celda; se corta en MAX_SPREADSHEET_CELLS."""
    try:
        import xlrd
    except ImportError:
        raise Unreadable("falta instalar xlrd para leer libros .xls") from None
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise Unreadable(f"no se pudo abrir el libro .xls: {exc}") from None
    lines = []
    cells_seen = 0
    for sheet in workbook.sheets():
        lines.append(f"Hoja: {sheet.name}; filas: {sheet.nrows}; columnas: {sheet.ncols}")
        for row_number in range(sheet.nrows):
            values = []
            for column in range(sheet.ncols):
                cells_seen += 1
                if cells_seen > MAX_SPREADSHEET_CELLS:
                    lines.append(f"LECTURA PARCIAL: se superó el límite de {MAX_SPREADSHEET_CELLS} celdas")
                    workbook.release_resources()
                    return "\n".join(lines)
                value = sheet.cell_value(row_number, column)
                if value != "":
                    values.append(f"fila {row_number + 1}, columna {column + 1}={value}")
            if values:
                lines.append(" | ".join(values))
    workbook.release_resources()
    return "\n".join(lines)


def pptx_text(content: bytes) -> str:
    """Texto de una presentación PowerPoint: diapositivas, notas, SmartArt y OCR de imágenes incrustadas."""
    with zip_acotado(content) as presentation:
        slide_names = sorted(
            (name for name in presentation.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(Path(name).stem.removeprefix("slide")),
        )
        if not slide_names:
            raise Unreadable("presentación sin diapositivas legibles")
        lines = []
        for slide_number, name in enumerate(slide_names, 1):
            root = ElementTree.fromstring(presentation.read(name))
            paragraphs = []
            for node in root.iter():
                if local_name(node.tag) == "p":
                    text = "".join(child.text or "" for child in node.iter() if local_name(child.tag) == "t").strip()
                    if text:
                        paragraphs.append(text)
            lines.append(f"[Diapositiva {slide_number}]\n" + "\n".join(paragraphs))
        images = [name for name in presentation.namelist() if name.startswith("ppt/media/")]
        append_archive_image_ocr(presentation, images, lines)
        notes = [name for name in presentation.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)]
        for note_name in notes:
            note_root = ElementTree.fromstring(presentation.read(note_name))
            note_text = [
                "".join(node.itertext()).strip()
                for node in note_root.iter()
                if local_name(node.tag) == "t" and "".join(node.itertext()).strip()
            ]
            if note_text:
                lines.append(f"Notas del presentador ({note_name}): " + " | ".join(note_text))
        diagrams = [
            name for name in presentation.namelist() if name.startswith("ppt/diagrams/") and name.endswith(".xml")
        ]
        for diagram_name in diagrams:
            diagram_root = ElementTree.fromstring(presentation.read(diagram_name))
            diagram_text = [
                "".join(node.itertext()).strip()
                for node in diagram_root.iter()
                if local_name(node.tag) == "t" and "".join(node.itertext()).strip()
            ]
            if diagram_text:
                lines.append(f"COBERTURA VISUAL PARCIAL: texto SmartArt ({diagram_name}): " + " | ".join(diagram_text))
        charts = [name for name in presentation.namelist() if name.startswith("ppt/charts/")]
        if charts:
            lines.append("CONTENIDO DE GRÁFICOS NO EXTRAÍDO: " + ", ".join(charts))
        return "\n".join(lines)


def odt_text(content: bytes) -> str:
    """Texto de un documento OpenDocument de texto (.odt) y OCR de sus imágenes."""
    with zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for node in root.iter():
        if local_name(node.tag) in {"p", "h"}:
            text = " ".join("".join(node.itertext()).split())
            if text:
                lines.append(text)
    with zip_acotado(content) as document:
        images = [name for name in document.namelist() if name.startswith("Pictures/")]
        append_archive_image_ocr(document, images, lines)
    return plain_text("\n".join(lines).encode("utf-8"))


def ods_text(content: bytes) -> str:
    """Texto de una hoja OpenDocument (.ods) con fila, columna y fórmula de cada celda."""
    with zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for table_number, table in enumerate((n for n in root.iter() if local_name(n.tag) == "table"), 1):
        name = next(
            (value for key, value in table.attrib.items() if local_name(key) == "name"),
            f"Hoja {table_number}",
        )
        lines.append(f"Hoja: {name}")
        for row_number, row in enumerate((n for n in table.iter() if local_name(n.tag) == "table-row"), 1):
            values = []
            column_number = 0
            for cell in (n for n in row if local_name(n.tag) in {"table-cell", "covered-table-cell"}):
                repeat = next((int(v) for k, v in cell.attrib.items() if local_name(k) == "number-columns-repeated"), 1)
                value = " ".join(
                    "".join(p.itertext()).strip() for p in cell.iter() if local_name(p.tag) in {"p", "h"}
                ).strip()
                formula = next((v for k, v in cell.attrib.items() if local_name(k) == "formula"), "")
                if formula:
                    if not value:
                        lines.append(
                            f"FÓRMULA SIN VALOR CALCULADO: hoja {name}, fila {row_number}, columna {column_number + 1}"
                        )
                    value = f"{value} [fórmula: {formula}]".strip()
                for _ in range(min(repeat, 1000)):
                    column_number += 1
                    if value:
                        values.append(f"fila {row_number}, columna {column_number}={value}")
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def odp_text(content: bytes) -> str:
    """Texto de una presentación OpenDocument (.odp) y OCR de sus imágenes."""
    with zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    pages = [node for node in root.iter() if local_name(node.tag) == "page"]
    for number, page in enumerate(pages, 1):
        values = []
        for node in page.iter():
            if local_name(node.tag) in {"p", "h"}:
                value = " ".join("".join(node.itertext()).split())
                if value:
                    values.append(value)
        lines.append(f"[Diapositiva {number}]\n" + "\n".join(values))
    with zip_acotado(content) as document:
        images = [name for name in document.namelist() if name.startswith(("Pictures/", "ObjectReplacements/"))]
        append_archive_image_ocr(document, images, lines)
    return "\n".join(lines)


def _docx_paragraph_text(paragraph: ElementTree.Element) -> str:
    return "".join(
        "\t" if local_name(node.tag) == "tab" else "\n" if local_name(node.tag) in {"br", "cr"} else node.text or ""
        for node in paragraph.iter()
        if local_name(node.tag) in {"t", "tab", "br", "cr"}
    ).strip()


def docx_text(content: bytes) -> str:
    """Texto de un documento Word: párrafos y tablas numerados, encabezados, notas y cambios rastreados."""
    with zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("word/document.xml"))
    body = next((node for node in root.iter() if local_name(node.tag) == "body"), root)
    lines = []
    paragraph_number = table_number = 0
    for block in body:
        kind = local_name(block.tag)
        if kind == "p":
            paragraph_number += 1
            text = _docx_paragraph_text(block)
            if text:
                lines.append(f"Párrafo {paragraph_number}: {text}")
        elif kind == "tbl":
            table_number += 1
            _docx_table(block, table_number, lines)
    with zip_acotado(content) as document:
        names = document.namelist()
        for name in names:
            if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                header_root = ElementTree.fromstring(document.read(name))
                text = " ".join(
                    "".join(node.itertext()).strip() for node in header_root.iter() if local_name(node.tag) == "p"
                ).strip()
                if text:
                    kind = "Encabezado" if "/header" in name else "Pie de página"
                    lines.append(f"{kind} {name}: {text}")
        for name in names:
            if name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
                sidecar_root = ElementTree.fromstring(document.read(name))
                note_text = [
                    " ".join("".join(node.itertext()).split())
                    for node in sidecar_root.iter()
                    if local_name(node.tag) in {"footnote", "endnote", "comment"}
                ]
                if note_text:
                    lines.append(f"Contenido asociado ({name}): " + " | ".join(note_text))
        revisions = [node for node in root.iter() if local_name(node.tag) in {"ins", "del"}]
        if revisions:
            lines.append(
                "CAMBIOS RASTREADOS: comparar inserciones y eliminaciones antes de considerar definitivo el texto"
            )
        for revision in revisions:
            value = " ".join("".join(revision.itertext()).split())
            if value:
                lines.append(f"Revisión rastreada ({local_name(revision.tag)}): {value}")
        images = [name for name in names if name.startswith("word/media/")]
        append_archive_image_ocr(document, images, lines)
    return "\n".join(lines)


def _docx_table(block: ElementTree.Element, table_number: int, lines: list[str]) -> None:
    for row_number, row in enumerate((node for node in block if local_name(node.tag) == "tr"), 1):
        cells = []
        for cell_number, cell in enumerate((node for node in row if local_name(node.tag) == "tc"), 1):
            value = " ".join(_docx_paragraph_text(node) for node in cell.iter() if local_name(node.tag) == "p").strip()
            cells.append(f"columna {cell_number}={value}")
        if any(value.split("=", 1)[-1] for value in cells):
            lines.append(f"Tabla {table_number}, fila {row_number}: " + " | ".join(cells))


def xlsx_text(content: bytes) -> str:
    """Texto de un libro Excel con coordenada de cada celda, fórmulas, comentarios y avisos de cobertura parcial."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise Unreadable("falta instalar openpyxl para leer Excel") from None
    zip_acotado(content).close()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
    rows = []
    for sheet in workbook.worksheets:
        formula_sheet = formulas[sheet.title]
        oculta = "HOJA OCULTA; " if sheet.sheet_state != "visible" else ""
        rows.append(
            f"Hoja: {sheet.title}; visibilidad: {sheet.sheet_state}; {oculta}"
            f"rango declarado: {sheet.calculate_dimension()}"
        )
        cell_count = 0
        over_limit = False
        formula_rows = formula_sheet.iter_rows()
        for row_number, row in enumerate(sheet.iter_rows(), 1):
            if row_number > MAX_SPREADSHEET_CELLS:
                rows.append(f"LECTURA PARCIAL: hoja {sheet.title} supera {MAX_SPREADSHEET_CELLS} filas inspeccionadas")
                break
            formula_row = next(formula_rows, ())
            values = []
            for cell, formula_cell in zip(row, formula_row, strict=True):
                cell_count += 1
                if cell_count > MAX_SPREADSHEET_CELLS:
                    rows.append(
                        f"LECTURA PARCIAL: hoja {sheet.title} supera {MAX_SPREADSHEET_CELLS} celdas inspeccionadas"
                    )
                    over_limit = True
                    break
                value = _xlsx_value(cell, formula_cell, sheet, rows)
                if value is not None:
                    values.append(f"{cell.coordinate}={value}")
            if values:
                rows.append(" | ".join(values))
            if over_limit:
                break
    with zip_acotado(content) as archive:
        images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        append_archive_image_ocr(archive, images, rows)
        charts = [name for name in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
        for chart_name in charts:
            chart_root = ElementTree.fromstring(archive.read(chart_name))
            chart_text = [
                " ".join("".join(node.itertext()).split())
                for node in chart_root.iter()
                if local_name(node.tag) in {"f", "v"} and "".join(node.itertext()).strip()
            ]
            rows.append(f"Gráfico ({chart_name}); referencias y valores almacenados: " + "; ".join(chart_text[:80]))
        comments = [name for name in archive.namelist() if name.startswith("xl/comments")]
        for comment_name in comments:
            comment_root = ElementTree.fromstring(archive.read(comment_name))
            for comment in (node for node in comment_root.iter() if local_name(node.tag) == "comment"):
                reference = comment.attrib.get("ref", "celda sin referencia")
                author_id = comment.attrib.get("authorId", "")
                value = " ".join("".join(comment.itertext()).split())
                rows.append(f"Comentario ({comment_name}, {reference}, autor {author_id}): {value}")
        pivots = [name for name in archive.namelist() if name.startswith("xl/pivotTables/")]
        if pivots:
            rows.append(
                "COBERTURA PARCIAL: tablas dinámicas requieren interpretación especializada: " + ", ".join(pivots)
            )
        drawings = [
            name for name in archive.namelist() if name.startswith("xl/drawings/") and name.endswith(".xml")
        ]
        if drawings:
            rows.append("DIBUJOS/FORMAS DE HOJA PENDIENTES DE LECTURA VISUAL: " + ", ".join(drawings))
    return "\n".join(rows)


def _xlsx_value(cell: Any, formula_cell: Any, sheet: Any, rows: list[str]) -> str | None:
    if formula_cell.data_type == "f":
        value: str | None = f"{formula_cell.value} [valor guardado: {cell.value}]"
        if cell.value is None:
            rows.append(f"FÓRMULA SIN VALOR CALCULADO: hoja {sheet.title}, celda {cell.coordinate}")
    else:
        value = str(cell.value) if cell.value is not None else None
    return value
