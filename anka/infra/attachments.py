"""Lectores de documentos y archivos comprimidos y coordinación de extracción."""

from __future__ import annotations
from typing import Any
import posixpath
from anka.dominio.attachments import (
    attachment_reviewed_sections as attachment_reviewed_sections,
    build_attachment_evidence as build_attachment_evidence,
)
import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from anka.dominio.attachments import (
    MAX_ATTACHMENT_BYTES,
    MAX_ATTACHMENT_TEXT,
    MAX_WORK_PACKET_ATTACHMENT_BYTES,
    MAX_EVIDENCE_ITEMS,
    MAX_ARCHIVE_FILES,
    MAX_ARCHIVE_EXPANDED_BYTES,
    MAX_SPREADSHEET_CELLS,
    MAX_OCR_PAGES,
    MAX_OCR_IMAGES,
    MAX_OCR_PIXELS,
    TEXT_ATTACHMENT_EXTENSIONS,
)
from anka.dominio.lectores_texto import Unreadable, plain_text, html_text, rtf_text
from anka.dominio.lectores_geograficos import (
    local_name,
    kmz_text,
    kml_text_file,
    wpml_text,
    gpx_text,
    qgc_waypoints_text,
    qgc_plan_text,
    geojson_text,
    csv_text,
)
from anka.infra.lectores_imagen import ocr_image, ocr_pil_image, image_text, office_legacy_text


def _zip_acotado(content: bytes) -> zipfile.ZipFile:
    """Abre un ZIP (o un documento Office/OpenDocument, que lo es) rechazando bombas de descompresión.

    Los tamaños se leen del directorio central sin descomprimir nada; `zipfile` no lee más de lo declarado,
    así que el tope vale aunque la cabecera mienta. Sin esto, un adjunto de 10 MB podría expandirse a gigabytes.
    """
    archive = zipfile.ZipFile(BytesIO(content))
    entries = [item for item in archive.infolist() if not item.is_dir()]
    if len(entries) > MAX_ARCHIVE_FILES or sum(item.file_size for item in entries) > MAX_ARCHIVE_EXPANDED_BYTES:
        archive.close()
        raise Unreadable("el archivo excede el límite de entradas o de tamaño expandido; requiere revisión manual")
    return archive


def _xls_text(content: bytes) -> str:
    try:
        import xlrd
    except ImportError:
        raise Unreadable("falta instalar xlrd para leer libros .xls") from None
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise Unreadable(f"no se pudo abrir el libro .xls: {exc}") from None
    lines = []
    cells_seen = 0
    for sheet in workbook.sheets():
        lines.append(f"Hoja: {sheet.name}; filas: {sheet.nrows}; columnas: {sheet.ncols}")
        for row_number in range(sheet.nrows):
            values = []
            for column in range(sheet.ncols):
                cells_seen += 1
                if cells_seen > MAX_SPREADSHEET_CELLS:
                    lines.append(f"LECTURA PARCIAL: se superó el límite de {MAX_SPREADSHEET_CELLS} celdas")
                    workbook.release_resources()
                    return "\n".join(lines)
                value = sheet.cell_value(row_number, column)
                if value != "":
                    values.append(f"fila {row_number + 1}, columna {column + 1}={value}")
            if values:
                lines.append(" | ".join(values))
    workbook.release_resources()
    return "\n".join(lines)


def _shapefile_zip_text(content: bytes) -> str:
    try:
        import shapefile
    except ImportError:
        raise Unreadable("falta instalar pyshp para leer shapefiles") from None
    with _zip_acotado(content) as archive:
        groups: dict[str, dict[str, str]] = {}
        for name in archive.namelist():
            suffix = Path(name).suffix.casefold()
            if suffix in {".shp", ".shx", ".dbf", ".prj", ".cpg"}:
                groups.setdefault(str(Path(name).with_suffix("")), {})[suffix] = name
        fragments = []
        for layer, files in groups.items():
            if not {".shp", ".shx", ".dbf"}.issubset(files):
                continue
            reader = shapefile.Reader(
                shp=BytesIO(archive.read(files[".shp"])),
                shx=BytesIO(archive.read(files[".shx"])),
                dbf=BytesIO(archive.read(files[".dbf"])),
            )
            fields = [field[0] for field in reader.fields[1:]]
            fragments.append(
                f"Capa shapefile: {layer}; geometría: {reader.shapeTypeName}; "
                f"elementos: {len(reader)}; campos: {', '.join(fields)}"
            )
            if reader.bbox:
                fragments.append("Extensión XY (mín X, mín Y, máx X, máx Y): " + ", ".join(map(str, reader.bbox)))
            for index, record in enumerate(reader.iterRecords()):
                if index >= MAX_EVIDENCE_ITEMS:
                    fragments.append(f"LECTURA PARCIAL: se muestran {MAX_EVIDENCE_ITEMS} registros de atributos")
                    break
                fragments.append(
                    f"Registro {index + 1}: "
                    + "; ".join(f"{field}={value}" for field, value in zip(fields, record, strict=True))
                )
            if ".prj" in files:
                projection = archive.read(files[".prj"]).decode("utf-8", errors="replace")
                fragments.append("Sistema de coordenadas (PRJ): " + projection)
            else:
                fragments.append("SISTEMA DE COORDENADAS NO DECLARADO: el shapefile no incluye archivo .prj")
            reader.close()
        if not fragments:
            raise Unreadable("ZIP sin shapefile completo; se requieren archivos .shp, .shx y .dbf")
        return "\n".join(fragments)


def _pptx_text(content: bytes) -> str:
    with _zip_acotado(content) as presentation:
        slide_names = sorted(
            (name for name in presentation.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(Path(name).stem.removeprefix("slide")),
        )
        if not slide_names:
            raise Unreadable("presentación sin diapositivas legibles")
        lines = []
        for slide_number, name in enumerate(slide_names, 1):
            root = ElementTree.fromstring(presentation.read(name))
            paragraphs = []
            for node in root.iter():
                if local_name(node.tag) == "p":
                    text = "".join(child.text or "" for child in node.iter() if local_name(child.tag) == "t").strip()
                    if text:
                        paragraphs.append(text)
            lines.append(f"[Diapositiva {slide_number}]\n" + "\n".join(paragraphs))
        images = [name for name in presentation.namelist() if name.startswith("ppt/media/")]
        _append_archive_image_ocr(presentation, images, lines)
        notes = [name for name in presentation.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)]
        for note_name in notes:
            note_root = ElementTree.fromstring(presentation.read(note_name))
            note_text = [
                "".join(node.itertext()).strip()
                for node in note_root.iter()
                if local_name(node.tag) == "t" and "".join(node.itertext()).strip()
            ]
            if note_text:
                lines.append(f"Notas del presentador ({note_name}): " + " | ".join(note_text))
        diagrams = [
            name for name in presentation.namelist() if name.startswith("ppt/diagrams/") and name.endswith(".xml")
        ]
        for diagram_name in diagrams:
            diagram_root = ElementTree.fromstring(presentation.read(diagram_name))
            diagram_text = [
                "".join(node.itertext()).strip()
                for node in diagram_root.iter()
                if local_name(node.tag) == "t" and "".join(node.itertext()).strip()
            ]
            if diagram_text:
                lines.append(f"COBERTURA VISUAL PARCIAL: texto SmartArt ({diagram_name}): " + " | ".join(diagram_text))
        charts = [name for name in presentation.namelist() if name.startswith("ppt/charts/")]
        if charts:
            lines.append("CONTENIDO DE GRÁFICOS NO EXTRAÍDO: " + ", ".join(charts))
        return "\n".join(lines)


def _odt_text(content: bytes) -> str:
    with _zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for node in root.iter():
        if local_name(node.tag) in {"p", "h"}:
            text = " ".join("".join(node.itertext()).split())
            if text:
                lines.append(text)
    with _zip_acotado(content) as document:
        images = [name for name in document.namelist() if name.startswith("Pictures/")]
        _append_archive_image_ocr(document, images, lines)
    return plain_text("\n".join(lines).encode("utf-8"))


def _ods_text(content: bytes) -> str:
    with _zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for table_number, table in enumerate((n for n in root.iter() if local_name(n.tag) == "table"), 1):
        name = next(
            (value for key, value in table.attrib.items() if local_name(key) == "name"),
            f"Hoja {table_number}",
        )
        lines.append(f"Hoja: {name}")
        for row_number, row in enumerate((n for n in table.iter() if local_name(n.tag) == "table-row"), 1):
            values = []
            column_number = 0
            for cell in (n for n in row if local_name(n.tag) in {"table-cell", "covered-table-cell"}):
                repeat = next((int(v) for k, v in cell.attrib.items() if local_name(k) == "number-columns-repeated"), 1)
                value = " ".join(
                    "".join(p.itertext()).strip() for p in cell.iter() if local_name(p.tag) in {"p", "h"}
                ).strip()
                formula = next((v for k, v in cell.attrib.items() if local_name(k) == "formula"), "")
                if formula:
                    if not value:
                        lines.append(
                            f"FÓRMULA SIN VALOR CALCULADO: hoja {name}, fila {row_number}, columna {column_number + 1}"
                        )
                    value = f"{value} [fórmula: {formula}]".strip()
                for _ in range(min(repeat, 1000)):
                    column_number += 1
                    if value:
                        values.append(f"fila {row_number}, columna {column_number}={value}")
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _odp_text(content: bytes) -> str:
    with _zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    pages = [node for node in root.iter() if local_name(node.tag) == "page"]
    for number, page in enumerate(pages, 1):
        values = []
        for node in page.iter():
            if local_name(node.tag) in {"p", "h"}:
                value = " ".join("".join(node.itertext()).split())
                if value:
                    values.append(value)
        lines.append(f"[Diapositiva {number}]\n" + "\n".join(values))
    with _zip_acotado(content) as document:
        images = [name for name in document.namelist() if name.startswith(("Pictures/", "ObjectReplacements/"))]
        _append_archive_image_ocr(document, images, lines)
    return "\n".join(lines)


def _append_archive_image_ocr(archive: zipfile.ZipFile, images: list[str], lines: list[str]) -> None:
    for image_name in images[:MAX_OCR_IMAGES]:
        try:
            lines.append(f"[Imagen incrustada {image_name}; texto OCR]\n{ocr_image(archive.read(image_name))}")
            lines.append(
                f"COBERTURA VISUAL PARCIAL: el texto de {image_name} se reconoció por OCR; "
                "relaciones de tabla/figura no interpretadas automáticamente"
            )
        except Unreadable as exc:
            lines.append(f"CONTENIDO GRÁFICO NO INSPECCIONADO ({image_name}): {exc}")
    if len(images) > MAX_OCR_IMAGES:
        lines.append(f"LECTURA PARCIAL: documento con más de {MAX_OCR_IMAGES} imágenes incrustadas")


def _zip_text(content: bytes) -> str:
    """Inspecciona paquetes ZIP sin extraerlos a disco; limita expansión y cantidad de entradas."""
    with zipfile.ZipFile(BytesIO(content)) as archive:
        entries = [item for item in archive.infolist() if not item.is_dir()]
        supported_extensions = TEXT_ATTACHMENT_EXTENSIONS | {
            ".pdf",
            ".doc",
            ".docx",
            ".xls",
            ".xlsx",
            ".xlsm",
            ".ppt",
            ".pptx",
            ".odt",
            ".ods",
            ".odp",
            ".rtf",
            ".html",
            ".htm",
            ".kmz",
            ".shp",
        }
        shapefile_sidecars = {".shp", ".shx", ".dbf", ".prj", ".cpg"}
        has_shapefiles = any(Path(item.filename).suffix.casefold() == ".shp" for item in entries)
        supported = [
            item
            for item in entries
            if Path(item.filename).suffix.casefold() in supported_extensions
            and Path(item.filename).suffix.casefold() != ".shp"
        ]
        if len(entries) > MAX_ARCHIVE_FILES or sum(item.file_size for item in entries) > MAX_ARCHIVE_EXPANDED_BYTES:
            raise Unreadable("paquete ZIP excede el límite de archivos o tamaño expandido; requiere revisión manual")
        if not supported and not has_shapefiles:
            raise Unreadable("el paquete ZIP no contiene documentos de formatos legibles")
        fragments = [_shapefile_zip_text(content)] if has_shapefiles else []
        unreadable = []
        unsupported = [
            item.filename
            for item in entries
            if Path(item.filename).suffix.casefold() not in supported_extensions
            and Path(item.filename).suffix.casefold() not in shapefile_sidecars
        ]
        for item in supported:
            safe_name = posixpath.normpath(item.filename)
            if safe_name.startswith("../") or safe_name.startswith("/"):
                continue
            nested = extract_attachment_text(safe_name, "", archive.read(item))
            text, reason = nested
            fragments.append(f"Archivo dentro de ZIP: {safe_name}\n{text or reason}")
            if reason:
                unreadable.append(f"{safe_name}: {reason}")
        if unreadable:
            fragments.append("LECTURA PARCIAL DEL ZIP; elementos que requieren revisión: " + " | ".join(unreadable))
        if unsupported:
            fragments.append(
                "LECTURA PARCIAL DEL ZIP; formatos internos no inspeccionados: " + ", ".join(unsupported[:100])
            )
        return "\n\n".join(fragments)


def _pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise Unreadable("falta instalar pypdf para leer PDF") from None
    reader = PdfReader(BytesIO(content))
    page_texts: list[tuple[str, bool]] = []
    if len(reader.pages) > MAX_OCR_PAGES:
        raise Unreadable(f"PDF supera el límite de {MAX_OCR_PAGES} páginas; requiere revisión manual")
    try:
        import pypdfium2 as pdfium

        pdf_document = pdfium.PdfDocument(content)
    except ImportError:
        pdfium = None
        pdf_document = None
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pdf_page = pdf_document[page_number - 1] if pdf_document is not None else None
        has_visual_content = _pdf_visual_content(pdf_page, pdfium, page_texts)
        needs_ocr = not text.strip() or len(text.strip()) < 100 or has_visual_content
        if needs_ocr and pdf_page is not None:
            try:
                bitmap = pdf_page.render(scale=3)
                image = bitmap.to_pil()
                if image.width * image.height > MAX_OCR_PIXELS:
                    raise Unreadable("página PDF excede el límite de píxeles para OCR")
                ocr = ocr_pil_image(image)
                native_compact = re.sub(r"\W+", "", text.casefold())
                unique_ocr = [
                    line
                    for line in ocr.splitlines()
                    if re.sub(r"\W+", "", line.casefold()) and re.sub(r"\W+", "", line.casefold()) not in native_compact
                ]
                if unique_ocr:
                    text = (text.rstrip() + "\n[OCR visual adicional]\n" + "\n".join(unique_ocr)).strip()
                if has_visual_content:
                    text += (
                        "\nCOBERTURA VISUAL PARCIAL: OCR de la página ejecutado; "
                        "relaciones de tabla/figura requieren interpretación visual"
                    )
                has_visual_content = False
            except Unreadable as exc:
                text = (text.rstrip() + f"\nCONTENIDO VISUAL NO INSPECCIONADO: {exc}").strip()
                has_visual_content = True
        elif needs_ocr:
            message = "\nCONTENIDO VISUAL NO INSPECCIONADO: falta pypdfium2 para renderizar la página"
            text = (text.rstrip() + message).strip()
            has_visual_content = True
        if pdf_page is not None:
            pdf_page.close()
        page_texts.append((text, has_visual_content))
    if pdf_document is not None:
        pdf_document.close()
    if not any(page.strip() for page, _ in page_texts):
        raise Unreadable("PDF sin texto extraíble ni texto legible por OCR")
    return "\n".join(
        f"[Página {number}{'; sin texto extraíble' if not page.strip() else ''}"
        f"{'; contiene imágenes no interpretadas por OCR' if has_images else ''}]\n{page}"
        for number, (page, has_images) in enumerate(page_texts, 1)
    )


def _docx_paragraph_text(paragraph: ElementTree.Element) -> str:
    return "".join(
        "\t" if local_name(node.tag) == "tab" else "\n" if local_name(node.tag) in {"br", "cr"} else node.text or ""
        for node in paragraph.iter()
        if local_name(node.tag) in {"t", "tab", "br", "cr"}
    ).strip()


def _docx_text(content: bytes) -> str:
    with _zip_acotado(content) as document:
        root = ElementTree.fromstring(document.read("word/document.xml"))
    body = next((node for node in root.iter() if local_name(node.tag) == "body"), root)
    lines = []
    paragraph_number = table_number = 0
    for block in body:
        kind = local_name(block.tag)
        if kind == "p":
            paragraph_number += 1
            text = _docx_paragraph_text(block)
            if text:
                lines.append(f"Párrafo {paragraph_number}: {text}")
        elif kind == "tbl":
            table_number += 1
            _docx_table(block, table_number, lines)
    with _zip_acotado(content) as document:
        names = document.namelist()
        for name in names:
            if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                header_root = ElementTree.fromstring(document.read(name))
                text = " ".join(
                    "".join(node.itertext()) for node in header_root.iter() if local_name(node.tag) in {"p", "t"}
                ).strip()
                if text:
                    kind = "Encabezado" if "/header" in name else "Pie de página"
                    lines.append(f"{kind} {name}: {text}")
        for name in names:
            if name in {"word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"}:
                sidecar_root = ElementTree.fromstring(document.read(name))
                note_text = [
                    " ".join("".join(node.itertext()).split())
                    for node in sidecar_root.iter()
                    if local_name(node.tag) in {"footnote", "endnote", "comment"}
                ]
                if note_text:
                    lines.append(f"Contenido asociado ({name}): " + " | ".join(note_text))
        revisions = [node for node in root.iter() if local_name(node.tag) in {"ins", "del"}]
        if revisions:
            lines.append(
                "CAMBIOS RASTREADOS: comparar inserciones y eliminaciones antes de considerar definitivo el texto"
            )
        for revision in revisions:
            value = " ".join("".join(revision.itertext()).split())
            if value:
                lines.append(f"Revisión rastreada ({local_name(revision.tag)}): {value}")
        images = [name for name in names if name.startswith("word/media/")]
        _append_archive_image_ocr(document, images, lines)
    return "\n".join(lines)


def _xlsx_text(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise Unreadable("falta instalar openpyxl para leer Excel") from None
    _zip_acotado(content).close()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
    rows = []
    for sheet in workbook.worksheets:
        formula_sheet = formulas[sheet.title]
        rows.append(
            f"Hoja: {sheet.title}; visibilidad: {sheet.sheet_state}; rango declarado: {sheet.calculate_dimension()}"
        )
        cell_count = 0
        formula_rows = formula_sheet.iter_rows()
        for row_number, row in enumerate(sheet.iter_rows(), 1):
            if row_number > MAX_SPREADSHEET_CELLS:
                rows.append(f"LECTURA PARCIAL: hoja {sheet.title} supera {MAX_SPREADSHEET_CELLS} filas inspeccionadas")
                break
            formula_row = next(formula_rows, ())
            values = []
            for cell, formula_cell in zip(row, formula_row, strict=True):
                cell_count += 1
                if cell_count > MAX_SPREADSHEET_CELLS:
                    rows.append(
                        f"LECTURA PARCIAL: hoja {sheet.title} supera {MAX_SPREADSHEET_CELLS} celdas inspeccionadas"
                    )
                    break
                value = _xlsx_value(cell, formula_cell, sheet, rows)
                if value is not None:
                    values.append(f"{cell.coordinate}={value}")
            if values:
                rows.append(" | ".join(values))
        with _zip_acotado(content) as archive:
            images = [name for name in archive.namelist() if name.startswith("xl/media/")]
            _append_archive_image_ocr(archive, images, rows)
            charts = [name for name in archive.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", name)]
            for chart_name in charts:
                chart_root = ElementTree.fromstring(archive.read(chart_name))
                chart_text = [
                    " ".join("".join(node.itertext()).split())
                    for node in chart_root.iter()
                    if local_name(node.tag) in {"f", "v"} and "".join(node.itertext()).strip()
                ]
                rows.append(f"Gráfico ({chart_name}); referencias y valores almacenados: " + "; ".join(chart_text[:80]))
            comments = [name for name in archive.namelist() if name.startswith("xl/comments")]
            for comment_name in comments:
                comment_root = ElementTree.fromstring(archive.read(comment_name))
                for comment in (node for node in comment_root.iter() if local_name(node.tag) == "comment"):
                    reference = comment.attrib.get("ref", "celda sin referencia")
                    author_id = comment.attrib.get("authorId", "")
                    value = " ".join("".join(comment.itertext()).split())
                    rows.append(f"Comentario ({comment_name}, {reference}, autor {author_id}): {value}")
            pivots = [name for name in archive.namelist() if name.startswith("xl/pivotTables/")]
            if pivots:
                rows.append(
                    "COBERTURA PARCIAL: tablas dinámicas requieren interpretación especializada: " + ", ".join(pivots)
                )
            drawings = [
                name for name in archive.namelist() if name.startswith("xl/drawings/") and name.endswith(".xml")
            ]
            if drawings:
                rows.append("DIBUJOS/FORMAS DE HOJA PENDIENTES DE LECTURA VISUAL: " + ", ".join(drawings))
    return "\n".join(rows)


def _read_other_format(filename: str, mime_type: str, content: bytes, extension: str) -> str:
    """Extrae contenido y conserva localizadores para poder citar la fuente."""
    readers = {
        ".html": html_text,
        ".htm": html_text,
        ".rtf": rtf_text,
        ".gpx": gpx_text,
        ".waypoints": qgc_waypoints_text,
        ".plan": qgc_plan_text,
        ".geojson": geojson_text,
        ".csv": csv_text,
        ".pptx": _pptx_text,
        ".odt": _odt_text,
        ".ods": _ods_text,
        ".odp": _odp_text,
        ".zip": _zip_text,
    }
    if extension in readers:
        return readers[extension](content)
    if extension in {".kml", ".wpml"}:
        return {".kml": kml_text_file, ".wpml": wpml_text}[extension](Path(filename).name, content)
    if extension in TEXT_ATTACHMENT_EXTENSIONS or mime_type.startswith("text/"):
        return plain_text(content)
    if extension == ".kmz" or mime_type == "application/vnd.google-earth.kmz":
        return kmz_text(content)
    if extension == ".pdf" or mime_type == "application/pdf":
        return _pdf_text(content)
    if extension == ".docx":
        return _docx_text(content)
    if extension in {".xlsx", ".xlsm"}:
        return _xlsx_text(content)
    if extension == ".dwg":
        raise Unreadable("DWG requiere exportación a PDF, DXF o imagen; revise el plano original")
    if extension in {".dwg", ".dxf"}:
        raise Unreadable("plano CAD requiere revisión manual")
    return ""


def _read_by_format(filename: str, mime_type: str, content: bytes) -> str:
    extension = Path(filename).suffix.casefold()
    if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}:
        return image_text(filename, content)
    if extension == ".xls":
        return _xls_text(content)
    if extension in {".doc", ".ppt"}:
        return office_legacy_text(extension, content)
    if extension == ".shp":
        raise Unreadable("shapefile incompleto; comprima .shp, .shx y .dbf juntos en un ZIP")
    return _read_other_format(filename, mime_type, content, extension)


def _readable(filename: str, mime_type: str) -> bool:
    extension = Path(filename).suffix.casefold()
    return (
        extension in TEXT_ATTACHMENT_EXTENSIONS
        or mime_type.startswith("text/")
        or extension
        in {
            ".kmz",
            ".pdf",
            ".docx",
            ".xlsx",
            ".xlsm",
            ".pptx",
            ".odt",
            ".ods",
            ".odp",
            ".rtf",
            ".html",
            ".htm",
            ".zip",
            ".xls",
            ".doc",
            ".ppt",
            ".shp",
            ".png",
            ".jpg",
            ".jpeg",
            ".tif",
            ".tiff",
            ".bmp",
            ".webp",
            ".gif",
        }
        or mime_type in {"application/vnd.google-earth.kmz", "application/pdf"}
    )


def extract_attachment_text(filename: str, mime_type: str, content: bytes) -> tuple[str, str]:
    """Devuelve texto y motivo; el texto mantiene referencias a páginas, filas o elementos."""
    if len(content) > MAX_ATTACHMENT_BYTES:
        return "", f"supera el límite de {MAX_ATTACHMENT_BYTES // (1024 * 1024)} MB"
    if not _readable(filename, mime_type):
        extension = Path(filename).suffix.casefold()
        if extension in {".dwg", ".dxf"}:
            return "", "plano CAD requiere revisión manual"
        return "", f"formato no soportado ({extension or mime_type or 'desconocido'})"
    reason = ""
    try:
        text = _read_by_format(filename, mime_type, content).strip()
    except Unreadable as exc:
        return "", str(exc)
    except Exception as exc:
        return "", f"no se pudo leer: {exc}"
    if not text:
        return "", "el archivo no contiene texto extraíble"
    blank_pages = re.findall(r"\[Página (\d+); sin texto extraíble", text)
    if blank_pages:
        reason = "lectura parcial: páginas sin texto extraíble (posible escaneo u hoja vacía): " + ", ".join(
            blank_pages
        )
    visual_limit = (
        "contiene imágenes no interpretadas por OCR" in text
        or "CONTENIDO GRÁFICO" in text
        or "LECTURA PARCIAL" in text
        or "sin texto alternativo; requiere inspección visual" in text
        or "[Imagen:" in text
        or "CONTENIDO VISUAL NO INSPECCIONADO" in text
        or "CONTENIDO DE GRÁFICOS NO EXTRAÍDO" in text
        or "COBERTURA VISUAL PARCIAL" in text
        or "COBERTURA PARCIAL:" in text
        or "CAMBIOS RASTREADOS:" in text
    )
    if visual_limit:
        reason = (
            "lectura parcial: hay imágenes, gráficos, comentarios o secciones que requieren inspección visual/manual"
        )
    if len(text) > MAX_ATTACHMENT_TEXT:
        return text[:MAX_ATTACHMENT_TEXT], (
            f"lectura parcial: se recortó a {MAX_ATTACHMENT_TEXT:,} caracteres "
            f"de {len(text):,}; revise el archivo completo"
        )
    if "LECTURA PARCIAL DEL ZIP" in text:
        reason = "lectura parcial: uno o más archivos dentro del ZIP no se pudieron extraer por completo"
    if "SISTEMA DE COORDENADAS NO DECLARADO" in text:
        reason = "lectura parcial: shapefile sin .prj; revisar referencia espacial antes de interpretar coordenadas"
    return text, reason


def extract_attachment_text_for_work_packet(filename: str, mime_type: str, content: bytes) -> tuple[str, str]:
    """Extrae el texto completo disponible para un paquete local que se adjunta a Work.

    Usa los mismos lectores por formato, pero permite archivos mayores y no recorta a
    MAX_ATTACHMENT_TEXT. Los límites propios de cada lector (páginas, celdas y ZIP)
    siguen vigentes y se devuelven como motivo de lectura parcial.
    """
    if len(content) > MAX_WORK_PACKET_ATTACHMENT_BYTES:
        return "", (f"supera el límite del paquete local de {MAX_WORK_PACKET_ATTACHMENT_BYTES // (1024 * 1024)} MiB")
    if not _readable(filename, mime_type):
        extension = Path(filename).suffix.casefold()
        if extension in {".dwg", ".dxf"}:
            return "", "plano CAD requiere revisión manual"
        return "", f"formato no soportado ({extension or mime_type or 'desconocido'})"
    try:
        text = _read_by_format(filename, mime_type, content).strip()
    except Unreadable as exc:
        return "", str(exc)
    except Exception as exc:
        return "", f"no se pudo leer ({type(exc).__name__})"
    if not text:
        return "", "el archivo no contiene texto extraíble"

    reasons = []
    blank_pages = re.findall(r"\[Página (\d+); sin texto extraíble", text)
    if blank_pages:
        reasons.append("páginas sin texto extraíble: " + ", ".join(blank_pages))
    if any(
        marker in text
        for marker in (
            "LECTURA PARCIAL",
            "CONTENIDO VISUAL NO INSPECCIONADO",
            "CONTENIDO GRÁFICO",
            "sin texto alternativo; requiere inspección visual",
            "CONTENIDO DE GRÁFICOS NO EXTRAÍDO",
            "COBERTURA VISUAL PARCIAL",
            "COBERTURA PARCIAL:",
            "CAMBIOS RASTREADOS:",
        )
    ):
        reasons.append("hay contenido, gráficos o secciones que requieren inspección visual/manual")
    if "SISTEMA DE COORDENADAS NO DECLARADO" in text:
        reasons.append("shapefile sin .prj; revisar referencia espacial antes de interpretar coordenadas")
    return text, "; ".join(dict.fromkeys(reasons))


def _pdf_visual_content(pdf_page: Any, pdfium: Any, page_texts: list[tuple[str, bool]]) -> bool:
    has_visual_content = False
    if pdf_page is not None:
        try:
            page_width, page_height = pdf_page.get_size()
            page_area = page_width * page_height
            image_objects = pdf_page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE])
            for image_object in image_objects:
                left, bottom, right, top = image_object.get_bounds()
                image_width, image_height = max(0, right - left), max(0, top - bottom)
                image_area = image_width * image_height
                if image_area >= page_area * 0.025 or (image_width >= page_width * 0.35 and image_height >= 65):
                    has_visual_content = True
                    break
        except Exception as exc:
            page_texts.append((f"COBERTURA PARCIAL: inspección visual PDF falló ({type(exc).__name__})", True))
            has_visual_content = True
    return has_visual_content


def _docx_table(block: ElementTree.Element, table_number: int, lines: list[str]) -> None:
    for row_number, row in enumerate((node for node in block if local_name(node.tag) == "tr"), 1):
        cells = []
        for cell_number, cell in enumerate((node for node in row if local_name(node.tag) == "tc"), 1):
            value = " ".join(_docx_paragraph_text(node) for node in cell.iter() if local_name(node.tag) == "p").strip()
            cells.append(f"columna {cell_number}={value}")
        if any(value.split("=", 1)[-1] for value in cells):
            lines.append(f"Tabla {table_number}, fila {row_number}: " + " | ".join(cells))


def _xlsx_value(cell: Any, formula_cell: Any, sheet: Any, rows: list[str]) -> str | None:
    if formula_cell.data_type == "f":
        value: str | None = f"{formula_cell.value} [valor guardado: {cell.value}]"
        if cell.value is None:
            rows.append(f"FÓRMULA SIN VALOR CALCULADO: hoja {sheet.title}, celda {cell.coordinate}")
    else:
        value = str(cell.value) if cell.value is not None else None
    return value
