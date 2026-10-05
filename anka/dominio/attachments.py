"""Extracción local y trazable de texto desde adjuntos (sin llamar al modelo)."""

from __future__ import annotations

import html
from html.parser import HTMLParser
from functools import lru_cache
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_ATTACHMENT_TEXT = 30000
MAX_EVIDENCE_ITEMS = 60
MAX_EVIDENCE_CHARS = 700
MAX_ARCHIVE_FILES = 200
MAX_ARCHIVE_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_SPREADSHEET_CELLS = 50000
MAX_OCR_PAGES = 200
MAX_OCR_IMAGES = 100
MAX_OCR_PIXELS = 40_000_000
TEXT_ATTACHMENT_EXTENSIONS = {
    ".csv", ".geojson", ".json", ".kml", ".md", ".tab", ".tsv", ".txt", ".xml",
}
_KML_GEOMETRIES = {"Point", "LineString", "Polygon", "Track", "Model"}


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _child_text(element: ElementTree.Element, name: str) -> str:
    for child in element:
        if _local_name(child.tag) == name:
            return " ".join("".join(child.itertext()).split())
    return ""


def _kml_coordinates(element: ElementTree.Element) -> list[tuple[float, float]]:
    coordinates = []
    for node in element.iter():
        kind = _local_name(node.tag)
        if kind == "coord" and node.text:
            parts = node.text.split()
            if len(parts) >= 2:
                try:
                    coordinates.append((float(parts[0]), float(parts[1])))
                except ValueError:
                    pass
            continue
        if kind != "coordinates" or not node.text:
            continue
        for item in node.text.split():
            parts = item.split(",")
            if len(parts) < 2:
                continue
            try:
                longitude, latitude = float(parts[0]), float(parts[1])
            except ValueError:
                continue
            coordinates.append((longitude, latitude))
    return coordinates


def _kmz_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        kml_files = [name for name in archive.namelist() if name.casefold().endswith(".kml")]
        if not kml_files:
            raise ValueError("el KMZ no contiene un archivo KML")
        return "\n\n".join(_kml_text_file(name, archive.read(name)) for name in kml_files)


def _kml_text_file(name: str, content: bytes) -> str:
    root = ElementTree.fromstring(content)
    placemarks = [node for node in root.iter() if _local_name(node.tag) == "Placemark"]
    lines = [f"Archivo KML: {name}", f"Elementos geográficos: {len(placemarks)}"]
    all_coordinates: list[tuple[float, float]] = []
    for number, placemark in enumerate(placemarks, 1):
        label = _child_text(placemark, "name") or f"Elemento {number}"
        lines.append(f"Elemento: {label}")
        description = _child_text(placemark, "description")
        if description:
            clean = html.unescape(re.sub(r"<[^>]+>", " ", description))
            lines.append(f"Descripción: {' '.join(clean.split())[:500]}")
        attributes = []
        for node in placemark.iter():
            if _local_name(node.tag) in {"Data", "SimpleData"}:
                key = node.attrib.get("name", "")
                value = _child_text(node, "value") or " ".join("".join(node.itertext()).split())
                if key and value:
                    attributes.append(f"{key}={value}")
        if attributes:
            lines.append("Datos asociados: " + "; ".join(attributes)[:1000])
        geometries = [
            node for node in placemark.iter()
            if _local_name(node.tag) in _KML_GEOMETRIES
            and not any(_local_name(child.tag) in _KML_GEOMETRIES for child in node)
        ]
        if not geometries:
            geometries = [placemark]
        for geometry in geometries:
            kind = _local_name(geometry.tag)
            coordinates = _kml_coordinates(geometry)
            all_coordinates.extend(coordinates)
            lines.append(f"Geometría: {kind}; coordenadas: {len(coordinates)}")
            if coordinates:
                samples = "; ".join(f"{lon:.8f},{lat:.8f}" for lon, lat in coordinates[:8])
                lines.append(f"Muestras (longitud, latitud): {samples}")
    if all_coordinates:
        west = min(point[0] for point in all_coordinates)
        south = min(point[1] for point in all_coordinates)
        east = max(point[0] for point in all_coordinates)
        north = max(point[1] for point in all_coordinates)
        lines.append(
            "Extensión geográfica KML, oeste/sur/este/norte (longitud/latitud): "
            f"{west:.8f}, {south:.8f}, {east:.8f}, {north:.8f}"
        )
        if any(abs(lon) > 180 or abs(lat) > 90 for lon, lat in all_coordinates):
            lines.append("Advertencia: hay coordenadas fuera de los rangos usuales de longitud/latitud.")
    return "\n".join(lines)


class _Unreadable(Exception):
    """El adjunto no se puede leer; el motivo se muestra en el dashboard."""


def _plain_text(content: bytes) -> str:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    lines = text.splitlines()
    return "\n".join(f"Línea {number}: {line}" for number, line in enumerate(lines, 1))


class _VisibleHtml(HTMLParser):
    """Conserva texto y límites básicos de bloques; descarta etiquetas activas y estilos."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript", "svg"}:
            self.hidden_depth += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4"} and not self.hidden_depth:
            self.parts.append("\n")
        if tag == "img":
            alt = dict(attrs).get("alt") or ""
            self.parts.append(f"\n[Imagen: {alt or 'sin texto alternativo; requiere inspección visual'}]\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self.hidden_depth:
            self.hidden_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"} and not self.hidden_depth:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.hidden_depth and data.strip():
            self.parts.append(data.strip())

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "img":
            alt = dict(attrs).get("alt") or ""
            self.parts.append(f"\n[Imagen: {alt or 'sin texto alternativo; requiere inspección visual'}]\n")


def _html_text(content: bytes) -> str:
    parser = _VisibleHtml()
    parser.feed(content.decode("utf-8-sig", errors="replace"))
    lines = [line.strip() for line in " ".join(parser.parts).splitlines() if line.strip()]
    return _plain_text("\n".join(lines).encode("utf-8"))


def _rtf_text(content: bytes) -> str:
    raw = content.decode("latin-1")
    has_images = r"\pict" in raw
    raw = re.sub(
        r"\\'[0-9a-fA-F]{2}",
        lambda m: bytes([int(m.group()[2:], 16)]).decode("cp1252", errors="replace"), raw,
    )
    raw = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), raw)
    raw = re.sub(r"\\(par|line|tab)\b", lambda m: "\n" if m.group(1) != "tab" else "\t", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", raw)
    raw = raw.replace("{", "").replace("}", "")
    if has_images:
        raw += "\nCONTENIDO GRÁFICO RTF NO INSPECCIONADO"
    return _plain_text(raw.encode("utf-8"))


@lru_cache(maxsize=1)
def _ocr_engine():
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        raise _Unreadable("falta instalar rapidocr-onnxruntime para leer imágenes y PDF escaneados") from None
    return RapidOCR()


def _ocr_image(content: bytes) -> str:
    try:
        from PIL import Image, ImageOps
        import numpy as np
    except ImportError:
        raise _Unreadable("falta instalar Pillow y numpy para leer imágenes") from None
    try:
        with Image.open(BytesIO(content)) as source:
            if source.width * source.height > MAX_OCR_PIXELS:
                raise _Unreadable("imagen excede el límite de píxeles para OCR; requiere revisión manual")
            image = ImageOps.exif_transpose(source).convert("RGB")
            return _ocr_pil_image(image, np)
    except _Unreadable:
        raise
    except Exception as exc:
        raise _Unreadable(f"no se pudo ejecutar OCR sobre la imagen: {exc}") from None


def _ocr_pil_image(image, numpy_module=None) -> str:
    """Reconoce una imagen PIL ya decodificada, como las páginas renderizadas de PDFium."""
    if image.width * image.height > MAX_OCR_PIXELS:
        raise _Unreadable("imagen excede el límite de píxeles para OCR; requiere revisión manual")
    if numpy_module is None:
        try:
            import numpy as numpy_module
        except ImportError:
            raise _Unreadable("falta instalar numpy para leer imágenes") from None
    try:
        result, _ = _ocr_engine()(numpy_module.asarray(image.convert("RGB")))
    except Exception as exc:
        raise _Unreadable(f"no se pudo ejecutar OCR sobre la imagen: {exc}") from None
    lines = []
    if result:
        for entry in result:
            if len(entry) >= 3 and float(entry[2]) >= 0.25 and str(entry[1]).strip():
                lines.append(str(entry[1]).strip())
    if not lines:
        raise _Unreadable("OCR no detectó texto legible en la imagen; requiere inspección visual")
    return "\n".join(lines)


def _image_text(filename: str, content: bytes) -> str:
    return f"[Imagen: {Path(filename).name}; texto OCR]\n{_ocr_image(content)}"


def _office_legacy_text(extension: str, content: bytes) -> str:
    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter and extension == ".doc":
        converter = shutil.which("antiword")
    if not converter:
        raise _Unreadable(
            f"formato Office antiguo ({extension}); instale LibreOffice para convertirlo a texto"
        )
    with tempfile.TemporaryDirectory(prefix="anka-adjunto-") as directory:
        source = Path(directory) / f"documento{extension}"
        source.write_bytes(content)
        if Path(converter).name.casefold().startswith("antiword"):
            result = subprocess.run([converter, str(source)], capture_output=True, timeout=60, check=False)
            if result.returncode:
                raise _Unreadable("antiword no pudo convertir el documento Word antiguo")
            return _plain_text(result.stdout)
        result = subprocess.run(
            [converter, "--headless", "--convert-to", "txt:Text", "--outdir", directory, str(source)],
            capture_output=True, timeout=120, check=False,
        )
        converted = source.with_suffix(".txt")
        if result.returncode or not converted.exists():
            raise _Unreadable(f"LibreOffice no pudo convertir {extension} a texto")
        return _plain_text(converted.read_bytes())


def _xls_text(content: bytes) -> str:
    try:
        import xlrd
    except ImportError:
        raise _Unreadable("falta instalar xlrd para leer libros .xls") from None
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as exc:
        raise _Unreadable(f"no se pudo abrir el libro .xls: {exc}") from None
    lines = []
    cells_seen = 0
    for sheet in workbook.sheets():
        lines.append(f"Hoja: {sheet.name}; filas: {sheet.nrows}; columnas: {sheet.ncols}")
        for row_number in range(sheet.nrows):
            values = []
            for column in range(sheet.ncols):
                cells_seen += 1
                if cells_seen > MAX_SPREADSHEET_CELLS:
                    lines.append(f"LECTURA PARCIAL: se superó el límite de {MAX_SPREADSHEET_CELLS} celdas")
                    workbook.release_resources()
                    return "\n".join(lines)
                value = sheet.cell_value(row_number, column)
                if value != "":
                    values.append(f"fila {row_number + 1}, columna {column + 1}={value}")
            if values:
                lines.append(" | ".join(values))
    workbook.release_resources()
    return "\n".join(lines)


def _shapefile_zip_text(content: bytes) -> str:
    try:
        import shapefile
    except ImportError:
        raise _Unreadable("falta instalar pyshp para leer shapefiles") from None
    with zipfile.ZipFile(BytesIO(content)) as archive:
        groups: dict[str, dict[str, str]] = {}
        for name in archive.namelist():
            suffix = Path(name).suffix.casefold()
            if suffix in {".shp", ".shx", ".dbf", ".prj", ".cpg"}:
                groups.setdefault(str(Path(name).with_suffix("")), {})[suffix] = name
        fragments = []
        for layer, files in groups.items():
            if not {".shp", ".shx", ".dbf"}.issubset(files):
                continue
            reader = shapefile.Reader(
                shp=BytesIO(archive.read(files[".shp"])),
                shx=BytesIO(archive.read(files[".shx"])),
                dbf=BytesIO(archive.read(files[".dbf"])),
            )
            fields = [field[0] for field in reader.fields[1:]]
            fragments.append(
                f"Capa shapefile: {layer}; geometría: {reader.shapeTypeName}; "
                f"elementos: {len(reader)}; campos: {', '.join(fields)}"
            )
            if reader.bbox:
                fragments.append("Extensión XY (mín X, mín Y, máx X, máx Y): " + ", ".join(map(str, reader.bbox)))
            for index, record in enumerate(reader.iterRecords()):
                if index >= MAX_EVIDENCE_ITEMS:
                    fragments.append(f"LECTURA PARCIAL: se muestran {MAX_EVIDENCE_ITEMS} registros de atributos")
                    break
                fragments.append(f"Registro {index + 1}: " + "; ".join(
                    f"{field}={value}" for field, value in zip(fields, record, strict=True)
                ))
            if ".prj" in files:
                projection = archive.read(files[".prj"]).decode("utf-8", errors="replace")
                fragments.append("Sistema de coordenadas (PRJ): " + projection)
            reader.close()
        if not fragments:
            raise _Unreadable("ZIP sin shapefile completo; se requieren archivos .shp, .shx y .dbf")
        return "\n".join(fragments)


def _pptx_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as presentation:
        slide_names = sorted(
            (name for name in presentation.namelist()
             if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
            key=lambda name: int(re.search(r"slide(\d+)", name).group(1)),
        )
        if not slide_names:
            raise _Unreadable("presentación sin diapositivas legibles")
        lines = []
        for slide_number, name in enumerate(slide_names, 1):
            root = ElementTree.fromstring(presentation.read(name))
            paragraphs = []
            for node in root.iter():
                if _local_name(node.tag) == "p":
                    text = "".join(
                        child.text or "" for child in node.iter() if _local_name(child.tag) == "t"
                    ).strip()
                    if text:
                        paragraphs.append(text)
            lines.append(f"[Diapositiva {slide_number}]\n" + "\n".join(paragraphs))
        images = [name for name in presentation.namelist() if name.startswith("ppt/media/")]
        _append_archive_image_ocr(presentation, images, lines)
        charts = [name for name in presentation.namelist() if name.startswith("ppt/charts/")]
        if charts:
            lines.append("CONTENIDO DE GRÁFICOS NO EXTRAÍDO: " + ", ".join(charts))
        return "\n".join(lines)


def _odt_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for node in root.iter():
        if _local_name(node.tag) in {"p", "h"}:
            text = " ".join("".join(node.itertext()).split())
            if text:
                lines.append(text)
    with zipfile.ZipFile(BytesIO(content)) as document:
        images = [name for name in document.namelist() if name.startswith("Pictures/")]
        _append_archive_image_ocr(document, images, lines)
    return _plain_text("\n".join(lines).encode("utf-8"))


def _ods_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    for table_number, table in enumerate((n for n in root.iter() if _local_name(n.tag) == "table"), 1):
        name = next(
            (value for key, value in table.attrib.items() if _local_name(key) == "name"),
            f"Hoja {table_number}",
        )
        lines.append(f"Hoja: {name}")
        for row_number, row in enumerate((n for n in table.iter() if _local_name(n.tag) == "table-row"), 1):
            values = []
            column_number = 0
            for cell in (n for n in row if _local_name(n.tag) in {"table-cell", "covered-table-cell"}):
                repeat = next((int(v) for k, v in cell.attrib.items()
                               if _local_name(k) == "number-columns-repeated"), 1)
                value = " ".join("".join(p.itertext()).strip() for p in cell.iter()
                                 if _local_name(p.tag) in {"p", "h"}).strip()
                formula = next((v for k, v in cell.attrib.items() if _local_name(k) == "formula"), "")
                if formula:
                    if not value:
                        lines.append(
                            f"FÓRMULA SIN VALOR CALCULADO: hoja {name}, "
                            f"fila {row_number}, columna {column_number + 1}"
                        )
                    value = f"{value} [fórmula: {formula}]".strip()
                for _ in range(min(repeat, 1000)):
                    column_number += 1
                    if value:
                        values.append(f"fila {row_number}, columna {column_number}={value}")
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _odp_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as document:
        root = ElementTree.fromstring(document.read("content.xml"))
    lines = []
    pages = [node for node in root.iter() if _local_name(node.tag) == "page"]
    for number, page in enumerate(pages, 1):
        values = []
        for node in page.iter():
            if _local_name(node.tag) in {"p", "h"}:
                value = " ".join("".join(node.itertext()).split())
                if value:
                    values.append(value)
        lines.append(f"[Diapositiva {number}]\n" + "\n".join(values))
    with zipfile.ZipFile(BytesIO(content)) as document:
        images = [name for name in document.namelist() if name.startswith(("Pictures/", "ObjectReplacements/"))]
        _append_archive_image_ocr(document, images, lines)
    return "\n".join(lines)


def _append_archive_image_ocr(archive: zipfile.ZipFile, images: list[str], lines: list[str]) -> None:
    for image_name in images[:MAX_OCR_IMAGES]:
        try:
            lines.append(f"[Imagen incrustada {image_name}; texto OCR]\n{_ocr_image(archive.read(image_name))}")
        except _Unreadable as exc:
            lines.append(f"CONTENIDO GRÁFICO NO INSPECCIONADO ({image_name}): {exc}")
    if len(images) > MAX_OCR_IMAGES:
        lines.append(f"LECTURA PARCIAL: documento con más de {MAX_OCR_IMAGES} imágenes incrustadas")


def _zip_text(content: bytes) -> str:
    """Inspecciona paquetes ZIP sin extraerlos a disco; limita expansión y cantidad de entradas."""
    with zipfile.ZipFile(BytesIO(content)) as archive:
        entries = [item for item in archive.infolist() if not item.is_dir()]
        supported_extensions = TEXT_ATTACHMENT_EXTENSIONS | {
            ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".odt", ".ods", ".odp", ".rtf", ".html", ".htm", ".kml"
        }
        supported = [item for item in entries if Path(item.filename).suffix.casefold() in supported_extensions]
        if len(entries) > MAX_ARCHIVE_FILES or sum(item.file_size for item in entries) > MAX_ARCHIVE_EXPANDED_BYTES:
            raise _Unreadable("paquete ZIP excede el límite de archivos o tamaño expandido; requiere revisión manual")
        if not supported:
            raise _Unreadable("el paquete ZIP no contiene documentos de formatos legibles")
        fragments = []
        unreadable = []
        unsupported = [
            item.filename for item in entries
            if Path(item.filename).suffix.casefold() not in supported_extensions
        ]
        for item in supported:
            safe_name = posixpath.normpath(item.filename)
            if safe_name.startswith("../") or safe_name.startswith("/"):
                continue
            nested = extract_attachment_text(safe_name, "", archive.read(item))
            text, reason = nested
            fragments.append(f"Archivo dentro de ZIP: {safe_name}\n{text or reason}")
            if reason:
                unreadable.append(f"{safe_name}: {reason}")
        if unreadable:
            fragments.append("LECTURA PARCIAL DEL ZIP; elementos que requieren revisión: " + " | ".join(unreadable))
        if unsupported:
            fragments.append(
                "LECTURA PARCIAL DEL ZIP; formatos internos no inspeccionados: "
                + ", ".join(unsupported[:100])
            )
        return "\n\n".join(fragments)


def _pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise _Unreadable("falta instalar pypdf para leer PDF") from None
    reader = PdfReader(BytesIO(content))
    page_texts = []
    if len(reader.pages) > MAX_OCR_PAGES:
        raise _Unreadable(f"PDF supera el límite de {MAX_OCR_PAGES} páginas; requiere revisión manual")
    for page_number, page in enumerate(reader.pages, 1):
        try:
            has_images = bool(getattr(page, "images", []))
        except Exception:
            has_images = True  # no podemos certificar la cobertura visual de esta página
        text = page.extract_text() or ""
        if not text.strip() or has_images:
            try:
                import pypdfium2 as pdfium
                pdf_document = pdfium.PdfDocument(content)
                pdf_page = pdf_document[page_number - 1]
                bitmap = pdf_page.render(scale=2)
                image = bitmap.to_pil()
                if image.width * image.height > MAX_OCR_PIXELS:
                    raise _Unreadable("página PDF excede el límite de píxeles para OCR")
                ocr = _ocr_pil_image(image)
                text = (text.rstrip() + "\n[Texto reconocido por OCR]\n" + ocr).strip()
                pdf_page.close()
                pdf_document.close()
                has_images = False
            except _Unreadable as exc:
                text = (text.rstrip() + f"\nCONTENIDO VISUAL NO INSPECCIONADO: {exc}").strip()
                has_images = True
            except ImportError:
                message = "\nCONTENIDO VISUAL NO INSPECCIONADO: falta pypdfium2 para renderizar esta página"
                text = (text.rstrip() + message).strip()
                has_images = True
        page_texts.append((text, has_images))
    if not any(page.strip() for page, _ in page_texts):
        raise _Unreadable("PDF sin texto extraíble ni texto legible por OCR")
    return "\n".join(
        f"[Página {number}{'; sin texto extraíble' if not page.strip() else ''}"
        f"{'; contiene imágenes no interpretadas por OCR' if has_images else ''}]\n{page}"
        for number, (page, has_images) in enumerate(page_texts, 1)
    )


def _docx_paragraph_text(paragraph: ElementTree.Element) -> str:
    return "".join(
        "\t" if _local_name(node.tag) == "tab" else "\n" if _local_name(node.tag) in {"br", "cr"}
        else node.text or ""
        for node in paragraph.iter()
        if _local_name(node.tag) in {"t", "tab", "br", "cr"}
    ).strip()


def _docx_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as document:
        root = ElementTree.fromstring(document.read("word/document.xml"))
    body = next((node for node in root.iter() if _local_name(node.tag) == "body"), root)
    lines = []
    paragraph_number = table_number = 0
    for block in body:
        kind = _local_name(block.tag)
        if kind == "p":
            paragraph_number += 1
            text = _docx_paragraph_text(block)
            if text:
                lines.append(f"Párrafo {paragraph_number}: {text}")
        elif kind == "tbl":
            table_number += 1
            for row_number, row in enumerate(
                (node for node in block if _local_name(node.tag) == "tr"), 1
            ):
                cells = []
                for cell_number, cell in enumerate(
                    (node for node in row if _local_name(node.tag) == "tc"), 1
                ):
                    value = " ".join(
                        _docx_paragraph_text(node) for node in cell.iter()
                        if _local_name(node.tag) == "p"
                    ).strip()
                    cells.append(f"columna {cell_number}={value}")
                if any(value.split("=", 1)[-1] for value in cells):
                    lines.append(f"Tabla {table_number}, fila {row_number}: " + " | ".join(cells))
    with zipfile.ZipFile(BytesIO(content)) as document:
        names = document.namelist()
        for name in names:
            if re.fullmatch(r"word/(header|footer)\d+\.xml", name):
                root = ElementTree.fromstring(document.read(name))
                text = " ".join("".join(node.itertext()) for node in root.iter()
                               if _local_name(node.tag) in {"p", "t"}).strip()
                if text:
                    kind = "Encabezado" if "/header" in name else "Pie de página"
                    lines.append(f"{kind} {name}: {text}")
        images = [name for name in names if name.startswith("word/media/")]
        if images:
            lines.append("CONTENIDO GRÁFICO NO INSPECCIONADO (sin OCR): " + ", ".join(images))
    return "\n".join(lines)


def _xlsx_text(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise _Unreadable("falta instalar openpyxl para leer Excel") from None
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
    rows = []
    for sheet in workbook.worksheets:
        formula_sheet = formulas[sheet.title]
        rows.append(
            f"Hoja: {sheet.title}; visibilidad: {sheet.sheet_state}; "
            f"rango declarado: {sheet.calculate_dimension()}"
        )
        cell_count = 0
        formula_rows = formula_sheet.iter_rows()
        for row_number, row in enumerate(sheet.iter_rows(), 1):
            if row_number > MAX_SPREADSHEET_CELLS:
                rows.append(f"LECTURA PARCIAL: hoja {sheet.title} supera {MAX_SPREADSHEET_CELLS} filas inspeccionadas")
                break
            formula_row = next(formula_rows, ())
            values = []
            for cell, formula_cell in zip(row, formula_row, strict=True):
                cell_count += 1
                if cell_count > MAX_SPREADSHEET_CELLS:
                    rows.append(
                        f"LECTURA PARCIAL: hoja {sheet.title} supera "
                        f"{MAX_SPREADSHEET_CELLS} celdas inspeccionadas"
                    )
                    break
                if formula_cell.data_type == "f":
                    value = f"{formula_cell.value} [valor guardado: {cell.value}]"
                    if cell.value is None:
                        rows.append(
                            f"FÓRMULA SIN VALOR CALCULADO: hoja {sheet.title}, "
                            f"celda {cell.coordinate}"
                        )
                else:
                    value = cell.value
                if value is not None:
                    values.append(f"{cell.coordinate}={value}")
            if values:
                rows.append(" | ".join(values))
        with zipfile.ZipFile(BytesIO(content)) as archive:
            omitted = [name for name in archive.namelist()
                       if name.startswith(("xl/charts/", "xl/drawings/", "xl/comments"))]
        if omitted:
            rows.append("CONTENIDO GRÁFICO/COMENTARIOS NO EXTRAÍDO: " + ", ".join(omitted))
    return "\n".join(rows)


def _read_other_format(filename: str, mime_type: str, content: bytes, extension: str) -> str:
    """Extrae contenido y conserva localizadores para poder citar la fuente."""
    if extension == ".kml":
        return _kml_text_file(Path(filename).name, content)
    if extension in {".html", ".htm"}:
        return _html_text(content)
    if extension == ".rtf":
        return _rtf_text(content)
    if extension in {".pptx", ".odt", ".ods", ".odp"}:
        return {
            ".pptx": lambda: _pptx_text(content),
            ".odt": lambda: _odt_text(content),
            ".ods": lambda: _ods_text(content),
            ".odp": lambda: _odp_text(content),
        }[extension]()
    if extension == ".zip":
        return _zip_text(content)
    if extension in TEXT_ATTACHMENT_EXTENSIONS or mime_type.startswith("text/"):
        return _plain_text(content)
    if extension == ".kmz" or mime_type == "application/vnd.google-earth.kmz":
        return _kmz_text(content)
    if extension == ".pdf" or mime_type == "application/pdf":
        return _pdf_text(content)
    if extension == ".docx":
        return _docx_text(content)
    if extension in {".xlsx", ".xlsm"}:
        return _xlsx_text(content)
    if extension == ".dwg":
        raise _Unreadable("DWG requiere exportación a PDF, DXF o imagen; revise el plano original")
    if extension in {".dwg", ".dxf"}:
        raise _Unreadable("plano CAD requiere revisión manual")
    return ""


def _read_by_format(filename: str, mime_type: str, content: bytes) -> str:
    extension = Path(filename).suffix.casefold()
    if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}:
        return _image_text(filename, content)
    if extension == ".xls":
        return _xls_text(content)
    if extension in {".doc", ".ppt"}:
        return _office_legacy_text(extension, content)
    if extension == ".shp":
        raise _Unreadable("shapefile incompleto; comprima .shp, .shx y .dbf juntos en un ZIP")
    return _read_other_format(filename, mime_type, content, extension)


def _evidence_location(line: str, current: str) -> tuple[str, str]:
    match = re.match(r"^\[(Página\s+\d+)(?:;([^\]]+))?\]\s*(.*)$", line)
    if match:
        detail = " ".join(part for part in (match.group(2), match.group(3)) if part)
        return match.group(1), detail.strip()
    if line.startswith("Hoja: "):
        return line.split(";", 1)[0], ""
    if line.startswith("Archivo KML:"):
        return line, ""
    if line.startswith("Archivo dentro de ZIP:"):
        return line, ""
    if line.startswith("LECTURA PARCIAL DEL ZIP"):
        return "Cobertura del ZIP", line
    if line.startswith("Elementos geográficos:"):
        return "Resumen KML", line
    if line.startswith(("Geometría:", "Muestras (longitud, latitud):", "Extensión geográfica KML")):
        return current, line
    if line.startswith("[") and line.endswith("]"):
        return line.strip("[]"), ""
    if line.startswith("Párrafo ") or line.startswith("Tabla ") or line.startswith("Elemento: "):
        return line.split(":", 1)[0], line.partition(":")[2].strip()
    if re.match(r"^[A-Z]{1,3}\d+=", line):
        return current, line
    return current, line


def attachment_reviewed_sections(text: str) -> list[str]:
    """Enumera todos los localizadores del texto extraído, sin el límite de citas de la ficha."""
    sections = []
    current = "Contenido extraído"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        location, _ = _evidence_location(line, current)
        if location != current:
            current = location
            if current not in sections:
                sections.append(current)
    return sections or ["Contenido sin paginación"]


def build_attachment_evidence(text: str) -> list[dict[str, str]]:
    """Resume fragmentos con localizadores estables para la ficha del caso."""
    evidence: list[dict[str, str]] = []
    current_location = "Contenido extraído"
    excerpt: list[str] = []

    def flush() -> None:
        value = " ".join(part.strip() for part in excerpt if part.strip()).strip()
        if value and len(evidence) < MAX_EVIDENCE_ITEMS:
            evidence.append({"topic": "extracto", "value": "", "location": current_location,
                             "excerpt": value[:MAX_EVIDENCE_CHARS]})

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        location, content = _evidence_location(line, current_location)
        if location != current_location:
            flush()
            excerpt.clear()
            current_location = location
        if content:
            excerpt.append(content)
    flush()
    return evidence


def _readable(filename: str, mime_type: str) -> bool:
    extension = Path(filename).suffix.casefold()
    return (
        extension in TEXT_ATTACHMENT_EXTENSIONS
        or mime_type.startswith("text/")
        or extension in {
            ".kmz", ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".odt", ".ods", ".odp",
            ".rtf", ".html", ".htm", ".zip", ".xls", ".doc", ".ppt", ".shp",
            ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif",
        }
        or mime_type in {"application/vnd.google-earth.kmz", "application/pdf"}
    )


def extract_attachment_text(filename: str, mime_type: str, content: bytes) -> tuple[str, str]:
    """Devuelve texto y motivo; el texto mantiene referencias a páginas, filas o elementos."""
    if len(content) > MAX_ATTACHMENT_BYTES:
        return "", f"supera el límite de {MAX_ATTACHMENT_BYTES // (1024 * 1024)} MB"
    if not _readable(filename, mime_type):
        extension = Path(filename).suffix.casefold()
        if extension in {".dwg", ".dxf"}:
            return "", "plano CAD requiere revisión manual"
        return "", f"formato no soportado ({extension or mime_type or 'desconocido'})"
    reason = ""
    try:
        text = _read_by_format(filename, mime_type, content).strip()
    except _Unreadable as exc:
        return "", str(exc)
    except Exception as exc:
        return "", f"no se pudo leer: {exc}"
    if not text:
        return "", "el archivo no contiene texto extraíble"
    blank_pages = re.findall(r"\[Página (\d+); sin texto extraíble", text)
    if blank_pages:
        reason = (
            "lectura parcial: páginas sin texto extraíble (posible escaneo u hoja vacía): "
            + ", ".join(blank_pages)
        )
    visual_limit = (
        "contiene imágenes no interpretadas por OCR" in text
        or "CONTENIDO GRÁFICO" in text
        or "LECTURA PARCIAL" in text
        or "sin texto alternativo; requiere inspección visual" in text
        or "[Imagen:" in text
        or "CONTENIDO VISUAL NO INSPECCIONADO" in text
        or "CONTENIDO DE GRÁFICOS NO EXTRAÍDO" in text
    )
    if visual_limit:
        reason = (
            "lectura parcial: hay imágenes, gráficos, comentarios o secciones "
            "que requieren inspección visual/manual"
        )
    if len(text) > MAX_ATTACHMENT_TEXT:
        return text[:MAX_ATTACHMENT_TEXT], (
            f"lectura parcial: se recortó a {MAX_ATTACHMENT_TEXT:,} caracteres "
            f"de {len(text):,}; revise el archivo completo"
        )
    if "LECTURA PARCIAL DEL ZIP" in text:
        reason = "lectura parcial: uno o más archivos dentro del ZIP no se pudieron extraer por completo"
    return text, reason
