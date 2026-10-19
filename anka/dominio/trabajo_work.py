"""Un trabajo Python → Work: su identidad, la referencia verificable al mensaje y la hoja `.xlsx` que Work lee (T-49).

Reglas puras: sin Drive, sin SQLite y sin reloj propio.
"""

from __future__ import annotations

import hashlib
import json
from collections.abc import Mapping, Sequence
from email.utils import parseaddr
from io import BytesIO
from typing import Any, NamedTuple

from anka.dominio import contexto_previo
from anka.dominio.analisis import Adjunto
from anka.dominio.reception import NOISE_SIGNALS, mentions_commercial_terms, noise_signal_id

PREFIJO_DE_TRABAJO = "ANKAJOB-"
TIPO_ANALIZAR = "analizar_solicitud"
TIPOS = (TIPO_ANALIZAR,)
VERSION_DEL_CONTRATO = "ANKA_DASHBOARD_JSON_V1"

PENDIENTE_DE_PUBLICAR = "pendiente_de_publicar"
PUBLICANDO = "publicando"
PUBLICADO = "publicado"
CONCILIADO = "conciliado"
AMBIGUO = "ambiguo"
CANCELADO = "cancelado"
ESTADOS = (PENDIENTE_DE_PUBLICAR, PUBLICANDO, PUBLICADO, CONCILIADO, AMBIGUO, CANCELADO)
# Estados desde los que una importación de Work cierra el trabajo (T-49, punto 8).
CONCILIABLES = (PENDIENTE_DE_PUBLICAR, PUBLICANDO, PUBLICADO)

CLAVE_DE_MARCA = "anka_job_id"      # propiedad del archivo en Drive (`appProperties`): lo que lo ata a su trabajo
HOJAS = ("TRABAJO", "ADJUNTOS_ESPERADOS", "VERIFICACION")


def id_de_trabajo(account: str, thread_id: str, message_id: str, job_type: str, contract_version: str) -> str:
    """`ANKAJOB-` + 16 hexadecimales de SHA-256: el mismo mensaje, tipo y versión dan siempre el mismo trabajo."""
    partes = [str(parte or "").strip() for parte in (account, thread_id, message_id, job_type, contract_version)]
    if not all(partes):
        raise ValueError("Falta un dato para identificar el trabajo.")
    huella = hashlib.sha256("|".join(partes).encode("utf-8")).hexdigest()
    return f"{PREFIJO_DE_TRABAJO}{huella[:16]}"


def nombre_de_archivo(job_id: str) -> str:
    """Nombre de la hoja en Drive; solo una pista, la identidad es la marca `anka_job_id`."""
    return f"ANKA_JOB_{job_id}"


def referencia_del_mensaje(fila: Mapping[str, Any]) -> dict[str, Any]:
    """Los metadatos del mensaje que Work compara con Gmail. Nunca el cuerpo ni el texto de los adjuntos."""
    adjuntos = []
    for item in _lista(fila.get("attachments_json")):
        if isinstance(item, dict):
            adjuntos.append({"filename": str(item.get("filename") or ""), "mime_type": str(item.get("mime_type") or ""),
                             "size_bytes": item.get("size_bytes") if isinstance(item.get("size_bytes"), int) else None,
                             "gmail_attachment_id": str(item.get("gmail_attachment_id") or "")})
    return {"gmail_thread_id": str(fila.get("thread_id") or ""), "gmail_message_id": str(fila.get("message_id") or ""),
            "sender": str(fila.get("sender") or ""), "subject": str(fila.get("subject") or ""),
            "internal_date": str(fila.get("internal_date") or ""), "header_date": str(fila.get("header_date") or ""),
            "attachments": adjuntos}


def huella_de_la_referencia(referencia: Mapping[str, Any]) -> str:
    """SHA-256 del JSON canónico de la referencia: si cambia algo del mensaje, cambia la huella."""
    canonico = json.dumps(referencia, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonico.encode("utf-8")).hexdigest()


def _lista(valor: object) -> list[Any]:
    if isinstance(valor, list):
        return valor
    try:
        cargado = json.loads(valor) if isinstance(valor, str) and valor else []
    except ValueError:
        return []
    return cargado if isinstance(cargado, list) else []


INSTRUCCION_PARA_WORK = (
    "Esta hoja es de solo lectura. Abre en Gmail el hilo `thread_id` y el mensaje `message_id`; antes de analizar "
    "compara remitente, asunto, fecha y adjuntos con la hoja VERIFICACION. Si algo no coincide, no analices y avisa. "
    "Devuelve el borrador de siempre con [ANKA_DASHBOARD_JSON_V1]. No escribas en esta hoja. La "
    "clasificacion_preliminar es solo una pista de Python: confírmala o corrígela en el campo opcional "
    "`classification` "
    "(SOLICITUD_NUEVA, SEGUIMIENTO, NO_SOLICITUD o REVISION_MANUAL) y no te limites a palabras clave.")

PASOS_DE_VERIFICACION = (
    ("1", "El hilo de Gmail con el thread_id de la hoja existe y contiene el message_id de la hoja."),
    ("2", "Remitente, asunto y fecha del mensaje coinciden con la referencia (TRABAJO)."),
    ("3", "Los adjuntos del mensaje coinciden con ADJUNTOS_ESPERADOS (nombre y tamaño)."),
    ("4", "Solo se analiza ese mensaje; los demás mensajes del hilo son contexto, no son este trabajo."),
    ("5", "El borrador devuelto lleva el mismo thread_id y message_id que esta hoja."),
)


def construir_libro_del_trabajo(trabajo: Mapping[str, Any]) -> bytes:
    """El `.xlsx` del trabajo: la hoja TRABAJO (campo/valor), ADJUNTOS_ESPERADOS y VERIFICACION."""
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet

    referencia = trabajo.get("referencia") or {}
    libro = Workbook()
    principal = libro.active
    assert isinstance(principal, Worksheet)
    principal.title = HOJAS[0]
    filas: list[tuple[str, str]] = [
        ("job_id", str(trabajo.get("job_id") or "")), ("thread_id", str(trabajo.get("thread_id") or "")),
        ("message_id", str(trabajo.get("message_id") or "")), ("tipo_de_trabajo", str(trabajo.get("job_type") or "")),
        ("estado_del_trabajo", str(trabajo.get("state") or "")),
        ("version_del_contrato", str(trabajo.get("contract_version") or "")),
        ("fecha_de_creacion", str(trabajo.get("created_at") or "")),
        ("ref_sha256", str(trabajo.get("ref_sha256") or "")),
        ("remitente", str(referencia.get("sender") or "")), ("asunto", str(referencia.get("subject") or "")),
        ("fecha_interna_gmail", str(referencia.get("internal_date") or "")),
        ("fecha_del_encabezado", str(referencia.get("header_date") or "")),
        ("clasificacion_preliminar", str(trabajo.get("clasificacion_preliminar") or "")),
        ("motivo_preliminar", str(trabajo.get("motivo_preliminar") or "")),
        ("instruccion_para_work", INSTRUCCION_PARA_WORK),
    ]
    paquete = str(trabajo.get("paquete_version") or "")
    if paquete:                                  # solo los trabajos creados con versión de paquete (T-59, punto 14)
        filas.insert(6, ("version_del_paquete", paquete))
    contexto = trabajo.get("contexto_previo")
    if isinstance(contexto, contexto_previo.Contexto):
        filas.append(("instruccion_de_seguimiento", contexto_previo.INSTRUCCION_DE_SEGUIMIENTO))
    principal.append(["campo", "valor"])
    for fila in filas:
        principal.append(list(fila))
    _estilo(principal, (28, 110))
    adjuntos = libro.create_sheet(HOJAS[1])
    adjuntos.append(["nombre", "tipo", "tamano_bytes", "id_de_adjunto_gmail"])
    adjuntos_esperados: Sequence[Mapping[str, Any]] = referencia.get("attachments") or []
    for adj in adjuntos_esperados:
        adjuntos.append([adj.get("filename", ""), adj.get("mime_type", ""),
                         adj.get("size_bytes") if adj.get("size_bytes") is not None else "",
                         adj.get("gmail_attachment_id", "")])
    _estilo(adjuntos, (40, 30, 16, 50))
    pasos = libro.create_sheet(HOJAS[2])
    pasos.append(["paso", "que_debe_comprobar_work"])
    for paso in PASOS_DE_VERIFICACION:
        pasos.append(list(paso))
    _estilo(pasos, (8, 110))
    if isinstance(contexto, contexto_previo.Contexto):
        for nombre, encabezados, datos in contexto_previo.hojas(contexto):
            hoja_de_contexto = libro.create_sheet(nombre)
            hoja_de_contexto.append(list(encabezados))
            for dato in datos:
                hoja_de_contexto.append(list(dato))
            _estilo(hoja_de_contexto, (24, 110) if len(encabezados) == 2 else (22,) * len(encabezados))
    for hoja in libro.worksheets:     # texto, nunca fórmulas: openpyxl toma por fórmula lo que empieza con "="
        for fila_de_celdas in hoja.iter_rows():
            for celda in fila_de_celdas:
                if isinstance(celda.value, str) and celda.value.startswith("="):
                    celda.value = "'" + celda.value
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _estilo(hoja: Any, anchos: Sequence[int]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for posicion, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(posicion)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")


# --- Elegibilidad (T-51, D-155) --------------------------------------------------------------------------------------

ELEGIBLE = "elegible"
EXCLUIDO = "excluido"
REVISION = "revision"
SENAL_NOTIFICACION_DE_CHATGPT = "notificacion_de_chatgpt"
# Dominios desde los que las tareas de ChatGPT avisan de sus ejecuciones: nunca son una solicitud comercial.
DOMINIOS_DE_NOTIFICACION_DE_CHATGPT = frozenset({"tm.openai.com"})
MOTIVOS_DE_EXCLUSION = {**NOISE_SIGNALS,
                        SENAL_NOTIFICACION_DE_CHATGPT: "notificación automática de una tarea de ChatGPT"}
MOTIVO_AVISO_COMERCIAL = "aviso_automatico_con_terminos_comerciales"
MOTIVO_SIN_REMITENTE = "remitente_desconocido"
MOTIVO_SIN_CONFIGURACION = "terminos_comerciales_no_disponibles"


class Veredicto(NamedTuple):
    """Resultado de la elegibilidad: `estado` (elegible, excluido o revision), `motivo_id` y su texto genérico."""

    estado: str
    motivo_id: str
    motivo: str

    @property
    def elegible(self) -> bool:
        """¿Se puede crear y publicar el trabajo?"""
        return self.estado == ELEGIBLE


def _mensaje_de(remitente: str, asunto: str = "") -> dict[str, Any]:
    """Un mensaje con solo las cabeceras que el clasificador de R1 mira, armado desde lo que guarda la cola."""
    cabeceras = [{"name": "From", "value": remitente}]
    if asunto:
        cabeceras.append({"name": "Subject", "value": asunto})
    return {"payload": {"headers": cabeceras}}


def _es_notificacion_de_chatgpt(direccion: str) -> bool:
    dominio = direccion.rpartition("@")[2].casefold()
    return any(dominio == base or dominio.endswith("." + base) for base in DOMINIOS_DE_NOTIFICACION_DE_CHATGPT)


def _senal_de_automatizacion(fila: Mapping[str, Any], remitentes: Sequence[str]) -> str | None:
    """La señal de ruido del mensaje: la que guardó el receptor o, si no hay, la que se ve en el remitente guardado.

    Con varios remitentes solo cuenta si TODOS son automáticos (R1)."""
    try:
        evidencia = json.loads(str(fila.get("evidence_json") or "{}"))
    except ValueError:
        evidencia = {}
    guardada = evidencia.get("noise_signal") if isinstance(evidencia, dict) else None
    if guardada in MOTIVOS_DE_EXCLUSION:
        return str(guardada)
    senales = [SENAL_NOTIFICACION_DE_CHATGPT if _es_notificacion_de_chatgpt(r) else noise_signal_id(_mensaje_de(r))
               for r in remitentes]
    return senales[0] if senales and all(senales) else None


def evaluar_elegibilidad(fila: Mapping[str, Any], terminos_comerciales: Sequence[str] | None = ()) -> Veredicto:
    """¿Se puede crear un trabajo Work para este mensaje de la cola de recepción? (T-51)

    Reutiliza R1: solo el ruido evidente se excluye; un aviso automático que habla de cotizaciones, o una fila sin
    remitente, queda en `revision` y nunca se descarta en silencio. Un remitente humano siempre es elegible.
    `terminos_comerciales=None` dice que no se pudieron leer: un aviso automático no se puede descartar sin ellos.
    """
    remitentes = [direccion for direccion in (parseaddr(parte)[1] for parte in str(fila.get("sender") or "").split(","))
                  if "@" in direccion]
    if not remitentes:
        return Veredicto(REVISION, MOTIVO_SIN_REMITENTE, "no se pudo leer el remitente del mensaje")
    senal = _senal_de_automatizacion(fila, remitentes)
    if senal is None:
        return Veredicto(ELEGIBLE, "", "")
    if senal != SENAL_NOTIFICACION_DE_CHATGPT:
        if terminos_comerciales is None:
            return Veredicto(REVISION, MOTIVO_SIN_CONFIGURACION, "no se pudieron leer los términos comerciales")
        adjuntos: list[Adjunto] = [{"filename": str(a.get("filename") or ""), "text": ""}
                                   for a in _lista(fila.get("attachments_json")) if isinstance(a, dict)]
        mensaje = _mensaje_de(remitentes[0], str(fila.get("subject") or ""))
        if mentions_commercial_terms([mensaje], list(terminos_comerciales), adjuntos):
            return Veredicto(REVISION, MOTIVO_AVISO_COMERCIAL, "aviso automático que menciona términos comerciales")
    return Veredicto(EXCLUIDO, senal, MOTIVOS_DE_EXCLUSION[senal])
