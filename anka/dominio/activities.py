"""Edicion y exportacion de la tabla de actividades (sin acceso a Gmail ni Drive)."""

from __future__ import annotations

import hashlib
import json
import re

from anka.dominio.errores import ErrorDeValidacion
from anka.dominio.analisis import Actividad

from collections.abc import Mapping, Sequence
from copy import copy
from io import BytesIO
from typing import Any, cast

ACTIVITY_COLUMNS = [
    ("activity", "Actividad"),
    ("servicio_tecnica", "Servicio/técnica"),
    ("ubicacion", "Ubicación"),
    ("quantity", "Cantidad"),
    ("unit", "Unidad"),
    ("entregable", "Entregable"),
    ("plazo_restriccion", "Plazo/restricción"),
    ("fuente", "Fuente"),
    ("estado", "Estado"),
    ("notes", "Notas"),
]
ESTADOS = ("confirmado", "pendiente", "contradictorio")


def normalize_estado(value: object) -> str:
    """Lleva cualquier texto de estado a uno de los tres valores (R12).

    «confirmado; soporte pendiente» o «confirmado en fuentes; revisión pendiente» no están
    confirmados del todo: si el texto menciona algo pendiente o por aclarar, es `pendiente`.
    """
    text = " ".join(str(value or "").casefold().split())
    if "contradic" in text:
        return "contradictorio"
    if any(mark in text for mark in _MARCAS_DE_DUDA):
        return "pendiente"
    return "confirmado" if text.startswith("confirmado") else "pendiente"


# Textos que, junto a «confirmado», dicen que el dato aún no está confirmado del todo (R12). Los verbos van con su
# preposición («por verificar», «a validar») para que «confirmado, verificado» siga siendo confirmado.
_MARCAS_DE_DUDA = (
    "pendiente", "revisi", "revisar", "requiere", "aclar", "falta", "no indicad",
    "por verific", "a verific", "sin verific", "por valid", "a valid", "sin valid",
    "por defin", "a defin", "sin defin", "por confirm", "sin confirm", "sin soporte",
    "supuest", "asumid", "dudos", "estimad",
)
_UNIDADES = {
    "ha": "ha", "hectarea": "ha", "hectareas": "ha", "hectárea": "ha", "hectáreas": "ha",
    "km": "km", "kilometro": "km", "kilometros": "km", "kilómetro": "km", "kilómetros": "km",
    "m": "m", "mt": "m", "mts": "m", "metro": "m", "metros": "m",
    "km2": "km²", "km²": "km²", "kilometros cuadrados": "km²", "kilómetros cuadrados": "km²",
    "m2": "m²", "m²": "m²", "metros cuadrados": "m²",
}

_UNIT_PATTERNS = (
    ("km²", re.compile(r"(?<!\w)(?:km²|km2|kil[oó]metros? cuadrados)(?!\w)", re.IGNORECASE)),
    ("m²", re.compile(r"(?<!\w)(?:m²|m2|metros? cuadrados)(?!\w)", re.IGNORECASE)),
    ("ha", re.compile(r"(?<!\w)(?:ha|hect[aá]reas?)(?!\w)", re.IGNORECASE)),
    ("km", re.compile(r"(?<!\w)(?:km|kil[oó]metros?)(?!\w)", re.IGNORECASE)),
    ("m", re.compile(r"(?<!\w)(?:m|mt|mts|metros?)(?!\w)", re.IGNORECASE)),
)


def unidades_en_cantidad(quantity: object) -> set[str]:
    """Unidades distintas mencionadas explícitamente en una cantidad compuesta."""
    text = str(quantity or "")
    return {unit for unit, pattern in _UNIT_PATTERNS if pattern.search(text)}


def unidad_canonica(unidad: str) -> str:
    """La unidad escrita de una sola forma («Ha», «hectáreas» → «ha»); una que no se reconoce queda como vino."""
    return _UNIDADES.get(unidad.strip().rstrip(".").casefold(), unidad)


def limpiar_cantidad_y_unidad(row: dict) -> None:
    """Evita que la unidad aparezca dos veces: el modelo suele escribir «544.5 ha» en la cantidad y «ha» en la unidad.

    Si la cantidad es solo «número + unidad», se queda con el número; si ya menciona todas las unidades (varias
    magnitudes en un texto), se vacía la unidad. Una cantidad sin la unidad no se toca. Al final la unidad se escribe
    de una sola forma (`unidad_canonica`).
    """
    quantity, unit = str(row.get("quantity") or "").strip(), str(row.get("unit") or "").strip()
    if quantity and unit:
        _quitar_unidad_repetida(row, quantity, unit)
    if row.get("unit"):
        row["unit"] = unidad_canonica(str(row["unit"]))


def _quitar_unidad_repetida(row: dict, quantity: str, unit: str) -> None:
    plain = re.fullmatch(rf"([\d.,]+)\s*{re.escape(unit)}\.?", quantity, re.IGNORECASE)
    if plain:
        row["quantity"] = plain.group(1)
        return
    tokens = [token for token in re.split(r"\s*(?:[,;/]|\sy\s)\s*", unit) if token]
    if tokens and all(re.search(rf"(?<!\w){re.escape(token)}(?!\w)", quantity, re.IGNORECASE) for token in tokens):
        row["unit"] = ""


MAX_ROWS = 200      # tope de filas que se acepta de un formulario: una tabla desmesurada no se guarda
MAX_CELL = 2000      # tope de caracteres por celda, por la misma razón


def huella_de_tabla(activities: Sequence[Mapping[str, Any]]) -> str:
    """Identifica el contenido de la tabla tal como la ve el formulario (`base_version`, U-34.5).

    Dos tablas con las mismas celdas tienen la misma huella; cualquier cambio de celdas, filas u orden la cambia."""
    celdas = [[str(item.get(key) or "").strip() for key, _ in ACTIVITY_COLUMNS] for item in activities]
    return hashlib.sha256(json.dumps(celdas, ensure_ascii=False).encode("utf-8")).hexdigest()[:16]


def parse_activities_form(form: dict[str, list[str]]) -> list[Actividad]:
    """Reconstruye las actividades desde el formulario; ignora filas borradas o vacías.

    Rechaza un formulario sin `row_count` (llegó incompleto: guardarlo borraría la tabla), una celda de más de
    `MAX_CELL` caracteres y una fila con datos pero sin nombre de actividad (U-34.5)."""
    try:
        count = int((form.get("row_count") or [""])[0])
    except ValueError:
        raise ErrorDeValidacion("El formulario de la tabla llegó incompleto, así que no se guardó nada. Recarga el "
                                "caso y vuelve a intentarlo.") from None
    if count < 0 or count > MAX_ROWS:
        raise ErrorDeValidacion(f"La tabla admite hasta {MAX_ROWS} actividades.")
    activities = []
    for index in range(count + 1):  # +1: fila nueva al final
        if form.get(f"a{index}_del", [""])[0]:
            continue
        item: dict[str, str] = {
            key: (form.get(f"a{index}_{key}", [""])[0] or "").strip() for key, _ in ACTIVITY_COLUMNS}
        fila = "la fila nueva" if index == count else f"la fila {index + 1}"
        for key, etiqueta in ACTIVITY_COLUMNS:
            if len(item[key]) > MAX_CELL:
                raise ErrorDeValidacion(
                    f"«{etiqueta}» de {fila} tiene más de {MAX_CELL} caracteres: acórtalo para guardar.")
        if not any(item[key] for key, _ in ACTIVITY_COLUMNS if key != "estado"):
            continue
        if not item["activity"]:
            raise ErrorDeValidacion(f"{fila[0].upper() + fila[1:]} tiene datos pero no tiene nombre de actividad: "
                                    "escríbelo o quítala.")
        if item["estado"] not in ESTADOS:
            item["estado"] = "pendiente"
        activities.append(cast(Actividad, item))
    return activities


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    # evita que Excel interprete el texto como fórmula, también si va precedida de espacios o tabuladores («\t=1»)
    return "'" + text if text.lstrip(" \t\r\n")[:1] in ("=", "+", "-", "@") else text


def build_workbook(quotes: Sequence[Mapping[str, Any]]) -> bytes:
    """Excel con una hoja «Actividades»: una fila por actividad con su proyecto, cliente y estado, coloreada.

    Se escribe en modo de solo escritura (U-48): el tiempo crece en línea recta con las filas, no al cuadrado. El
    contenido y el formato esencial (encabezado en negrita sobre verde, color por estado, texto ajustado, filtro,
    primera fila y columnas fijas, anchos) son los de siempre."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {
        "confirmado": PatternFill("solid", fgColor="D9EAD3"),
        "pendiente": PatternFill("solid", fgColor="FFF2CC"),
        "contradictorio": PatternFill("solid", fgColor="F4CCCC"),
    }
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Actividades")
    widths = [14, 28, 24, 26, 34, 22, 22, 12, 12, 26, 26, 30, 14, 30]
    for position, width in enumerate(widths, 1):          # los anchos van antes de la primera fila
        sheet.column_dimensions[get_column_letter(position)].width = width
    sheet.freeze_panes = "E2"
    headers = ["Hilo", "Proyecto", "Cliente", "Estado del caso"] + [label for _, label in ACTIVITY_COLUMNS]
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="126A62")
    header_alignment = Alignment(vertical="center", wrap_text=True)
    head_cells = []
    for header in headers:
        head_cell = WriteOnlyCell(sheet, value=header)
        head_cell.font, head_cell.fill, head_cell.alignment = header_font, header_fill, header_alignment
        head_cells.append(head_cell)
    sheet.append(head_cells)
    # El estilo se resuelve una vez por tipo de celda y se copia: asignar `alignment` o `fill` celda a celda recalcula
    # y busca el estilo cada vez, y es lo que más tardaba (copiar el arreglo de estilo ya resuelto es instantáneo).
    modelo_normal: Any = WriteOnlyCell(sheet, value="")
    modelo_normal.alignment = Alignment(vertical="top", wrap_text=True)
    estilo_normal = modelo_normal._style
    estilos_de_estado = {}
    for nombre, relleno in fills.items():
        modelo: Any = WriteOnlyCell(sheet, value="")
        modelo.alignment, modelo.fill = Alignment(vertical="top", wrap_text=True), relleno
        estilos_de_estado[nombre] = modelo._style
    state_position = 4 + [key for key, _ in ACTIVITY_COLUMNS].index("estado")
    last_row = 1
    for quote in quotes:
        analysis = quote.get("analysis") or {}
        base = [
            quote.get("thread_id", ""),
            analysis.get("project_name") or "",
            quote.get("client_name") or "",
            quote.get("status_label") or quote.get("status") or "",
        ]
        for item in analysis.get("activities") or [{}]:
            row = [_cell(value) for value in base] + [
                _cell(item.get(key) or ("no indicado" if key != "estado" and item else ""))
                for key, _ in ACTIVITY_COLUMNS
            ]
            cells = []
            for position, value in enumerate(row):
                body_cell: Any = WriteOnlyCell(sheet, value=value)
                body_cell._style = copy(estilos_de_estado.get(value, estilo_normal)
                                        if position == state_position else estilo_normal)
                cells.append(body_cell)
            sheet.append(cells)
            last_row += 1
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
