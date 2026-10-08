"""El libro de cálculo de un caso (Excel / Google Sheets): la tabla vigente y, aparte, lo que no debe perderse.

Cuatro hojas:
  - «Actividades»: la tabla de ahora (es lo que se comparte o se trabaja).
  - «Solicitud inicial»: el primer correo del cliente tal como llegó y la tabla que salió del primer análisis.
  - «Historial»: cada versión de la tabla (fecha, origen y qué cambió), sin los motivos internos.
  - «Léeme»: de dónde sale el libro y por qué lo que se edite aquí se pierde en la próxima actualización.

Es una función pura: recibe los datos y devuelve los bytes del `.xlsx`. La fuente de verdad sigue siendo el dashboard.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
from typing import Any

from anka.dominio.activities import ACTIVITY_COLUMNS, normalize_estado

_ORIGEN = {"model": "Análisis automático", "previa": "Tabla anterior a la primera edición",
           "dashboard": "Edición en el dashboard"}
MAX_TEXTO = 5000            # el correo inicial se corta aquí (una celda de Sheets admite 50.000; no hace falta más)
HOJAS = ("Actividades", "Solicitud inicial", "Historial", "Léeme")


def _celda(value: object) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@") else text    # que Excel no lo interprete como fórmula


def _tabla(sheet: Any, activities: Sequence[Mapping[str, Any]], first_row: int = 1) -> int:
    """Escribe la tabla de actividades desde `first_row`; devuelve la fila siguiente a la última."""
    from openpyxl.styles import Alignment, Font, PatternFill

    colores = {"confirmado": "D9EAD3", "pendiente": "FFF2CC", "contradictorio": "F4CCCC"}
    for column, (_, label) in enumerate(ACTIVITY_COLUMNS, 1):
        cell = sheet.cell(row=first_row, column=column, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="126A62")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    row = first_row + 1
    for item in activities or []:
        for column, (key, _) in enumerate(ACTIVITY_COLUMNS, 1):
            value = item.get(key) if isinstance(item, Mapping) else ""
            cell = sheet.cell(row=row, column=column, value=_celda(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key == "estado":
                color = colores.get(normalize_estado(str(value or "")))
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)
        row += 1
    return row


def _anchos(sheet: Any, widths: Sequence[int]) -> None:
    from openpyxl.utils import get_column_letter

    for position, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(position)].width = width


def build_case_workbook(caso: Mapping[str, Any], versiones: Sequence[Mapping[str, Any]],
                        primer_mensaje: Mapping[str, Any] | None, actualizado: str) -> bytes:
    """Libro del caso: la tabla vigente, la solicitud inicial, el historial de versiones y una nota de uso."""
    import json

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.worksheet import Worksheet

    analysis = caso.get("analysis") or {}
    libro = Workbook()
    actual = libro.active
    assert isinstance(actual, Worksheet)
    actual.title = HOJAS[0]
    _tabla(actual, analysis.get("activities") or [])
    actual.freeze_panes = "A2"
    _anchos(actual, [34, 22, 22, 12, 12, 26, 26, 30, 14, 30])

    inicial = libro.create_sheet(HOJAS[1])
    negrita = Font(bold=True)
    inicial["A1"] = "Solicitud inicial del cliente (tal como llegó)"
    inicial["A1"].font = Font(bold=True, size=13)
    if primer_mensaje:
        datos = (("De", "sender"), ("Asunto", "subject"), ("Fecha", "occurred_at"))
        for fila, (etiqueta, clave) in enumerate(datos, 2):
            inicial.cell(row=fila, column=1, value=etiqueta).font = negrita
            inicial.cell(row=fila, column=2, value=_celda(primer_mensaje.get(clave)))
        cuerpo = str(primer_mensaje.get("body") or "")
        inicial["A5"] = "Mensaje"
        inicial["A5"].font = negrita
        inicial["B5"] = _celda(cuerpo[:MAX_TEXTO] + (" […]" if len(cuerpo) > MAX_TEXTO else ""))
        inicial["B5"].alignment = Alignment(vertical="top", wrap_text=True)
        inicial.row_dimensions[5].height = 240
    else:
        inicial["A2"] = "El texto del primer correo no está guardado en el dashboard."
    primera = next((v for v in versiones), None)
    inicial["A7"] = "Tabla que salió del primer análisis"
    inicial["A7"].font = Font(bold=True, size=13)
    if primera:
        _tabla(inicial, json.loads(primera.get("activities_json") or "[]"), first_row=8)
    else:
        inicial["A8"] = "Sin versiones guardadas todavía."
    _anchos(inicial, [34, 60, 22, 12, 12, 26, 26, 30, 14, 30])

    historial = libro.create_sheet(HOJAS[2])
    historial.append(["Versión", "Fecha", "Origen", "Qué cambió"])
    for cell in historial[1]:
        cell.font = negrita
    for version in versiones:
        historial.append([version.get("version"), _celda(version.get("created_at")),
                          _ORIGEN.get(str(version.get("source")), _celda(version.get("source"))),
                          _celda(version.get("summary"))])
    _anchos(historial, [10, 26, 34, 90])

    leeme = libro.create_sheet(HOJAS[3])
    for fila, texto in enumerate((
        f"Proyecto: {analysis.get('project_name') or 'Solicitud sin nombre'}",
        f"Cliente: {caso.get('client_name') or ''}",
        f"Código de la oferta: {caso.get('quote_code') or 'aún sin emitir'}",
        f"Última actualización desde el dashboard: {actualizado}",
        "",
        "Esta hoja se actualiza sola cada vez que se guarda la tabla en el dashboard de Anka.",
        "Para cambiar algo, edítalo en el dashboard: lo que se escriba directamente aquí se sobrescribe en la próxima "
        "actualización (Google conserva las versiones anteriores en Archivo > Historial de versiones).",
        "«Solicitud inicial» y «Historial» conservan cómo empezó el caso y cada cambio de la tabla.",
    ), 1):
        leeme.cell(row=fila, column=1, value=_celda(texto))
    _anchos(leeme, [120])

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
