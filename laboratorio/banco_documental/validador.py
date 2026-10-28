"""Validador estructural de la respuesta de Work en el laboratorio (T-75). No está conectado al importador de ANKA.

Separa tres capas de problemas y nunca repara nada:
- sintaxis: los bytes no son UTF-8 ni JSON válido (error exacto con línea, columna y posición);
- esquema: es JSON pero falta un campo obligatorio o un tipo no es el esperado;
- contenido: la estructura es válida pero algo no cuadra con el documento citado (fuente inexistente, valor distinto).

Una respuesta `VALIDA` o `VALIDA_CON_ADVERTENCIAS` solo está **estructuralmente aceptada**: no significa que la lectura sea
correcta. Esa verificación técnico-documental la hace la evaluación contra una referencia o una persona. Los bytes
originales se conservan siempre, tal como llegaron, y lo rechazado queda en cuarentena con su error.

Uso: python -m laboratorio.banco_documental.validador <respuesta> <carpeta_destino> [--caso C1] [--documento archivo]
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

SINTAXIS_INVALIDA = "SINTAXIS_INVALIDA"
ESQUEMA_INVALIDO = "ESQUEMA_INVALIDO"
CONTENIDO_INCONSISTENTE = "CONTENIDO_INCONSISTENTE"
VALIDA_CON_ADVERTENCIAS = "VALIDA_CON_ADVERTENCIAS"
VALIDA = "VALIDA"
ACEPTADOS = (VALIDA, VALIDA_CON_ADVERTENCIAS)

ERROR, ADVERTENCIA = "error", "advertencia"
_CLAVES_DE_ORIGEN = ("pagina", "hoja", "celda", "fila", "tabla")
_LISTAS_DE_TEXTO = ("entregables", "restricciones", "contradicciones", "preguntas_cliente", "revisiones_internas",
                    "elementos_visuales", "limitaciones_de_lectura")
_OPCIONALES = {"hojas": list, "evidencia_visual": list, "calidad_ocr": dict, "valores_historicos_u_obsoletos": list,
               "otros_datos_numericos": list}
_CASO = re.compile(r"^C[1-9]\d*$")
_CELDA = re.compile(r"^[A-Z]{1,3}[1-9]\d*(?::[A-Z]{1,3}[1-9]\d*)?$")


class ClaveDuplicada(ValueError):
    """Un objeto JSON repite una clave: el segundo valor taparía al primero sin avisar."""


@dataclass(frozen=True)
class Hallazgo:
    """Un problema o aviso: de qué capa es, dónde está y qué ocurre."""

    nivel: str
    capa: str
    ruta: str
    codigo: str
    detalle: str

    def como_dict(self) -> dict[str, str]:
        """El hallazgo como diccionario serializable."""
        return {"nivel": self.nivel, "capa": self.capa, "ruta": self.ruta, "codigo": self.codigo,
                "detalle": self.detalle}


@dataclass(frozen=True)
class Resultado:
    """El veredicto estructural sobre unos bytes; `datos` solo existe si el JSON se pudo leer."""

    estado: str
    sha256: str
    tamano: int
    hallazgos: tuple[Hallazgo, ...] = ()
    error_de_sintaxis: dict[str, Any] | None = None
    datos: dict[str, Any] | None = field(default=None, repr=False)

    @property
    def aceptada(self) -> bool:
        """Estructuralmente aceptada. NO equivale a lectura correcta ni a aprobación técnica."""
        return self.estado in ACEPTADOS

    def errores(self) -> list[Hallazgo]:
        """Los hallazgos de nivel error."""
        return [h for h in self.hallazgos if h.nivel == ERROR]

    def como_dict(self) -> dict[str, Any]:
        """El informe completo, sin el contenido de la respuesta."""
        return {"estado": self.estado, "aceptada": self.aceptada, "sha256": self.sha256, "bytes": self.tamano,
                "error_de_sintaxis": self.error_de_sintaxis, "hallazgos": [h.como_dict() for h in self.hallazgos],
                "aviso": "aceptada significa solo estructura válida; la lectura documental no está verificada"}


def sha256(contenido: bytes) -> str:
    """La huella SHA-256 en hexadecimal."""
    return hashlib.sha256(contenido).hexdigest()


def _sin_duplicados(pares: list[tuple[str, Any]]) -> dict[str, Any]:
    claves = [clave for clave, _ in pares]
    repetidas = sorted({c for c in claves if claves.count(c) > 1})
    if repetidas:
        raise ClaveDuplicada(f"clave repetida en un objeto: {', '.join(repetidas)}")
    return dict(pares)


def _constante_no_valida(nombre: str) -> Any:
    raise ValueError(f"constante no permitida en JSON: {nombre}")


def _linea_y_columna(texto: str, posicion: int) -> tuple[int, int]:
    antes = texto[:posicion]
    return antes.count("\n") + 1, posicion - (antes.rfind("\n") + 1) + 1


def leer(contenido: bytes) -> tuple[Any, dict[str, Any] | None]:
    """Lee los bytes como UTF-8 y JSON, sin reparar nada: devuelve (valor, None) o (None, error exacto)."""
    try:
        texto = contenido.decode("utf-8")
    except UnicodeDecodeError as error:
        return None, {"tipo": "utf8", "mensaje": str(error), "posicion_en_bytes": error.start}
    if texto.startswith("﻿"):
        return None, {"tipo": "bom", "mensaje": "el archivo empieza con una marca BOM; no es JSON estricto",
                      "posicion": 0}
    try:
        return json.loads(texto, object_pairs_hook=_sin_duplicados, parse_constant=_constante_no_valida), None
    except json.JSONDecodeError as error:
        linea = texto.splitlines()[error.lineno - 1] if texto.splitlines() else ""
        return None, {"tipo": "json", "mensaje": error.msg, "linea": error.lineno, "columna": error.colno,
                      "posicion": error.pos, "texto_de_la_linea": linea[:200],
                      "explicacion": f"{error.msg}: line {error.lineno} column {error.colno} (char {error.pos})"}
    except ClaveDuplicada as error:
        return None, {"tipo": "clave_duplicada", "mensaje": str(error)}
    except ValueError as error:
        return None, {"tipo": "constante", "mensaje": str(error)}
    except RecursionError:
        return None, {"tipo": "profundidad", "mensaje": "anidación excesiva"}


def _es_entero(valor: Any) -> bool:
    return isinstance(valor, int) and not isinstance(valor, bool)


def _texto_no_vacio(valor: Any) -> bool:
    return isinstance(valor, str) and valor.strip() != ""


def _validar_actividad(indice: int, actividad: Any, hallazgos: list[Hallazgo]) -> None:
    ruta = f"actividades[{indice}]"
    if not isinstance(actividad, dict):
        hallazgos.append(Hallazgo(ERROR, "esquema", ruta, "tipo_incorrecto", "cada actividad debe ser un objeto"))
        return
    for clave in ("nombre", "cantidad", "unidad"):
        valor = actividad.get(clave)
        if clave not in actividad:
            hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.{clave}", "campo_obligatorio_ausente",
                                      f"falta «{clave}»"))
        elif not _texto_no_vacio(valor):
            motivo = ("la cantidad debe escribirse como texto literal (no como número) para no perder separadores ni "
                      "ceros" if clave == "cantidad" and isinstance(valor, int | float) and not isinstance(valor, bool)
                      else f"«{clave}» debe ser un texto no vacío")
            hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.{clave}", "tipo_incorrecto", motivo))
    if "entregable" in actividad and not isinstance(actividad["entregable"], str):
        hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.entregable", "tipo_incorrecto", "debe ser texto"))
    origen = actividad.get("origen")
    if origen is None:
        hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.origen", "campo_obligatorio_ausente",
                                  "toda actividad cita su origen (página, hoja o celda)"))
        return
    if not isinstance(origen, dict) or not any(clave in origen for clave in _CLAVES_DE_ORIGEN):
        hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.origen", "origen_sin_localizador",
                                  f"debe ser un objeto con al menos una de {', '.join(_CLAVES_DE_ORIGEN)}"))
        return
    if "pagina" in origen and (not _es_entero(origen["pagina"]) or origen["pagina"] < 1):
        hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.origen.pagina", "tipo_incorrecto",
                                  "la página es un entero desde 1"))
    for clave in ("hoja", "celda"):
        if clave in origen and not _texto_no_vacio(origen[clave]):
            hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.origen.{clave}", "tipo_incorrecto",
                                      f"«{clave}» debe ser texto"))
    for clave in ("tabla", "fila"):
        if clave in origen and not (_texto_no_vacio(origen[clave]) or _es_entero(origen[clave])):
            hallazgos.append(Hallazgo(ERROR, "esquema", f"{ruta}.origen.{clave}", "tipo_incorrecto",
                                      f"«{clave}» debe ser texto o entero"))


def _validar_caso(datos: dict[str, Any], caso_esperado: str | None) -> list[Hallazgo]:
    caso = datos.get("caso")
    if "caso" not in datos:
        return [Hallazgo(ERROR, "esquema", "caso", "campo_obligatorio_ausente", "falta «caso»")]
    if not isinstance(caso, str) or not _CASO.match(caso):
        return [Hallazgo(ERROR, "esquema", "caso", "tipo_incorrecto", "«caso» debe ser un texto como C1")]
    if caso_esperado is not None and caso != caso_esperado:
        return [Hallazgo(ERROR, "esquema", "caso", "caso_distinto", f"la respuesta es del caso {caso}, no de {caso_esperado}")]
    return []


def validar_esquema(datos: Any, caso_esperado: str | None = None) -> list[Hallazgo]:
    """Comprueba campos obligatorios y tipos; lo opcional con tipo distinto solo advierte."""
    if not isinstance(datos, dict):
        return [Hallazgo(ERROR, "esquema", "$", "tipo_incorrecto", "la respuesta debe ser un objeto JSON")]
    hallazgos = _validar_caso(datos, caso_esperado)
    if "actividades" not in datos:
        hallazgos.append(Hallazgo(ERROR, "esquema", "actividades", "campo_obligatorio_ausente", "falta «actividades»"))
    elif not isinstance(datos["actividades"], list):
        hallazgos.append(Hallazgo(ERROR, "esquema", "actividades", "tipo_incorrecto", "debe ser una lista"))
    else:
        for indice, actividad in enumerate(datos["actividades"]):
            _validar_actividad(indice, actividad, hallazgos)
    for clave in _LISTAS_DE_TEXTO:
        if clave not in datos:
            hallazgos.append(Hallazgo(ERROR, "esquema", clave, "campo_obligatorio_ausente", f"falta «{clave}»"))
        elif not isinstance(datos[clave], list):
            hallazgos.append(Hallazgo(ERROR, "esquema", clave, "tipo_incorrecto", "debe ser una lista de textos"))
        else:
            for indice, elemento in enumerate(datos[clave]):
                if not isinstance(elemento, str):
                    hallazgos.append(Hallazgo(ERROR, "esquema", f"{clave}[{indice}]", "tipo_incorrecto",
                                              f"debe ser texto, no {type(elemento).__name__}: el comparador lee texto"))
    for clave, tipo in _OPCIONALES.items():
        if clave in datos and not isinstance(datos[clave], tipo):
            hallazgos.append(Hallazgo(ADVERTENCIA, "esquema", clave, "tipo_inesperado_opcional",
                                      f"se esperaba {tipo.__name__}, llegó {type(datos[clave]).__name__}; no se evalúa"))
    conocidas = {"caso", "actividades", *_LISTAS_DE_TEXTO, *_OPCIONALES}
    for clave in sorted(set(datos) - conocidas):
        hallazgos.append(Hallazgo(ADVERTENCIA, "esquema", clave, "campo_desconocido", "no forma parte del formato"))
    return hallazgos


def _normal(texto: str) -> str:
    sin_tildes = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", sin_tildes.lower()).strip()


def _decimal(texto: str) -> Decimal | None:
    limpio = texto.strip().replace(" ", "")
    if re.fullmatch(r"\d+(?:[.,]\d+)?", limpio) is None:
        return None
    try:
        return Decimal(limpio.replace(",", "."))
    except InvalidOperation:
        return None


def validar_contenido(datos: dict[str, Any]) -> list[Hallazgo]:
    """Consistencia interna que no necesita el documento: duplicados y señales de lectura incompleta."""
    hallazgos: list[Hallazgo] = []
    vistas: dict[tuple[str, str, str], int] = {}
    for indice, actividad in enumerate(datos.get("actividades", [])):
        if not isinstance(actividad, dict):
            continue
        clave = (_normal(str(actividad.get("nombre", ""))), _normal(str(actividad.get("cantidad", ""))),
                 _normal(str(actividad.get("unidad", ""))))
        if clave in vistas:
            hallazgos.append(Hallazgo(ADVERTENCIA, "contenido", f"actividades[{indice}]", "actividad_duplicada",
                                      f"repite la actividad {vistas[clave]} con la misma cantidad y unidad"))
        vistas.setdefault(clave, indice)
    if not datos.get("actividades"):
        hallazgos.append(Hallazgo(ADVERTENCIA, "contenido", "actividades", "sin_actividades",
                                  "no hay actividades: confirmar que el documento no las traía"))
    if datos.get("limitaciones_de_lectura") == []:
        hallazgos.append(Hallazgo(ADVERTENCIA, "contenido", "limitaciones_de_lectura", "limitaciones_vacias",
                                  "una lista vacía no prueba lectura completa"))
    return hallazgos


def verificar_fuentes(datos: dict[str, Any], documento: bytes, nombre_documento: str) -> list[Hallazgo]:
    """Contrasta cada origen citado con el DOCUMENTO (no con una referencia): existencia y, en Excel, valor guardado."""
    hallazgos: list[Hallazgo] = []
    actividades = [a for a in datos.get("actividades", []) if isinstance(a, dict) and isinstance(a.get("origen"), dict)]
    nombre = nombre_documento.lower()
    if nombre.endswith(".xlsx"):
        from openpyxl import load_workbook

        valores = load_workbook(io.BytesIO(documento), data_only=True)
        for indice, actividad in enumerate(datos.get("actividades", [])):
            origen = actividad.get("origen") if isinstance(actividad, dict) else None
            if not isinstance(origen, dict) or "hoja" not in origen or "celda" not in origen:
                continue
            ruta = f"actividades[{indice}].origen"
            hoja, celda = str(origen["hoja"]), str(origen["celda"]).replace("$", "").upper()
            if hoja not in valores.sheetnames:
                hallazgos.append(Hallazgo(ERROR, "contenido", ruta, "fuente_inexistente", f"no existe la hoja «{hoja}»"))
            elif not _CELDA.match(celda) or ":" in celda:
                hallazgos.append(Hallazgo(ERROR, "contenido", ruta, "fuente_inexistente",
                                          f"«{celda}» no es una celda única válida"))
            else:
                guardado = valores[hoja][celda].value
                if guardado is None:
                    hallazgos.append(Hallazgo(ERROR, "contenido", ruta, "fuente_inexistente",
                                              f"la celda {hoja}!{celda} está vacía"))
                else:
                    citada, real = _decimal(str(actividad.get("cantidad", ""))), _decimal(str(guardado))
                    if citada is not None and real is not None and citada != real:
                        hallazgos.append(Hallazgo(ERROR, "contenido", ruta, "valor_no_coincide_con_la_fuente",
                                                  f"cita {citada} pero {hoja}!{celda} guarda {guardado}"))
    elif nombre.endswith(".pdf"):
        from pypdf import PdfReader

        paginas = len(PdfReader(io.BytesIO(documento)).pages)
        for indice, actividad in enumerate(actividades):
            pagina = actividad["origen"].get("pagina")
            if _es_entero(pagina) and pagina > paginas:
                hallazgos.append(Hallazgo(ERROR, "contenido", f"actividades[{indice}].origen.pagina",
                                          "fuente_inexistente", f"el documento tiene {paginas} páginas, no {pagina}"))
    return hallazgos


def validar(contenido: bytes, caso_esperado: str | None = None, documento: bytes | None = None,
            nombre_documento: str = "") -> Resultado:
    """El veredicto estructural de unos bytes. Nunca repara: lo que no es JSON válido queda como SINTAXIS_INVALIDA."""
    huella = sha256(contenido)
    datos, error = leer(contenido)
    if error is not None:
        detalle = error.get("explicacion") or error["mensaje"]
        return Resultado(SINTAXIS_INVALIDA, huella, len(contenido),
                         (Hallazgo(ERROR, "sintaxis", "$", error["tipo"], detalle),), error)
    hallazgos = validar_esquema(datos, caso_esperado)
    if any(h.nivel == ERROR for h in hallazgos):
        return Resultado(ESQUEMA_INVALIDO, huella, len(contenido), tuple(hallazgos))
    assert isinstance(datos, dict)
    hallazgos += validar_contenido(datos)
    if documento is not None:
        hallazgos += verificar_fuentes(datos, documento, nombre_documento)
    if any(h.nivel == ERROR for h in hallazgos):
        return Resultado(CONTENIDO_INCONSISTENTE, huella, len(contenido), tuple(hallazgos))
    estado = VALIDA_CON_ADVERTENCIAS if hallazgos else VALIDA
    return Resultado(estado, huella, len(contenido), tuple(hallazgos), None, datos)


def _escribir_una_vez(ruta: Path, contenido: bytes) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if ruta.exists():
        if ruta.read_bytes() != contenido:
            raise FileExistsError(f"{ruta} ya existe con otro contenido")
        return
    ruta.write_bytes(contenido)


def procesar(contenido: bytes, destino: Path, caso_esperado: str | None = None, documento: bytes | None = None,
             nombre_documento: str = "") -> Resultado:
    """Valida y archiva: `originales/` conserva SIEMPRE los bytes, `informes/` el veredicto, y solo lo aceptado llega a
    `validadas/` (copia exacta); lo demás va a `cuarentena/` con su motivo. Un consumidor lee únicamente `validadas/`."""
    resultado = validar(contenido, caso_esperado, documento, nombre_documento)
    _escribir_una_vez(destino / "originales" / f"{resultado.sha256}.raw", contenido)
    informe = (json.dumps(resultado.como_dict(), ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")
    carpeta = "validadas" if resultado.aceptada else "cuarentena"
    _escribir_una_vez(destino / "informes" / f"{resultado.sha256}.json", informe)
    if resultado.aceptada:
        _escribir_una_vez(destino / carpeta / f"{resultado.sha256}.json", contenido)
    else:
        _escribir_una_vez(destino / carpeta / f"{resultado.sha256}.motivo.json", informe)
    return resultado


def main(argumentos: list[str]) -> int:
    """Valida un archivo de respuesta y lo archiva; código de salida 0 solo si quedó estructuralmente aceptado."""
    opciones: dict[str, str | None] = {"--caso": None, "--documento": None}
    posicionales: list[str] = []
    iterador = iter(argumentos)
    for argumento in iterador:
        if argumento in opciones:
            opciones[argumento] = next(iterador, None)
        else:
            posicionales.append(argumento)
    if len(posicionales) != 2:
        print(__doc__)
        return 2
    documento = Path(opciones["--documento"]) if opciones["--documento"] else None
    resultado = procesar(Path(posicionales[0]).read_bytes(), Path(posicionales[1]), opciones["--caso"],
                         documento.read_bytes() if documento else None, documento.name if documento else "")
    print(json.dumps(resultado.como_dict(), ensure_ascii=False, indent=2))
    return 0 if resultado.aceptada else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
