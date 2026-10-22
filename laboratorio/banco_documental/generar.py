"""Genera los tres documentos ficticios del banco a partir de `datos.py` (T-71): mismos datos, mismos bytes.

El PDF digital se escribe a mano (texto con fuente estándar), el escaneado es una imagen JPEG con ruido de semilla fija y
el Excel lleva fórmulas con su valor ya calculado. Para el escaneado hace falta una fuente TrueType con tildes; se busca
en el equipo y su huella queda en el manifiesto, porque otra fuente cambiaría los píxeles.
"""

from __future__ import annotations

import hashlib
import io
import random
import re
import textwrap
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import cast

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont

from laboratorio.banco_documental import datos
from laboratorio.banco_documental.pdf_minimo import Pagina, Pdf

SEMILLA = 20261112
CALIDAD_JPEG = 72
DPI = 150
ANCHO_PX, ALTO_PX = 1240, 1754
FECHA_FIJA = datetime(2026, 11, 12, 8, 0, 0)
_FUENTES = {"normal": ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"),
            "negrita": ("arialbd.ttf", "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"),
            "cursiva": ("ariali.ttf", "DejaVuSans-Oblique.ttf", "LiberationSans-Italic.ttf")}
_CARPETAS_DE_FUENTES = (Path("C:/Windows/Fonts"), Path("/usr/share/fonts/truetype/dejavu"),
                        Path("/usr/share/fonts/truetype/liberation"), Path("/Library/Fonts"))


class SinFuente(RuntimeError):
    """No hay una fuente TrueType con tildes en el equipo: el escaneado no se puede generar fielmente."""


def ruta_de_fuente(estilo: str) -> Path:
    """La primera fuente TrueType del estilo que exista en el equipo."""
    for carpeta in _CARPETAS_DE_FUENTES:
        for nombre in _FUENTES[estilo]:
            if (carpeta / nombre).is_file():
                return carpeta / nombre
    raise SinFuente(f"No se encontró una fuente TrueType ({', '.join(_FUENTES[estilo])}) para el escaneado.")


def huella_de_fuentes() -> dict[str, str]:
    """El SHA-256 de cada fuente usada en el escaneado, para saber si otra máquina lo reproduciría igual."""
    return {estilo: hashlib.sha256(ruta_de_fuente(estilo).read_bytes()).hexdigest() for estilo in _FUENTES}


# ---------------------------------------------------------------------------------------------- C1: PDF digital
def _envolver(texto: str, tam: float, ancho: float) -> list[str]:
    return textwrap.wrap(texto, width=max(10, int(ancho / (tam * 0.5))))


def _parrafo(pagina: Pagina, x: float, y: float, texto: str, ancho: float = 495, tam: float = 9.5,
             negrita: bool = False, interlineado: float = 13) -> float:
    for linea in _envolver(texto, tam, ancho):
        pagina.texto(x, y, linea, tam, negrita)
        y += interlineado
    return y


def generar_c1() -> bytes:
    """El PDF digital de la solicitud de levantamiento topográfico (dos páginas, Tabla 1 en la primera)."""
    caso = datos.C1
    p = caso["parametros"]
    pdf = Pdf("Solicitud de cotizacion")
    uno = pdf.nueva_pagina()
    uno.texto(50, 60, "SOLICITUD DE COTIZACIÓN – SERVICIOS DE TOPOGRAFÍA", 14, True)
    uno.texto(50, 80, p["entidad"], 10)
    uno.texto(50, 94, f"Referencia: {p['codigo']}   |   Fecha: {p['fecha']}", 9)
    uno.linea(50, 102, 545, 102, 1)
    uno.texto(50, 122, "1. OBJETO", 10.5, True)
    y = _parrafo(uno, 50, 138, f"Se solicita cotización para el levantamiento topográfico del predio «{p['predio']}», "
                                f"ubicado en la {p['ubicacion']}, para el diseño del {p['proyecto']}.")
    uno.texto(50, y + 12, "2. DATOS DEL PREDIO", 10.5, True)
    y = _parrafo(uno, 50, y + 28, f"Área total del predio según escritura: {p['area_predio']} ha.")
    y = _parrafo(uno, 50, y, "El polígono del predio se enviará por separado en archivo KMZ.")
    uno.texto(50, y + 12, "3. ACTIVIDADES REQUERIDAS (Tabla 1)", 10.5, True)
    columnas = (50, 80, 300, 350, 425, 545)
    y += 30
    uno.rectangulo(columnas[0], y, columnas[-1] - columnas[0], 18, 0.88)
    for x, rotulo in zip(columnas, ("Ítem", "Actividad", "Cantidad", "Unidad", "Entregable"), strict=False):
        uno.texto(x + 3, y + 12, rotulo, 8.5, True)
    y += 18
    for indice, actividad in enumerate(caso["actividades"], 1):
        celdas = (str(indice), actividad["nombre"], actividad["cantidad_literal"], actividad["unidad_literal"],
                  actividad["entregable"])
        lineas = [_envolver(celda, 8, columnas[i + 1] - columnas[i] - 6) for i, celda in enumerate(celdas)]
        alto = 10 * max(len(trozo) for trozo in lineas) + 8
        for i, trozo in enumerate(lineas):
            for k, linea in enumerate(trozo):
                uno.texto(columnas[i] + 3, y + 11 + 10 * k, linea, 8)
        y += alto
        uno.linea(columnas[0], y, columnas[-1], y, 0.4)
    uno.texto(50, 800, "Página 1 de 2", 8)

    dos = pdf.nueva_pagina()
    dos.texto(50, 60, "4. CONDICIONES", 10.5, True)
    condiciones = (
        f"a) Plazo de entrega: {p['plazo_dias']} días calendario contados desde la orden de inicio.",
        f"b) Reunión de inicio programada para el {p['reunion']}.",
        "c) Acceso al predio de lunes a viernes entre 7:00 y 16:00; sábados solo con autorización previa.",
        f"d) Precisión horizontal requerida: {p['precision']} (posicionamiento GNSS).",
        f"e) Se estima un total de {p['puntos_en_texto']} puntos de control para el proyecto; el paquete de reseñas debe "
        f"incluir {p['fotos_por_punto']} fotografías por punto de control.",
        "f) El predio tiene presencia de ganado; se exige afiliación vigente a ARL y uso de EPP.",
    )
    y = 78
    for condicion in condiciones:
        y = _parrafo(dos, 50, y, condicion) + 3
    dos.texto(50, y + 14, "5. NOTAS", 10.5, True)
    y = _parrafo(dos, 50, y + 30, "Nota 6. El sistema de coordenadas se definirá en la reunión de inicio.")
    y = _parrafo(dos, 50, y + 3, f"{p['nota_siete']}. Las construcciones del sector sur se levantarán solo por fachada "
                                 "principal.")
    dos.texto(50, y + 14, "6. CONTACTO", 10.5, True)
    _parrafo(dos, 50, y + 30, p["contacto"])
    dos.texto(50, 800, "Página 2 de 2", 8)
    return pdf.serializar()


# ------------------------------------------------------------------------------------------- C2: PDF escaneado
def _fuente(estilo: str, tam: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(ruta_de_fuente(estilo)), tam)


def _texto_envuelto(d: ImageDraw.ImageDraw, x: int, y: int, texto: str, fuente: ImageFont.FreeTypeFont, ancho: int,
                    interlineado: int) -> int:
    linea = ""
    for palabra in texto.split():
        prueba = f"{linea} {palabra}".strip()
        if d.textlength(prueba, font=fuente) > ancho and linea:
            d.text((x, y), linea, font=fuente, fill=20)
            y += interlineado
            linea = palabra
        else:
            linea = prueba
    if linea:
        d.text((x, y), linea, font=fuente, fill=20)
        y += interlineado
    return y


def _a_mano(lienzo: Image.Image, azar: random.Random, xy: tuple[int, int], texto: str, tam: int, grados: float) -> None:
    """Escribe `texto` como una anotación a mano: caracteres desalineados, inclinada y con trazo desigual."""
    fuente = _fuente("cursiva", tam)
    capa = Image.new("L", (len(texto) * tam + 40, tam * 2), 255)
    d = ImageDraw.Draw(capa)
    x = 10
    for caracter in texto:
        d.text((x, 12 + azar.randint(-4, 4)), caracter, font=fuente, fill=azar.randint(40, 90))
        x += int(d.textlength(caracter, font=fuente)) + azar.randint(0, 3)
    girada = capa.rotate(grados, expand=True, fillcolor=255, resample=Image.Resampling.BICUBIC)
    lienzo.paste(ImageChops.darker(lienzo.crop((xy[0], xy[1], xy[0] + girada.width, xy[1] + girada.height)), girada), xy)


def _sello(lienzo: Image.Image, xy: tuple[int, int], texto: str) -> None:
    fuente = _fuente("negrita", 30)
    capa = Image.new("L", (520, 90), 255)
    d = ImageDraw.Draw(capa)
    d.rectangle((2, 2, 517, 87), outline=70, width=4)
    d.text((22, 24), texto, font=fuente, fill=70)
    girada = capa.rotate(7, expand=True, fillcolor=255, resample=Image.Resampling.BICUBIC)
    lienzo.paste(ImageChops.darker(lienzo.crop((xy[0], xy[1], xy[0] + girada.width, xy[1] + girada.height)), girada), xy)


def _cuadro(d: ImageDraw.ImageDraw, x0: int, y: int, columnas: tuple[int, ...], filas: list[tuple[str, ...]],
            altos: list[int], tam: int = 25) -> list[int]:
    """Dibuja una tabla con rejilla; devuelve el y superior de cada fila (la primera es el encabezado)."""
    normal, negrita = _fuente("normal", tam), _fuente("negrita", tam)
    arriba, y_actual = [], y
    for numero, (fila, alto) in enumerate(zip(filas, altos, strict=True)):
        arriba.append(y_actual)
        if numero == 0:
            d.rectangle((x0, y_actual, columnas[-1], y_actual + alto), fill=225)
        for i, celda in enumerate(fila):
            ancho = columnas[i + 1] - columnas[i] - 20 if i + 1 < len(columnas) else 200
            if i + 1 >= len(columnas):
                break
            _texto_envuelto(d, columnas[i] + 10, y_actual + 12, celda, negrita if numero == 0 else normal, ancho, tam + 6)
        y_actual += alto
    for y_linea in (*arriba, y_actual):
        d.line((x0, y_linea, columnas[-1], y_linea), fill=40, width=2)
    for x in columnas:
        d.line((x, arriba[0], x, y_actual), fill=40, width=2)
    return arriba


def _pagina_1_escaneada(azar: random.Random) -> Image.Image:
    p = datos.C2["parametros"]
    lienzo = Image.new("L", (ANCHO_PX, ALTO_PX), 255)
    d = ImageDraw.Draw(lienzo)
    d.text((100, 110), "TÉRMINOS DE REFERENCIA", font=_fuente("negrita", 44), fill=15)
    d.text((100, 175), "Levantamiento batimétrico – " + p["embalse"], font=_fuente("negrita", 30), fill=25)
    d.text((100, 225), p["entidad"], font=_fuente("normal", 26), fill=25)
    d.text((100, 262), f"TDR No. {p['numero']}   ·   Folio {p['folio']}", font=_fuente("normal", 26), fill=25)
    d.line((100, 312, 1140, 312), fill=30, width=3)
    _sello(lienzo, (600, 1420), f"COPIA CONTROLADA No. {p['copia']}")
    d.text((100, 340), "1. OBJETO", font=_fuente("negrita", 28), fill=20)
    y = _texto_envuelto(d, 100, 385, "Contratar el levantamiento batimétrico del " + p["embalse"] + " y la topografía de su "
                        "franja de ronda, conforme al alcance del Cuadro 1.", _fuente("normal", 26), 1040, 36)
    d.text((100, y + 24), "Cuadro 1. Alcance", font=_fuente("negrita", 28), fill=20)
    filas: list[tuple[str, ...]] = [("Código", "Actividad", "Magnitud", "Unidad")]
    for letra, actividad in zip("ABCD", datos.C2["actividades"], strict=True):
        filas.append((letra, actividad["nombre"], actividad["cantidad_literal"], actividad["unidad_literal"]))
    arriba = _cuadro(d, 100, y + 72, (100, 200, 800, 960, 1140), filas, [56, 130, 56, 56, 110])
    # la cifra impresa de la fila A queda tachada y se corrige a mano
    celda_y = arriba[1]
    d.line((812, celda_y + 20, 925, celda_y + 52), fill=30, width=4)
    _a_mano(lienzo, azar, (805, celda_y + 62), p["manuscrito"] + " ha", 34, 4)
    y_final = arriba[-1] + 110
    d.text((100, y_final + 24), "La franja de ronda se mide " + p["ronda"] + " a cada lado del eje del embalse.",
           font=_fuente("normal", 26), fill=20)
    d.text((100, 1650), "Página 1 de 2", font=_fuente("normal", 22), fill=60)
    return lienzo


def _hachurado(lienzo: Image.Image, poligono: list[tuple[int, int]]) -> None:
    mascara = Image.new("L", lienzo.size, 0)
    ImageDraw.Draw(mascara).polygon(poligono, fill=255)
    patron = Image.new("L", lienzo.size, 255)
    dp = ImageDraw.Draw(patron)
    for desplazamiento in range(-lienzo.height, lienzo.width, 14):
        dp.line((desplazamiento, lienzo.height, desplazamiento + lienzo.height, 0), fill=60, width=2)
    lienzo.paste(patron, (0, 0), mascara)
    ImageDraw.Draw(lienzo).polygon(poligono, outline=20, width=3)


def _pagina_2_escaneada() -> Image.Image:
    p = datos.C2["parametros"]
    lienzo = Image.new("L", (ANCHO_PX, ALTO_PX), 255)
    d = ImageDraw.Draw(lienzo)
    d.text((100, 110), "Figura 1. Croquis del embalse (sin escala)", font=_fuente("negrita", 30), fill=20)
    embalse = [(180, 520), (300, 430), (520, 400), (700, 330), (900, 380), (1020, 520), (960, 700), (800, 820),
               (640, 780), (520, 900), (340, 860), (220, 720)]
    d.polygon(embalse, fill=215, outline=20, width=4)
    _hachurado(lienzo, [(700, 330), (900, 380), (960, 470), (800, 520), (690, 440)])
    d.line((830, 440, 1080, 300), fill=20, width=3)
    d.text((960, 255), "Zona de captación", font=_fuente("normal", 26), fill=20)
    d.text((960, 285), p["excluida"] + " ha", font=_fuente("normal", 26), fill=20)
    # leyenda
    d.rectangle((100, 985, 1140, 1115), outline=30, width=2)
    d.rectangle((130, 1010, 190, 1050), fill=215, outline=20, width=3)
    d.text((215, 1012), "Área a levantar", font=_fuente("normal", 26), fill=20)
    _hachurado(lienzo, [(130, 1060), (190, 1060), (190, 1100), (130, 1100)])
    d.text((215, 1062), "Excluida del levantamiento", font=_fuente("normal", 26), fill=20)
    # flecha norte y escala
    d.line((1060, 700, 1060, 600), fill=20, width=5)
    d.polygon([(1060, 575), (1043, 615), (1077, 615)], fill=20)
    d.text((1050, 710), "N", font=_fuente("negrita", 28), fill=20)
    d.line((140, 920, 540, 920), fill=20, width=4)
    for x, rotulo in ((140, "0"), (340, "250"), (540, "500 m")):
        d.line((x, 905, x, 935), fill=20, width=3)
        d.text((x - 10, 940), rotulo, font=_fuente("normal", 22), fill=20)
    # cuadro 2 y notas
    d.text((100, 1150), "Cuadro 2. Áreas", font=_fuente("negrita", 28), fill=20)
    _cuadro(d, 100, 1195, (100, 760, 1000), [("Concepto", "Área (ha)"), ("Espejo de agua total", p["espejo"]),
                                              ("Zona de captación (ver Figura 1)", p["excluida"])], [56, 56, 56])
    notas = (f"1. La rampa de acceso para la embarcación tiene {p['rampa']} m de ancho.",
             f"2. Calado máximo de la embarcación: {p['calado']} m.",
             "3. El permiso de navegación está a cargo del contratista.",
             "4. Profundidad máxima esperada: 14 m.",
             "5. Ubicación del embalse: ver plano del Anexo 3.")
    for i, nota in enumerate(notas):
        d.text((100, 1395 + i * 40), nota, font=_fuente("normal", 25), fill=20)
    d.text((100, 1650), "Página 2 de 2", font=_fuente("normal", 22), fill=60)
    return lienzo


def _como_escaneo(pagina: Image.Image, azar: random.Random, grados: float) -> bytes:
    """Aplica inclinación leve, desenfoque y ruido de semilla fija, y devuelve el JPEG resultante."""
    inclinada = pagina.rotate(grados, resample=Image.Resampling.BICUBIC, fillcolor=250)
    suave = inclinada.filter(ImageFilter.GaussianBlur(0.7))
    ruido = Image.frombytes("L", pagina.size, azar.randbytes(pagina.width * pagina.height))
    mezcla = Image.blend(suave, ruido, 0.07)
    salida = io.BytesIO()
    mezcla.save(salida, format="JPEG", quality=CALIDAD_JPEG, optimize=False, progressive=False)
    return salida.getvalue()


def generar_c2() -> bytes:
    """El PDF escaneado: dos páginas que son solo una imagen, sin capa de texto."""
    azar = random.Random(SEMILLA)
    pdf = Pdf("Documento escaneado")
    for pagina, grados in ((_pagina_1_escaneada(azar), 0.35), (_pagina_2_escaneada(), -0.3)):
        pdf.pagina_de_imagen(_como_escaneo(pagina, azar, grados), ANCHO_PX, ALTO_PX)
    return pdf.serializar()


# ------------------------------------------------------------------------------------------------- C3: Excel
def _decimal_en_texto(valor: Decimal) -> str:
    texto = format(valor.normalize(), "f")
    return texto


def valores_calculados() -> dict[tuple[str, str], str]:
    """El resultado de cada fórmula del libro, calculado aquí con aritmética decimal exacta."""
    hectareas = {f"D{fila}": Decimal(area) / 10000 for fila, (_, _, area) in enumerate(datos.DETALLE_LOTES, 2)}
    total_m2 = sum((Decimal(area) for _, _, area in datos.DETALLE_LOTES), Decimal(0))
    total_ha = sum(hectareas.values(), Decimal(0))
    metros = sum((Decimal(longitud) for _, longitud in datos.DETALLE_VIAS), Decimal(0))
    calculados = {("Detalle", celda): _decimal_en_texto(valor) for celda, valor in hectareas.items()}
    calculados[("Detalle", "C8")] = _decimal_en_texto(total_m2)
    calculados[("Detalle", "D8")] = _decimal_en_texto(total_ha)
    calculados[("Resumen", "C4")] = _decimal_en_texto(total_ha)
    calculados[("Resumen", "C5")] = _decimal_en_texto(metros / 1000)
    calculados[("Resumen", "C6")] = _decimal_en_texto(Decimal(datos.SUPUESTOS[0][1]) * len(datos.DETALLE_LOTES))
    calculados[("Resumen", "C7")] = _decimal_en_texto(Decimal(len(datos.DETALLE_VIAS)))
    return calculados


def _libro_con_formulas() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.worksheet.worksheet import Worksheet

    libro = Workbook()
    libro.properties.creator = libro.properties.lastModifiedBy = "Autor Ejemplo"
    libro.properties.created = libro.properties.modified = FECHA_FIJA
    resumen = cast(Worksheet, libro.active)
    resumen.title = "Resumen"
    detalle = libro.create_sheet("Detalle")
    supuestos = libro.create_sheet("Supuestos")
    anterior = libro.create_sheet("Borrador_anterior")
    anterior.sheet_state = "hidden"

    resumen["A1"] = "Resumen de cantidades – " + datos.C3["parametros"]["proyecto"]
    for columna, rotulo in zip("ABCDE", ("Código", "Actividad", "Cantidad", "Unidad", "Observación"), strict=True):
        resumen[f"{columna}3"] = rotulo
    filas = (("R-01", "Levantamiento topográfico de lotes", "=Detalle!D8", "ha", "Suma de los lotes del Detalle"),
             ("R-02", "Levantamiento de ejes viales", "=SUM(Detalle!G2:G4)/1000", "km", "Longitud de tramos en metros / 1000"),
             ("R-03", "Puntos de control", "=Supuestos!B2*COUNTA(Detalle!A2:A7)", "puntos", "Puntos por lote × lotes"),
             ("R-04", "Perfiles de eje", "=COUNTA(Detalle!F2:F4)", "perfiles", "Un perfil por tramo"))
    for fila, valores in enumerate(filas, 4):
        for columna, valor in zip("ABCDE", valores, strict=True):
            resumen[f"{columna}{fila}"] = valor
    for celda, formato in (("C4", "0.0"), ("C5", "0.000"), ("C6", "0"), ("C7", "0")):
        resumen[celda].number_format = formato
    resumen["A9"], resumen["B9"], resumen["C9"], resumen["D9"] = "R-05", "Predios intervenidos", datos.PREDIOS_EN_RESUMEN, "lotes"
    resumen["E9"] = "Según la solicitud del cliente"
    resumen["C9"].comment = Comment("Pendiente de confirmar con el cliente.", "Autor Ejemplo")

    for columna, rotulo in zip("ABCDFG", ("Lote", "Vereda", "Área (m2)", "Área (ha)", "Tramo", "Longitud (m)"), strict=True):
        detalle[f"{columna}1"] = rotulo
    for fila, (lote, vereda, area) in enumerate(datos.DETALLE_LOTES, 2):
        detalle[f"A{fila}"], detalle[f"B{fila}"], detalle[f"C{fila}"] = lote, vereda, area
        detalle[f"D{fila}"] = f"=C{fila}/10000"
        detalle[f"D{fila}"].number_format = "0.000"
    detalle["A8"], detalle["C8"], detalle["D8"] = "Total", "=SUM(C2:C7)", "=SUM(D2:D7)"
    detalle["D8"].number_format = "0.000"
    for fila, (tramo, longitud) in enumerate(datos.DETALLE_VIAS, 2):
        detalle[f"F{fila}"], detalle[f"G{fila}"] = tramo, longitud

    for columna, rotulo in zip("ABC", ("Parámetro", "Valor", "Unidad"), strict=True):
        supuestos[f"{columna}1"] = rotulo
    for fila, (nombre, parametro, unidad) in enumerate(datos.SUPUESTOS, 2):
        supuestos[f"A{fila}"], supuestos[f"B{fila}"], supuestos[f"C{fila}"] = nombre, parametro, unidad

    for columna, rotulo in zip("ABCD", ("Código", "Actividad", "Cantidad", "Unidad"), strict=True):
        anterior[f"{columna}1"] = rotulo
    for fila, anteriores in enumerate(datos.BORRADOR_ANTERIOR, 2):
        for columna, previo in zip("ABCD", anteriores, strict=True):
            anterior[f"{columna}{fila}"] = previo
    for hoja, ancho in ((resumen, (8, 40, 12, 10, 36)), (detalle, (10, 16, 14, 12, 4, 10, 14)),
                        (supuestos, (30, 10, 10)), (anterior, (8, 40, 12, 10))):
        for columna, medida in zip("ABCDEFG", ancho, strict=False):
            hoja.column_dimensions[columna].width = medida
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _con_valores_en_las_formulas(contenido: bytes, calculados: dict[tuple[str, str], str]) -> bytes:
    """Escribe el valor calculado dentro de cada fórmula (openpyxl no lo hace) y fija las fechas del zip."""
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    hojas = {}
    libro_xml = entrada.read("xl/workbook.xml").decode("utf-8")
    relaciones = entrada.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    for nombre, identificador in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', libro_xml):
        destino = re.search(rf'<Relationship [^>]*Id="{identificador}"[^>]*Target="([^"]+)"', relaciones) or re.search(
            rf'<Relationship [^>]*Target="([^"]+)"[^>]*Id="{identificador}"', relaciones)
        if destino is None:
            raise ValueError(f"No se encontró la hoja {nombre} en el libro")
        hojas[("xl/" + destino.group(1).lstrip("/")).replace("xl//", "").replace("xl/xl/", "xl/")] = nombre
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as nuevo:
        for info in sorted(entrada.infolist(), key=lambda i: i.filename):
            datos_del_archivo = entrada.read(info.filename)
            if info.filename == "docProps/core.xml":
                fecha = FECHA_FIJA.strftime("%Y-%m-%dT%H:%M:%SZ")
                datos_del_archivo = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + fecha.encode(),
                                           datos_del_archivo)
            if info.filename in hojas:
                xml = datos_del_archivo.decode("utf-8")
                for (hoja, celda), valor in calculados.items():
                    if hoja != hojas[info.filename]:
                        continue
                    patron = rf'(<c r="{celda}"[^>]*>)(<f>[^<]*</f>)<v\s*/>'
                    xml, cambios = re.subn(patron, rf"\g<1>\g<2><v>{valor}</v>", xml)
                    if cambios != 1:
                        raise ValueError(f"No se pudo fijar el valor de {hoja}!{celda}")
                datos_del_archivo = xml.encode("utf-8")
            marca = zipfile.ZipInfo(info.filename, date_time=FECHA_FIJA.timetuple()[:6])
            marca.compress_type = zipfile.ZIP_DEFLATED
            marca.external_attr = info.external_attr
            nuevo.writestr(marca, datos_del_archivo)
    return salida.getvalue()


def generar_c3() -> bytes:
    """El libro Excel: cuatro hojas (una oculta), fórmulas con su valor calculado, un comentario y unidades."""
    return _con_valores_en_las_formulas(_libro_con_formulas(), valores_calculados())


GENERADORES = {datos.CASO_1: generar_c1, datos.CASO_2: generar_c2, datos.CASO_3: generar_c3}
