"""Cobertura de cifras del piloto: cuántas cifras de los adjuntos y del correo llegan a la tabla de Work (PLAN_PILOTO_LECTURA §5).

Solo lee: la base de ANKA (en una copia en memoria, de solo lectura) y los textos que el flujo ya guarda de cada hilo
en `manifiestos/`. No abre Gmail, no escribe en la base ni modifica ningún caso. El detalle (que contiene frases de los
clientes) se escribe en una hoja de Excel fuera del repositorio y en la pantalla solo salen conteos.

    python -m laboratorio.piloto.cobertura [--salida CARPETA] [hilo ...]
"""

from __future__ import annotations

import argparse
import sqlite3
import sys
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from anka.dominio.areas_calculadas import AVISO_DE_DATO
from anka.dominio.caso import load_analysis
from anka.dominio.cifras_de_pdf import cifras_aun_sin_mencionar, cifras_de_alcance
from anka.infra import instantaneas_de_manifiesto, settings

CLASES = ("real", "contexto", "ruido")   # cómo clasifica una persona cada cifra que falta
_ENCABEZADO_DE_CIFRAS = ("Hilo", "Archivo", "Cifra", "Dónde dice", "¿La tabla la menciona?", "Clasificación (persona)")


def copia_de_la_base(ruta: Path) -> sqlite3.Connection:
    """Una copia en memoria de la base, hecha con una conexión de solo lectura (respeta el modo WAL del dashboard)."""
    origen = sqlite3.connect(f"file:{ruta.as_posix()}?mode=ro", uri=True, timeout=10)
    try:
        copia = sqlite3.connect(":memory:")
        origen.backup(copia)
    finally:
        origen.close()
    copia.row_factory = sqlite3.Row
    return copia


def _sin_lineas_calculadas(texto: str) -> str:
    """Quita lo que calculó el programa (áreas de un KML): no son cifras que dijo el cliente."""
    return "\n".join(linea for linea in texto.splitlines() if AVISO_DE_DATO not in linea)


def cifras_del_hilo(textos: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Las cifras de alcance del correo y de cada adjunto con texto, sin repetir entre archivos."""
    vistas: list[tuple[set[float], str]] = []
    cifras = cifras_de_alcance("correo", _sin_lineas_calculadas(str(textos.get("correo") or "")), vistas)
    for adjunto in textos.get("adjuntos") or []:
        if isinstance(adjunto, Mapping) and adjunto.get("text"):
            cifras += cifras_de_alcance(str(adjunto.get("filename") or "?"),
                                        _sin_lineas_calculadas(str(adjunto["text"])), vistas)
    return [dict(c) for c in cifras]


def medir_hilo(analisis: Mapping[str, Any], textos: Mapping[str, Any]) -> dict[str, Any]:
    """La cobertura de un hilo: cifras totales, las que la tabla de Work no menciona y la fidelidad que ya comprobó T-38."""
    actividades = [a for a in analisis.get("activities") or [] if isinstance(a, Mapping)]
    cifras = cifras_del_hilo(textos)
    faltan = cifras_aun_sin_mencionar(cifras, actividades)
    ids_que_faltan = {id(c) for c in faltan}
    evidencia = analisis.get("evidence_check") or {}
    adjuntos = [a for a in textos.get("adjuntos") or [] if isinstance(a, Mapping)]
    return {
        "actividades": len(actividades), "cifras": len(cifras), "cifras_que_faltan": len(faltan),
        "cobertura": None if not cifras else 1 - len(faltan) / len(cifras),
        "detalle": [{**c, "en_tabla": id(c) not in ids_que_faltan} for c in cifras],
        "adjuntos": len(adjuntos), "adjuntos_no_legibles": sum(a.get("coverage") == "no_legible" for a in adjuntos),
        "adjuntos_parciales": sum(a.get("coverage") == "parcial" for a in adjuntos),
        "comprobadas": int(evidencia.get("comprobadas") or 0), "diferencias": len(evidencia.get("diferencias") or []),
        "no_comprobables": len(evidencia.get("no_comprobables") or []),
        "estado_t38": str(evidencia.get("estado") or "sin comprobar"),
    }


def medir(base: sqlite3.Connection, directorio_de_textos: Path,
          hilos: Sequence[str] = ()) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    """Mide los hilos pedidos (o todos los de la base) y devuelve (mediciones, hilos omitidos con su motivo)."""
    mediciones: dict[str, dict[str, Any]] = {}
    omitidos: dict[str, str] = {}
    for fila in base.execute("SELECT thread_id, analysis_json FROM quotes ORDER BY first_seen_at"):
        hilo = str(fila["thread_id"])
        if hilos and hilo not in hilos:
            continue
        analisis = load_analysis(fila["analysis_json"])
        textos = instantaneas_de_manifiesto.cargar_textos(directorio_de_textos, hilo)
        if not analisis.get("activities"):
            omitidos[hilo] = "sin tabla de actividades"
        elif textos is None:
            omitidos[hilo] = "sin textos guardados del hilo"
        elif textos.get("incompleto"):
            omitidos[hilo] = "textos incompletos (un adjunto no se pudo bajar)"
        else:
            mediciones[hilo] = medir_hilo(analisis, textos)
    for hilo in hilos:
        if hilo not in mediciones and hilo not in omitidos:
            omitidos[hilo] = "no está en la base"
    return mediciones, omitidos


def resumen(mediciones: Mapping[str, Mapping[str, Any]], omitidos: Mapping[str, str]) -> list[str]:
    """Las líneas que se imprimen: solo conteos, nunca texto de un cliente."""
    cifras = sum(m["cifras"] for m in mediciones.values())
    faltan = sum(m["cifras_que_faltan"] for m in mediciones.values())
    lineas = [f"Hilos medidos: {len(mediciones)} | omitidos: {len(omitidos)}"]
    for motivo in sorted(set(omitidos.values())):
        lineas.append(f"  omitidos por «{motivo}»: {sum(v == motivo for v in omitidos.values())}")
    if cifras:
        lineas.append(f"Cifras con unidad de alcance: {cifras} | la tabla de Work no menciona: {faltan} "
                      f"| cobertura global: {1 - faltan / cifras:.0%} (aún sin clasificar real, contexto o ruido)")
    lineas.append(f"Fidelidad (T-38): {sum(m['comprobadas'] for m in mediciones.values())} cifras comprobadas, "
                  f"{sum(m['diferencias'] for m in mediciones.values())} con diferencia, "
                  f"{sum(m['no_comprobables'] for m in mediciones.values())} no comprobables")
    lineas.append(f"Adjuntos: {sum(m['adjuntos'] for m in mediciones.values())} | no legibles: "
                  f"{sum(m['adjuntos_no_legibles'] for m in mediciones.values())} | parciales: "
                  f"{sum(m['adjuntos_parciales'] for m in mediciones.values())}")
    return lineas


def escribir_hoja(ruta: Path, mediciones: Mapping[str, Mapping[str, Any]]) -> None:
    """La hoja del piloto: un resumen por hilo y el detalle de las cifras, con una columna para que una persona las clasifique."""
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Resumen"
    hoja.append(["Hilo", "Actividades", "Cifras", "No mencionadas", "Cobertura", "T-38 comprobadas", "Diferencias",
                 "No comprobables", "Adjuntos", "No legibles", "Parciales"])
    for hilo, m in mediciones.items():
        hoja.append([hilo, m["actividades"], m["cifras"], m["cifras_que_faltan"],
                     "" if m["cobertura"] is None else round(m["cobertura"], 3), m["comprobadas"], m["diferencias"],
                     m["no_comprobables"], m["adjuntos"], m["adjuntos_no_legibles"], m["adjuntos_parciales"]])
    detalle = libro.create_sheet("Cifras")
    detalle.append(list(_ENCABEZADO_DE_CIFRAS))
    for hilo, m in mediciones.items():
        for cifra in m["detalle"]:
            detalle.append([hilo, cifra["archivo"], cifra["cifra"], cifra["contexto"],
                            "sí" if cifra["en_tabla"] else "no", ""])
    validacion = DataValidation(type="list", formula1='"' + ",".join(CLASES) + '"', allow_blank=True)
    detalle.add_data_validation(validacion)
    validacion.add(f"F2:F{max(detalle.max_row, 2)}")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


def clasificadas(ruta: Path) -> dict[str, int]:
    """Cuenta las cifras que una persona clasificó en la hoja (solo las que la tabla no menciona)."""
    cuentas = dict.fromkeys(CLASES, 0)
    libro = load_workbook(ruta, read_only=True)
    try:
        for fila in libro["Cifras"].iter_rows(min_row=2, values_only=True):
            clase = str(fila[5] or "").strip()
            if fila[4] == "no" and clase in cuentas:
                cuentas[clase] += 1
    finally:
        libro.close()   # en Windows un archivo abierto no se puede mover ni borrar
    return cuentas


def principal(argumentos: Iterable[str] | None = None) -> int:
    """Punto de entrada: mide, imprime los conteos y deja la hoja en la carpeta de salida."""
    parser = argparse.ArgumentParser(description="Cobertura de cifras del piloto (solo lectura).")
    parser.add_argument("hilos", nargs="*", help="Hilos a medir; sin ninguno, todos los de la base.")
    parser.add_argument("--salida", type=Path, default=settings.DATA_DIR.parent / "AnkaCorreos_piloto",
                        help="Carpeta de la hoja con el detalle (fuera del repositorio).")
    parser.add_argument("--datos", type=Path, default=settings.DATA_DIR, help="Carpeta de datos de ANKA.")
    parser.add_argument("--clasificadas", type=Path, help="Cuenta lo que una persona clasificó en una hoja ya llena.")
    args = parser.parse_args(list(argumentos) if argumentos is not None else None)
    if args.clasificadas:
        cuentas = clasificadas(args.clasificadas)
        total = sum(cuentas.values())
        print("Cifras clasificadas:", ", ".join(f"{k}: {v}" for k, v in cuentas.items()))
        if total:
            print(f"Precisión de la ayuda (real ÷ clasificadas): {cuentas['real'] / total:.0%}")
        return 0
    base_de_datos = args.datos / "quotes.sqlite3"
    if not base_de_datos.is_file():
        print(f"No encuentro la base en {base_de_datos}.", file=sys.stderr)
        return 1
    mediciones, omitidos = medir(copia_de_la_base(base_de_datos), args.datos / "manifiestos", args.hilos)
    print("\n".join(resumen(mediciones, omitidos)))
    if mediciones:
        hoja = args.salida / "cobertura_piloto.xlsx"
        escribir_hoja(hoja, mediciones)
        print(f"Detalle (con frases de clientes, solo en este equipo): {hoja}")
    return 0


if __name__ == "__main__":
    sys.exit(principal())
